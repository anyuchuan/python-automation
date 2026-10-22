"""只读加载修订版Excel；风险边界、筛选与局部图查询均在此实现。"""
from __future__ import annotations
import csv
import hashlib
import io
import json
from collections import Counter, defaultdict, deque
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
STATUS_ORDER=('直接支撑','部分支撑','未见明确支撑','待人工审核')
EXPECTED_COUNTS={'直接支撑':1,'部分支撑':18,'未见明确支撑':14,'待人工审核':1}
MAIN_TYPES=('岗位候选样本','典型工作任务（摘要）','职业能力建议','教学项目','知识条目','实训任务')
FILE_NAMES=('知识图谱节点表.xlsx','知识图谱关系表.xlsx','职业能力课程匹配矩阵.xlsx','课程能力覆盖与缺口分析.xlsx')


class DataError(ValueError):
    pass


def rows(path:Path,sheet:str,required=()):
    w=load_workbook(path,data_only=False,read_only=True)
    try:
        if sheet not in w.sheetnames:raise DataError(f'{path.name} 缺少修订版工作表：{sheet}')
        it=w[sheet].values;header=next(it)
        if len(set(header))!=len(header) or not set(required)<=set(header):raise DataError(f'{path.name}/{sheet}字段不完整')
        return [dict(zip(header,r)) for r in it if any(v is not None for v in r)]
    finally:w.close()


def attributes(node):
    try:
        data=json.loads(node.get('来源属性JSON') or '{}')
        if not isinstance(data,dict):raise ValueError()
        return data
    except (ValueError,TypeError) as exc:raise DataError('节点属性JSON无效：'+str(node.get('节点ID'))) from exc


def validate_records(nodes,edges,isolated,capabilities,coverage,project_cells):
    ids=[n.get('节点ID') for n in nodes]
    if any(not i for i in ids) or len(set(ids))!=len(ids):raise DataError('节点ID为空或重复')
    idset=set(ids)
    edgeids=[e.get('关系ID') for e in edges+isolated]
    if any(not i for i in edgeids) or len(set(edgeids))!=len(edgeids):raise DataError('关系ID为空或重复')
    for e in edges+isolated:
        if e.get('起点ID') not in idset or e.get('终点ID') not in idset:raise DataError('图谱关系端点无效：'+str(e.get('关系ID')))
    cids=[c.get('能力ID') for c in capabilities]
    if len(cids)!=34 or len(set(cids))!=34:raise DataError('必须读取修订版34项三级能力')
    counts=dict(Counter(c.get('最终候选匹配状态') for c in capabilities))
    if counts!=EXPECTED_COUNTS:raise DataError('支撑等级不是PR #3修订版1/18/14/1；禁止回退5项直接支撑版本')
    actual_leaf={n['节点ID'] for n in nodes if n['节点类型']=='职业能力建议' and attributes(n).get('能力层级')=='三级'}
    if set(cids)!=actual_leaf:raise DataError('节点表三级能力与课程匹配矩阵不一致')
    for c in capabilities:
        if not c.get('课程支撑的具体技术行为') or not c.get('课程尚未明确覆盖的工作行为'):raise DataError('缺少修订版工作行为边界列')
    stats={r['匹配状态']:r for r in coverage if r['匹配状态'] in STATUS_ORDER}
    if set(stats)!=set(STATUS_ORDER):raise DataError('覆盖统计状态不完整')
    for key,count in counts.items():
        row=stats[key]
        if row['能力数']!=count or row['统计分母']!=34 or abs(row['比例']-count/34)>1e-9:raise DataError('覆盖统计数量、比例或分母不一致')
    projects={n['节点ID'] for n in nodes if n['节点类型']=='教学项目'}
    expected={(c,p) for c in cids for p in projects}
    actual=[(c.get('能力ID'),c.get('课程项目ID')) for c in project_cells]
    if len(projects)!=8 or len(actual)!=272 or len(set(actual))!=272 or set(actual)!=expected:raise DataError('34×8项目单元不完整')


@dataclass(frozen=True)
class Options:
    explore:bool=False
    high_risk:bool=False
    isolated:bool=False


@dataclass
class Bundle:
    nodes:list[dict]
    edges:list[dict]
    isolated:list[dict]
    capabilities:list[dict]
    matches:list[dict]
    coverage:list[dict]
    project_cells:list[dict]
    improvements:list[dict]
    evidence:list[dict]
    files:dict[str,Path]
    provenance:dict

    def __post_init__(self):
        self.by_id={n['节点ID']:n for n in self.nodes}
        self.attrs={n['节点ID']:attributes(n) for n in self.nodes}
        self.cap_by_id={c['能力ID']:c for c in self.capabilities}
        self.project_by_id={n['节点ID']:n for n in self.nodes if n['节点类型']=='教学项目'}
        self.evidence_by_id={e['证据ID']:e for e in self.evidence}


def load_bundle(root:Path=ROOT,verify_fingerprints=True)->Bundle:
    root=Path(root).resolve();source=root/'phase3/output';files={}
    for name in FILE_NAMES:
        found=list(source.rglob(name)) if source.exists() else []
        if len(found)!=1:raise DataError(f'第三阶段当前输出无法唯一读取：{name}；不搜索修订前档案')
        files[name]=found[0]
    n=rows(files[FILE_NAMES[0]],'全部节点',['节点ID','节点类型','来源属性JSON'])
    e=rows(files[FILE_NAMES[1]],'候选图谱关系',['关系ID','起点ID','终点ID','已人工确认'])
    q=rows(files[FILE_NAMES[1]],'隔离关系不作确认',['关系ID','起点ID','终点ID'])
    c=rows(files[FILE_NAMES[2]],'34项能力汇总',['能力ID','最终候选匹配状态','课程支撑的具体技术行为','课程尚未明确覆盖的工作行为'])
    m=rows(files[FILE_NAMES[2]],'多对多匹配明细',['能力ID','课程项目ID','知识点ID','匹配状态'])
    p=rows(files[FILE_NAMES[2]],'272项目单元复核',['能力ID','课程项目ID','项目层匹配等级'])
    s=rows(files[FILE_NAMES[3]],'能力状态统计',['匹配状态','能力数','统计分母','比例'])
    improve=rows(files[FILE_NAMES[3]],'未见支撑项建议分类',['能力ID','未见支撑建议分类','分类性质'])
    ev=rows(files[FILE_NAMES[1]],'PDF原文证据',['证据ID','PDF页码','原文摘录'])
    validate_records(n,e,q,c,s,p)
    manifest=json.loads((Path(__file__).parent/'input_snapshot.json').read_text())
    if not manifest.get('pr3_merged'):raise DataError('来源未确认合并PR #3修订版')
    expected={x['path']:x['sha'] for x in manifest['files']};fingerprints={}
    for name,path in files.items():
        raw=path.read_bytes();relative=path.relative_to(root).as_posix();git=hashlib.sha1(b'blob '+str(len(raw)).encode()+b'\0'+raw).hexdigest()
        fingerprints[relative]=hashlib.sha256(raw).hexdigest()
        if verify_fingerprints and expected.get(relative)!=git:raise DataError('研究文件与锁定main快照不同，请审阅并更新app/input_snapshot.json，禁止静默回退：'+relative)
    for edge in e+q:
        for eid in filter(None,str(edge.get('课程证据ID') or '').split(';')):
            if eid not in {x['证据ID'] for x in ev}:raise DataError('课程证据ID不存在')
    return Bundle(n,e,q,c,m,s,p,improve,ev,files,{**manifest,'loaded_sha256':fingerprints})


def edge_risk(edge,bundle:Bundle)->str:
    """高风险是展示策略；不是对研究真实性的新核验或置信分数。"""
    kind=str(edge.get('关系类型') or '');state=str(edge.get('审核或置信状态') or '')
    if '疑似' in state:return '含历史疑似任务能力关系，未核验'
    if '归档' in kind or '旧编码' in kind:return '历史原能力或编码研究映射'
    if kind=='任务候选关联新能力' and '研究推断' in str(edge.get('招聘证据类型') or ''):return '任务到能力含研究推断，未确认'
    if kind.startswith('能力候选') and ('；待人工审核' in state or '遮挡' in str(edge.get('课程证据可见性') or '')):return '课程实施落点或文本证据尚待判断'
    return ''


def node_visible(node,bundle,options):
    id=node['节点ID'];a=bundle.attrs[id]
    isolated_node='遮挡' in str(node.get('证据或核验状态') or '') or str(a.get('实训类型') or '').startswith('单列实践表候选')
    if isolated_node:return options.explore and options.isolated
    if node['节点类型'] in {'原职业能力（归档）','课程建议能力（无招聘证据）'}:return options.explore and options.high_risk
    return True


def visible_data(bundle,options=Options()):
    nodes=[n for n in bundle.nodes if node_visible(n,bundle,options)];ids={n['节点ID'] for n in nodes}
    edges=[]
    for source,isolated in [(bundle.edges,False),(bundle.isolated,True)]:
        if isolated and not (options.explore and options.isolated):continue
        for edge in source:
            if edge['起点ID'] not in ids or edge['终点ID'] not in ids:continue
            risk=edge_risk(edge,bundle)
            if risk and not isolated and not (options.explore and options.high_risk):continue
            edges.append({**edge,'展示风险说明':('隔离关系，未确认；' if isolated else '')+risk,'展示分组':'隔离关系' if isolated else '高风险推断' if risk else '默认候选关系'})
    return nodes,edges


def neighbors(ids,edges,hops=1):
    found=set(ids);adj=defaultdict(set)
    for e in edges:
        adj[e['起点ID']].add(e['终点ID']);adj[e['终点ID']].add(e['起点ID'])
    for _ in range(hops):found|={n for i in tuple(found) for n in adj[i]}
    return found


def project_scope(bundle,project_id,edges):
    if not project_id:return set(bundle.by_id)
    if project_id not in bundle.project_by_id:return set()
    seeds={project_id}|{id for id,a in bundle.attrs.items() if a.get('课程项目ID')==project_id}
    caps={r['能力ID'] for r in bundle.project_cells if r['课程项目ID']==project_id and r['项目层匹配等级'] in {'直接支撑','部分支撑'}}
    seeds|=caps
    tasks={e['起点ID'] for e in edges if e['关系类型']=='任务候选关联新能力' and e['终点ID'] in caps}
    jobs={e['起点ID'] for e in edges if e['关系类型']=='岗位摘要包含任务' and e['终点ID'] in tasks}
    ancestors={a for cap in caps for a in [cap[:2],cap[:3]] if a in bundle.by_id}
    return seeds|tasks|jobs|ancestors


def search_nodes(nodes,query):
    query=query.strip().casefold()
    if not query:return list(nodes)
    return [n for n in nodes if query in (' '.join(str(n.get(k) or '') for k in ['节点ID','节点名称','节点类型','源文件或条目','来源属性JSON'])).casefold()]


def filter_graph(bundle,options=Options(),types=None,levels=None,project=None,statuses=None,center=None,hops=1,max_nodes=70,show_peer_abilities=False,query=None):
    nodes,edges=visible_data(bundle,options);scope=project_scope(bundle,project,edges)
    types=set(types) if types is not None else None;levels=set(levels) if levels is not None else None
    statuses=set(statuses) if statuses is not None else set(STATUS_ORDER)
    if not statuses:return [],[],{'匹配节点数':0,'绘制节点数':0,'已截断':False}
    if statuses!=set(STATUS_ORDER):
        caps={id for id,c in bundle.cap_by_id.items() if c['最终候选匹配状态'] in statuses}
        scope&=neighbors(caps,edges,2)
    selected=[]
    for n in nodes:
        id=n['节点ID'];attrs=bundle.attrs[id]
        if id not in scope or types is not None and n['节点类型'] not in types:continue
        if levels is not None and n['节点类型'] in {'职业能力建议','原职业能力（归档）','课程建议能力（无招聘证据）'} and (attrs.get('能力层级') or attrs.get('原能力层级')) not in levels:continue
        if id in bundle.cap_by_id and bundle.cap_by_id[id]['最终候选匹配状态'] not in statuses:continue
        if center in bundle.cap_by_id and not show_peer_abilities and id in bundle.cap_by_id and id!=center:continue
        selected.append(n)
    if query and not center:selected=search_nodes(selected,query)
    ids={n['节点ID'] for n in selected};edges=[e for e in edges if e['起点ID'] in ids and e['终点ID'] in ids]
    if center:
        if center not in ids:return [],[],{'匹配节点数':0,'绘制节点数':0,'已截断':False}
        ids&=neighbors({center},edges,hops)
    ordered=sorted(ids,key=lambda id:(0 if id==center else 1,id))
    # 邻近中心的节点优先，截断有提示且下载与画面完全一致。
    if center:
        adj=defaultdict(set)
        for e in edges:adj[e['起点ID']].add(e['终点ID']);adj[e['终点ID']].add(e['起点ID'])
        dist={center:0};queue=deque([center])
        while queue:
            current=queue.popleft()
            for id in sorted(adj[current]):
                if id not in dist:dist[id]=dist[current]+1;queue.append(id)
        ordered.sort(key=lambda id:(dist.get(id,999),id))
    total=len(ordered);ordered=ordered[:max(1,max_nodes)];ids=set(ordered)
    result=[bundle.by_id[i] for i in ordered];edges=[e for e in edges if e['起点ID'] in ids and e['终点ID'] in ids]
    return result,edges,{'匹配节点数':total,'绘制节点数':len(result),'已截断':total>len(result)}


def detail(bundle,node_id,edges):
    node=bundle.by_id[node_id];attrs=bundle.attrs[node_id];c=bundle.cap_by_id.get(node_id,{})
    pages=set(filter(None,str(node.get('PDF页码') or '').split(';')))
    for m in bundle.matches:
        if m['能力ID']==node_id:pages.update(filter(None,str(m.get('PDF页码') or '').split(';')))
    related=sorted({e['终点ID'] if e['起点ID']==node_id else e['起点ID'] for e in edges if node_id in {e['起点ID'],e['终点ID']}})
    return {'编码':node_id,'名称':node['节点名称'],'类型':node['节点类型'],'来源':node['源文件或条目'],
            '证据类型':attrs.get('证据类型') or attrs.get('课程条目定位') or ('岗位摘要支持' if node['节点类型'] in MAIN_TYPES[:2] else '课程文本/研究结构化，见源属性'),
            '课程引用物理页码':';'.join(sorted(pages,key=lambda x:int(x) if x.isdigit() else 999)) or '未单列；见关联课程引文',
            '证据状态':node['证据或核验状态'],'审核状态':node['人工审核状态'],'课程支撑等级':c.get('最终候选匹配状态','本节点未单独定级'),
            '课程支撑的具体技术行为':c.get('课程支撑的具体技术行为',''),
            '课程尚未明确覆盖的工作行为':c.get('课程尚未明确覆盖的工作行为',''),
            '课程判断理由':c.get('课程支撑判断',''),'相关节点':related,'源属性':attrs}


def edge_style(edge):
    state=str(edge.get('审核或置信状态') or '')
    if edge.get('展示分组')=='隔离关系' or edge.get('展示风险说明'):return '高风险/隔离','#bb4c62','3 5'
    if '部分支撑' in state:return '部分支撑','#c68a28','8 4'
    if '；直接支撑' in state:return '直接支撑','#12806b',''
    if '能力候选' in str(edge.get('关系类型')):return '待审核','#8b68b0','2 5'
    if '研究推断' in str(edge.get('招聘证据类型')) or '候选' in str(edge.get('关系类型')):return '研究推断','#8b68b0','5 5'
    return '原文结构/摘要关联','#9ba7b1',''


def csv_bytes(records:Iterable[dict]):
    records=list(records)
    if not records:return '\ufeff'.encode('utf-8')
    fields=list(dict.fromkeys(k for r in records for k in r if not k.startswith('_')))
    out=io.StringIO(newline='');writer=csv.DictWriter(out,fieldnames=fields,extrasaction='ignore');writer.writeheader()
    for row in records:
        safe={k:("'"+str(v) if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for k,v in row.items() if k in fields}
        writer.writerow(safe)
    return out.getvalue().encode('utf-8-sig')
