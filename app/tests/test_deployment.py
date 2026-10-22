import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl.reader.excel import load_workbook as real_load_workbook
from app.data import ROOT,load_bundle

class DeploymentTests(unittest.TestCase):
 def test_loader_works_without_repository_working_directory(self):
  before=Path.cwd()
  with tempfile.TemporaryDirectory() as directory:
   try:
    os.chdir(directory);bundle=load_bundle();self.assertEqual(len(bundle.capabilities),34)
   finally:os.chdir(before)
 def test_only_existing_excel_inputs_no_pdf_or_research_pipeline(self):
  opened=[]
  def track(path,*args,**kwargs):
   path=Path(path).resolve();self.assertTrue(path.is_relative_to(ROOT/'phase3/output'))
   self.assertEqual(path.suffix,'.xlsx');opened.append(path)
   return real_load_workbook(path,*args,**kwargs)
  with patch('app.data.load_workbook',side_effect=track):bundle=load_bundle()
  self.assertEqual(set(opened),set(bundle.files.values()))
 def test_root_dependency_forwarding_and_binary_install(self):
  root=(ROOT/'requirements.txt').read_text();app=(ROOT/'app/requirements.txt').read_text()
  self.assertIn('-r app/requirements.txt',root);self.assertIn('--only-binary=:all:',app)
  self.assertNotIn('pdfplumber',app);self.assertNotIn('PyMuPDF',app)
 def test_cached_load_runs_once_across_ui_rerun(self):
  from streamlit.testing.v1 import AppTest
  with patch('app.data.load_bundle',wraps=load_bundle) as loader:
   at=AppTest.from_file(str(ROOT/'app/main.py'),default_timeout=30).run()
   self.assertFalse(at.exception);calls=loader.call_count
   at.run();self.assertEqual(loader.call_count,calls)
   self.assertLessEqual(calls,1)
 def test_diagnostic_from_external_working_directory(self):
  with tempfile.TemporaryDirectory() as directory:
   r=subprocess.run([sys.executable,str(ROOT/'app/check_deployment.py')],cwd=directory,capture_output=True,text=True,timeout=30)
  self.assertEqual(r.returncode,0,r.stderr);self.assertIn('"nodes": 314',r.stdout)

if __name__=='__main__':unittest.main()
