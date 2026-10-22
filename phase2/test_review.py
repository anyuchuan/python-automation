"""第二阶段研究边界、溯源和输出完整性测试。"""
import copy
import hashlib
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('phase2_review', ROOT/'phase2/review.py')
review = importlib.util.module_from_spec(spec)
spec.loader.exec_module(review)


class ReviewTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.temp.name)/'result'
        cls.hashes = {p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in review.INPUTS}
        cls.stats = review.run(ROOT,cls.out)
        cls.jobs,cls.tasks,cls.old,cls.originals,cls.duplicates = review.load_inputs(ROOT)
        cls.rules = json.loads((ROOT/'phase2/rules.json').read_text())
        cls.specs = {r['code']:r for r in cls.rules['leaves']}
        cls.checked,cls.proposals = review.relation_review(cls.jobs,cls.tasks,cls.originals,cls.rules)
        cls.nodes,cls.course = review.build_nodes(cls.rules,cls.proposals)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_actual_counts_and_snapshot(self):
        self.assertEqual(self.stats['实际输入数量'], {'岗位':50,'任务':79,'一级':5,'二级':15,'三级':26,'关系':188,'聚合页一级':5,'聚合页二级':15})
        self.assertEqual(self.stats['输入快照差异数'],0)
        self.assertEqual(self.stats['原三级优化动作数量'],{'拆分':9,'保留并规范':15,'合并':2})
        self.assertEqual(self.stats['优化三级数'],34)

    def test_all_four_workbooks_open_and_cells_nonformula(self):
        files = list(self.out.glob('*.xlsx'))
        self.assertEqual(len(files),4)
        for path in files:
            wb = load_workbook(path,data_only=False)
            self.assertTrue(wb.sheetnames)
            for ws in wb:
                self.assertGreaterEqual(ws.max_row,1)
                self.assertEqual(ws.freeze_panes,'A2')
                for row in ws:
                    self.assertFalse(any(c.data_type=='f' for c in row))
            wb.close()

    def test_first_stage_inputs_unchanged(self):
        self.assertEqual(self.hashes,{p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in review.INPUTS})
        for path in [ROOT/'output',ROOT/'output/subdir',ROOT/'scripts',ROOT/'tests/x']:
            with self.assertRaisesRegex(ValueError,'禁止'):
                review.run(ROOT,path)

    def test_all_old_codes_and_relations_preserved(self):
        mapping,audits = review.compare(self.old,self.rules,self.nodes)
        self.assertEqual(len(audits),46)
        self.assertEqual({a['原编码'] for a in audits},{n['code'] for n in review.old_hierarchy(self.old)})
        self.assertEqual(len(self.checked),188)
        self.assertEqual({r['原关系ID'] for r in self.checked},{r['原关系ID'] for r in self.proposals})
        for old, checked in zip(self.originals,self.checked):
            for key,value in old.items():
                self.assertEqual(checked[key],value)

    def test_every_relation_traces_to_original_task_job_and_ability(self):
        tasks={t['任务编码']:t for t in self.tasks}
        jobs={j['样本ID']:j for j in self.jobs}
        originals={r['原关系ID']:r for r in self.checked}
        for r in self.proposals:
            self.assertIn(r['新三级能力编码'],self.specs)
            self.assertEqual(tasks[r['任务ID']]['岗位ID'],r['岗位ID'])
            self.assertIn(r['任务证据'],jobs[r['岗位ID']]['岗位职责摘要（转述）'])
            self.assertEqual(r['任务证据'],tasks[r['任务ID']]['证据文本'])
            self.assertEqual(originals[r['原关系ID']]['三级能力编码'],r['原三级能力编码'])
            self.assertEqual(r['原始招聘链接'],jobs[r['岗位ID']]['原始招聘链接'])

    def test_merge_deduplicates_support_counts(self):
        query=next(n for n in self.nodes if n['新编码']=='N113')
        direct=[r for r in self.proposals if r['新三级能力编码']=='N113' and r['证据类型']=='岗位摘要支持']
        self.assertEqual(query['原三级能力编码'],'C113;C131')
        self.assertEqual(query['摘要支持岗位数（含疑似重复）'],len({r['岗位ID'] for r in direct}))
        self.assertEqual(query['摘要支持任务数'],len({r['任务ID'] for r in direct}))
        self.assertLess(query['摘要支持任务数'],len(direct))

    def test_split_does_not_invent_restore_from_backup(self):
        self.assertTrue(review.summary_support(self.specs['N321'],'数据备份'))
        self.assertFalse(review.summary_support(self.specs['N322'],'数据备份'))
        self.assertTrue(review.summary_support(self.specs['N322'],'数据库备份恢复'))
        self.assertFalse(review.summary_support(self.specs['N321'],'系统监控、备份与维护'))

    def test_flow_and_report_not_database_application_or_document(self):
        self.assertFalse(review.summary_support(self.specs['N521'],'优化数据提取分析流程'))
        self.assertTrue(review.summary_support(self.specs['N522'],'优化数据提取分析流程'))
        self.assertFalse(review.summary_support(self.specs['N412'],'开发数据报表'))
        self.assertFalse(review.summary_support(self.specs['N311'],'数据库日常运维'))

    def test_generic_analysis_and_statistics_do_not_imply_sql_or_design(self):
        self.assertFalse(review.summary_support(self.specs['N113'],'统计分析客流数据'))
        self.assertFalse(review.summary_support(self.specs['N211'],'每周统计公园运营数据'))
        self.assertFalse(review.summary_support(self.specs['N411'],'需求调研'))
        job={'任职要求摘要（不含学历经验）':'未明确SQL；SQL Server数据库管理'}
        t={'证据文本':'分析经营数据'}
        self.assertEqual(review.sql_qualifier(t,job),'未获得该任务使用SQL的明确摘要证据')
        job['任职要求摘要（不含学历经验）']='熟练SQL'
        self.assertIn('任务使用方式待核验',review.sql_qualifier(t,job))
        self.assertFalse(review.summary_support(self.specs['N113'],t['证据文本']))

    def test_no_original_evidence_or_approval_fabricated(self):
        self.assertEqual(self.stats['招聘原文明示关系数'],0)
        for r in self.proposals:
            self.assertNotEqual(r['证据类型'],'招聘原文明示')
            self.assertEqual(r['本轮岗位状态'],'待核验')
            self.assertEqual(r['人工审核状态'],'待人工审核')
        for n in self.nodes+self.course:
            self.assertEqual(n['已核验招聘原文岗位数'],0)
            self.assertEqual(n['人工审核状态'],'待人工审核')

    def test_course_suggestions_have_zero_links_and_frequencies(self):
        for s in self.course:
            self.assertFalse(s['对应任务ID'])
            self.assertFalse(s['来源岗位ID'])
            self.assertEqual(s['摘要支持岗位数（含疑似重复）'],0)
            self.assertEqual(s['摘要支持任务数'],0)
            self.assertEqual(s['招聘需求频次'],'不计入（课程建议）')
        self.assertEqual({s['新编码'] for s in self.course if s['能力层级']=='三级'},{'S111','S112','S113'})
        self.assertFalse({r['新三级能力编码'] for r in self.proposals}&{s['新编码'] for s in self.course})

    def test_every_node_has_required_fields_and_valid_parent(self):
        required={'新编码','能力层级','能力名称','可观察的工作行为','对应任务ID','来源岗位ID','证据类型','证据状态','课程适配类别','优化理由','人工审核状态'}
        ids={n['新编码'] for n in self.nodes+self.course}
        self.assertEqual(len(ids),len(self.nodes+self.course))
        taskids={t['任务编码'] for t in self.tasks}
        jobids={j['样本ID'] for j in self.jobs}
        for n in self.nodes+self.course:
            self.assertTrue(required<=n.keys())
            self.assertTrue(n['父编码']=='' or n['父编码'] in ids)
            self.assertTrue(set(filter(None,n['对应任务ID'].split(';')))<=taskids)
            self.assertTrue(set(filter(None,n['来源岗位ID'].split(';')))<=jobids)
            self.assertIn(n['课程适配类别'],review.COURSE.values())

    def test_parent_frequencies_are_id_unions_not_sum_of_children(self):
        for n in self.nodes:
            if n['能力层级']=='三级':
                continue
            direct=[r for r in self.proposals if r['新三级能力编码'].startswith(n['新编码']) and r['证据类型']=='岗位摘要支持']
            self.assertEqual(n['摘要支持岗位数（含疑似重复）'],len({r['岗位ID'] for r in direct}))

    def test_invalid_relation_or_rules_fail(self):
        broken=copy.deepcopy(self.rules)
        broken['leaves']=[s for s in broken['leaves'] if s['code']!='N113']
        with self.assertRaisesRegex(ValueError,'原能力编码不一致'):
            review.validate_rules(broken,self.old)
        bad=copy.deepcopy(self.proposals)
        bad[0]['任务ID']='T-NONEXISTENT'
        with self.assertRaisesRegex(ValueError,'外键'):
            review.validate_output(self.nodes,self.course,bad,self.checked,self.jobs,self.tasks)

    def test_suspected_relations_and_manual_questions_exist(self):
        issues=review.issues_list(self.jobs,self.duplicates,self.checked,review.compare(self.old,self.rules,self.nodes)[1],self.course,self.tasks)
        self.assertEqual(len({i['问题ID'] for i in issues}),len(issues))
        self.assertEqual(sum(i['问题类型']=='招聘证据核验' for i in issues),50)
        self.assertEqual(sum(i['问题类型']=='关系疑似误匹配' for i in issues),sum(r['疑似误匹配']=='是' for r in self.checked))
        self.assertTrue(all(i['人工审核状态']=='待人工审核' for i in issues))
        self.assertTrue(all(not i['审核人'] and not i['审核结论'] for i in issues))

    def test_repeat_run_same_cells_and_stats(self):
        second=Path(self.temp.name)/'rerun'
        self.assertEqual(review.run(ROOT,second),self.stats)
        for f in self.out.glob('*.xlsx'):
            a,b=load_workbook(f),load_workbook(second/f.name)
            self.assertEqual(a.sheetnames,b.sheetnames)
            for name in a.sheetnames:
                self.assertEqual(list(a[name].values),list(b[name].values))
            a.close();b.close()
        self.assertEqual((self.out/'第二阶段优化报告.md').read_text(),(second/'第二阶段优化报告.md').read_text())


if __name__=='__main__':
    unittest.main()
