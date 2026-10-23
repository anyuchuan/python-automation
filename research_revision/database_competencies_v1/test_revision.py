"""研究修订数据完整性与证据边界测试，不声称验证招聘真实性或研究有效性。"""
import copy
import importlib.util
import json
from pathlib import Path
import re
import tempfile
import unittest
from collections import Counter
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
spec=importlib.util.spec_from_file_location('research_revision_build',HERE/'build.py')
b=importlib.util.module_from_spec(spec);spec.loader.exec_module(b)

class RevisionTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory();cls.out=Path(cls.tmp.name)
  cls.summary=b.run(output=cls.out)
  cls.rules=json.loads((HERE/'rules.json').read_text())
  cls.selected,cls.inventory,cls.missing,cls.files=b.discover(b.ROOT)
  cls.ex,cls.typical,cls.audit,cls.req=b.extract_tasks(cls.selected,cls.rules)
  cls.core,cls.pending,cls.links,cls.details,cls.hier=b.build_capabilities(cls.rules,cls.ex)
 @classmethod
 def tearDownClass(cls):cls.tmp.cleanup()
 def test_inputs_detected_and_fields_missing(self):
  self.assertEqual(len(self.selected['jobs']['rows']),50);self.assertEqual(len(self.selected['tasks']['rows']),79)
  self.assertEqual(len(self.selected['old34']['rows']),34);self.assertEqual(len(self.selected['old26']['rows']),26)
  p=self.selected['jobs']['path'];rows={r['字段']:r for r in self.missing if r['文件']==p and r['工作表']=='招聘岗位主库'}
  self.assertEqual(rows['证据摘录']['缺失数'],50);self.assertEqual(rows['岗位发布日期']['缺失数'],44)
 def test_original_files_unchanged_including_app(self):
  manifest=json.loads((HERE/'input_snapshot.json').read_text())
  for f in manifest['repository_files']:self.assertEqual(b.git_sha(b.ROOT/f['path']),f['sha'],f['path'])
 def test_workbooks_open_and_nonempty(self):
  self.assertEqual(len(list(self.out.glob('*.xlsx'))),9)
  for p in self.out.glob('*.xlsx'):
   w=load_workbook(p,read_only=True,data_only=False)
   self.assertTrue(w.sheetnames)
   for s in w:self.assertGreaterEqual(sum(1 for _ in s.values),1)
   w.close()
 def test_exact_quotes_addresses_and_all_old_tasks_traced(self):
  jobs={r['样本ID']:r for r in self.selected['jobs']['rows']}
  for e in self.ex:
   self.assertIn(e['摘要证据文本'],jobs[e['岗位ID']][e['来源摘要字段']])
   self.assertEqual(self.selected['jobs']['rows'][e['源Excel行号']-2]['样本ID'],e['岗位ID'])
   self.assertEqual(e['招聘原文已核验'],'否');self.assertEqual(e['原始招聘描述'],'未提供')
  self.assertEqual(len(self.audit),79);self.assertTrue(all(r['子片段数']>=1 for r in self.audit))
  self.assertEqual(len(self.req),50);self.assertTrue(any(e['来源摘要字段'].startswith('任职') for e in self.ex))
 def test_task_ids_and_groups_cover_exactly(self):
  self.assertEqual(len({e['提取任务ID'] for e in self.ex}),len(self.ex))
  tids={t['典型工作任务ID'] for t in self.typical}
  self.assertEqual({e['归纳典型任务ID'] for e in self.ex},tids)
  self.assertEqual(sum(Counter(t['数据库专业相关性'] for t in self.typical).values()),len(self.typical))
  self.assertEqual(len(self.ex),179);self.assertEqual(len(tids),48)
 def test_core_has_explicit_summary_technical_actions(self):
  patterns={r['id']:r['support_pattern'] for r in self.rules['core_capabilities']}
  for r in self.details:
   if r['纳入核心候选']=='是':self.assertRegex(r['摘要证据文本'],patterns[r['能力/线索ID']])
  self.assertTrue(all(c['招聘原文明示证据数']==0 for c in self.core));self.assertEqual(len(self.core),5)
  altered=copy.deepcopy(self.ex)
  for e in altered:
   if e['归纳典型任务ID']=='TT005':e['摘要证据文本']='一般数据清洗'
  with self.assertRaisesRegex(ValueError,'缺明确技术'):b.build_capabilities(self.rules,altered)
 def test_no_unfounded_abilities_or_course_reverse_inference(self):
  names=';'.join(c['三级能力名称'] for c in self.core)
  for word in ['多表','完整性约束','表结构','Python','旅游服务','权限配置','记录更新']:self.assertNotIn(word,names)
  self.assertEqual(len(self.pending),6);self.assertTrue(all(c['能力编号'].startswith('R') and '不纳核心' in c['纳入范围'] for c in self.pending))
 def test_hierarchy_derived_from_core_and_valid(self):
  ids={n['节点ID'] for n in self.hier};self.assertEqual(len(ids),len(self.hier))
  self.assertEqual(Counter(n['层级'] for n in self.hier),{'一级':2,'二级':4,'三级':5})
  for n in self.hier:
   if n['父编码']:self.assertIn(n['父编码'],ids)
  self.assertEqual({n['节点ID'] for n in self.hier if n['层级']=='三级'},{c['能力编号'] for c in self.core})
 def test_all_relationship_endpoints_and_lineage(self):
  groups={g['典型工作任务ID'] for g in self.typical};caps={c['能力编号'] for c in self.core+self.pending};exids={e['提取任务ID'] for e in self.ex}
  for r in self.links:
   self.assertIn(r['典型工作任务ID'],groups);self.assertIn(r['三级能力/待审线索ID'],caps)
   self.assertTrue(set(r['提取任务ID'].split(';'))<=exids)
  logical=b.read(self.out/'07_三级能力逻辑关系表.xlsx','候选非层级关系')
  core={c['能力编号'] for c in self.core}
  for r in logical:
   self.assertIn(r['起点能力ID'],core);self.assertIn(r['终点能力ID'],core)
   self.assertIn('研究推断',r['关系类型']);self.assertTrue(r['适用条件与边界'])
   self.assertNotIn('权重',r);self.assertNotIn('关联强度',r)
 def test_duplicates_are_reported_not_dropped(self):
  jobs,des,dups=b.check_jobs(self.selected)
  self.assertEqual(len(jobs),50);self.assertEqual(len(des),100)
  pairs={r['重复组'] for r in dups if r['检查类型']=='单位与岗位同时重复'};self.assertEqual(len(pairs),2)
  raw={r['样本ID']:r for r in self.selected['jobs']['rows']}
  for j in jobs:
   self.assertEqual(j['核验状态'],raw[j['样本ID']]['核验状态']);self.assertEqual(j['证据质量（初评）'],raw[j['样本ID']]['证据质量（初评）'])
 def test_old_new_full_coverage_and_no_automatic_migration(self):
  n34,n26,course,reverse=b.compare(self.selected,self.core,self.pending)
  self.assertEqual(len(n34),34);self.assertEqual(len(n26),26)
  self.assertEqual(sum(r['操作']=='保留并收窄' for r in n34),5)
  self.assertEqual(sum(r['操作']=='排除出核心（保留原始与外围任务）' for r in n34),22)
  self.assertTrue(all(r['迁移旧任务能力关系'].startswith('否') for r in n34))
  self.assertEqual(self.summary['old_suspect_relations'],109)
 def test_semantic_reproducibility_and_forbidden_output(self):
  second=self.out/'again';again=b.run(output=second);self.assertEqual(again,self.summary)
  for name in self.summary['outputs']:
   w=load_workbook(self.out/name,read_only=True);sheets=w.sheetnames;w.close()
   for sheet in sheets:self.assertEqual(b.read(self.out/name,sheet),b.read(second/name,sheet))
  with self.assertRaisesRegex(ValueError,'禁止覆盖'):b.run(output=b.ROOT/'phase3/output')

if __name__=='__main__':unittest.main()
