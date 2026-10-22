"""只读原研究输入；在独立目录生成候选修订工作簿。仅依赖openpyxl。"""
from __future__ import annotations
import argparse
from collections import Counter,defaultdict
from difflib import SequenceMatcher
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import urlsplit,urlunsplit
from datetime import datetime
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
REVIEW='候选研究修订；待研究者人工审查；未验证、未定稿'
SUMMARY='岗位摘要支持；非招聘原文；候选岗位待核验'

def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def git_sha(path):
    raw=Path(path).read_bytes();return hashlib.sha1(b'blob '+str(len(raw)).encode()+b'\0'+raw).hexdigest()
def read(path,sheet):
    w=load_workbook(path,read_only=True,data_only=False)
    try:
        it=iter(w[sheet].values);header=next(it)
        return [dict(zip(header,r)) for r in it if any(v is not None for v in r)]
    finally:w.close()

def discover(root):
    inventory=[];fields=[];sheets={};source_files=[]
    for p in sorted(Path(root).rglob('*.xlsx')):
        if set(p.relative_to(root).parts)&{'work','research_revision','.git','.venv','revisions'}:continue
        source_files.append(p);w=load_workbook(p,read_only=True,data_only=False)
        try:
            for s in w:
                it=iter(s.values)
                try:h=next(it)
                except StopIteration:continue
                rows=[dict(zip(h,r)) for r in it if any(v is not None for v in r)]
                relative=p.relative_to(root).as_posix();sheets[(relative,s.title)]=rows
                inventory.append({'文件':relative,'工作表':s.title,'记录数':len(rows),'字段数':len(h),'字段':';'.join(map(str,h)),'SHA256':digest(p)})
                for k in h:
                    miss=sum(r.get(k) is None or str(r.get(k)).strip()=='' for r in rows)
                    fields.append({'文件':relative,'工作表':s.title,'字段':k,'记录数':len(rows),'缺失数':miss,'缺失比例':miss/len(rows) if rows else None})
        finally:w.close()
    signatures={
        'jobs':('招聘岗位主库',{'样本ID','岗位职责摘要（转述）','任职要求摘要（不含学历经验）'}),
        'tasks':('任务编码',{'任务编码','岗位ID','典型工作任务','证据文本'}),
        'old26':('三级能力指标',{'三级能力编码','三级能力'}),
        'old34':('三级岗位能力',{'新编码','父编码','能力层级'}),
        'old_all':('优化岗位能力全层级',{'新编码','父编码','能力层级'}),
        'old_course':('课程建议能力单列',{'新编码','父编码','能力层级'}),
        'review':('原关系逐条复核',{'原关系ID','疑似误匹配'})}
    selected={}
    for role,(sheet,required) in signatures.items():
        found=[(p,s) for (p,s),rows in sheets.items() if s==sheet and rows and required<=set(rows[0])]
        if len(found)!=1:raise ValueError(f'无法唯一识别输入 {role}: {found}')
        key=found[0];selected[role]={'path':key[0],'sheet':key[1],'rows':sheets[key]}
    return selected,inventory,fields,source_files

def list_ids(items):return ';'.join(sorted(set(str(i) for i in items if i)))
def normalized_link(text):
    if not text:return ''
    u=urlsplit(str(text).strip());return urlunsplit((u.scheme.lower(),u.netloc.lower(),u.path.rstrip('/'),u.query,''))

def check_jobs(selected):
    source=selected['jobs'];jobs=source['rows'];duplicates=[];statuses=[];descriptions=[];issues=[]
    for kind,fn,note in [
        ('招聘链接完全重复',lambda r:str(r.get('原始招聘链接') or '').strip(),'疑似重复；保留全部记录'),
        ('招聘链接规范化重复',lambda r:normalized_link(r.get('原始招聘链接')),'仅去片段和末尾斜杠，不删除query参数'),
        ('单位重复（不等于岗位重复）',lambda r:r.get('招聘单位'),'同单位不同职位通常不是重复样本'),
        ('岗位名称重复（不等于样本重复）',lambda r:r.get('岗位名称'),'同名岗位可能来自不同单位，不直接认定为重复样本'),
        ('单位与岗位同时重复',lambda r:(r.get('招聘单位'),r.get('岗位名称')),'疑似重复发布或不同来源，不能按独立需求计数')]:
        groups=defaultdict(list)
        for r in jobs:
            key=fn(r)
            if key:groups[key].append(r)
        for key,rs in groups.items():
            if len(rs)>1:
                for r in rs:duplicates.append({'重复组':kind+'-'+str(key),'检查类型':kind,'匹配键':str(key),'岗位ID':r['样本ID'],'招聘单位':r['招聘单位'],'岗位名称':r['岗位名称'],'原始招聘链接':r['原始招聘链接'],'说明':note,'处理':'不删除；待审查，不增加核验任务'})
    for i,r in enumerate(jobs,2):
        statuses.append({**r,'源文件':source['path'],'源工作表':source['sheet'],'源Excel行号':i,'本稿统一证据状态':'候选岗位待核验；原状态照录，未升级','招聘原文是否提供':'否；证据摘录为空，现有可用内容为摘要'})
        for field in ['岗位职责摘要（转述）','任职要求摘要（不含学历经验）']:
            descriptions.append({'描述ID':f"DESC-{r['样本ID']}-{'D' if field.startswith('岗位职责') else 'R'}",'岗位ID':r['样本ID'],'源文件':source['path'],'工作表':source['sheet'],'字段':field,'Excel行号':i,'现有摘要文本':r[field],'原始招聘描述':'未提供，不把摘要当原文','证据类型':SUMMARY,'原始招聘链接':r['原始招聘链接'],'原核验状态':r['核验状态']})
    for a_index,a in enumerate(jobs):
        for b in jobs[a_index+1:]:
            if a['招聘单位']!=b['招聘单位']:continue
            similarity=SequenceMatcher(None,str(a['岗位职责摘要（转述）']),str(b['岗位职责摘要（转述）'])).ratio()
            if similarity>=0.72 and a['岗位名称']!=b['岗位名称']:
                duplicates.append({'重复组':'同单位职责相似-'+a['样本ID']+'-'+b['样本ID'],'检查类型':'同单位职责相似（辅助候选）','匹配键':a['招聘单位'],'岗位ID':a['样本ID']+';'+b['样本ID'],'招聘单位':a['招聘单位'],'岗位名称':a['岗位名称']+' / '+b['岗位名称'],'原始招聘链接':a['原始招聘链接']+';'+b['原始招聘链接'],'说明':f'字符相似度{similarity:.3f}，阈值0.72仅用于提示，不是语义重复或独立职位判定','处理':'保留全部记录；人工判断，不进行招聘原文核验'})
    return statuses,descriptions,duplicates

def extract_tasks(selected,rules):
    jobs={r['样本ID']:r for r in selected['jobs']['rows']};old=selected['tasks']['rows'];groups={g['id']:g for g in rules['task_groups']};out=[]
    if set(rules['responsibilities'])!=set(jobs):raise ValueError('职责逐条审阅规则不覆盖全部岗位')
    for job in sorted(jobs):
        row=jobs[job]
        for key,field in [('responsibilities','岗位职责摘要（转述）'),('requirements','任职要求摘要（不含学历经验）')]:
            for seq,item in enumerate(rules[key].get(job,[]),1):
                span=item['span'];group=item['group'];text=row[field]
                if span not in text or group not in groups:raise ValueError(f'任务引文不存在: {job}/{field}/{span}')
                old_ids=[t['任务编码'] for t in old if t['岗位ID']==job and span in str(t['证据文本'])] if key=='responsibilities' else []
                noun=span in {'自动化报表','数据埋点','数据字典','可视化看板','商业分析报告','权限管理','数据备份','统计汇总','数据资源管理','备份','项目实施','OTA运营报表','文档','文旅数据资源管理','资源管理','市场调研','经营报告','数据看板','数据接入','可视化','收益管理','业务决策','建模','数据清洗','同步','SQL统计工具'}
                inference='动作/任务补足含研究推断，待审查' if noun or (job=='JOB008' and key=='requirements') else '摘要动作或对象片段；分解归纳仍为研究判断'
                out.append({'提取任务ID':f"EX-{job}-{'D' if key=='responsibilities' else 'R'}{seq:02}",'岗位ID':job,'原任务ID':list_ids(old_ids) or '无（从任职要求补提）' if key=='requirements' else list_ids(old_ids) or '无（职责补提）','原始招聘描述':'未提供','来源摘要字段':field,'摘要证据文本':span,'完整摘要上下文':text,'提取的工作任务':span,'提取活动维度（研究判断）':groups[group]['name'],'归纳典型任务ID':group,'归纳后的典型工作任务':groups[group]['name'],'任务提取判断':inference,'提取规则':'逐条审阅定位；并列动作按对象/产出拆分；同一上下文可分解为多个活动','语义合并依据':groups[group]['merge_reason'],'证据类型':SUMMARY,'招聘原文已核验':'否','原始招聘链接':row['原始招聘链接'],'原证据质量等级':row['证据质量（初评）'],'原核验状态':row['核验状态'],'源文件':selected['jobs']['path'],'源工作表':selected['jobs']['sheet'],'源Excel行号':next(i for i,r in enumerate(selected['jobs']['rows'],2) if r['样本ID']==job),'源单元格':get_column_letter(list(row).index(field)+1)+str(next(i for i,r in enumerate(selected['jobs']['rows'],2) if r['样本ID']==job)),'研究状态':REVIEW})
    # 岗位字段保持原貌；提取片段与合并任务名用不同列，避免伪装成原始描述。
    typical=[]
    for id,g in groups.items():
        members=[e for e in out if e['归纳典型任务ID']==id]
        if not members:continue
        typical.append({'典型工作任务ID':id,'典型工作任务':g['name'],'数据库专业相关性':g['db_relevance'],'相关性判断依据':g['relevance_reason'],'提取任务ID':list_ids(e['提取任务ID'] for e in members),'来源岗位ID':list_ids(e['岗位ID'] for e in members),'原任务ID':list_ids(x for e in members for x in str(e['原任务ID']).split(';') if x.startswith('T-')),'提取片段数（非需求频次）':len(members),'候选岗位数（未去疑似重复）':len({e['岗位ID'] for e in members}),'合并操作':g['merge_reason'],'证据类型':SUMMARY,'数据库相关性说明':'直接相关只表示摘要涉及数据库/SQL技术，不等于已获充分三级能力依据','研究状态':REVIEW})
    audit=[]
    for t in old:
        members=[e for e in out if t['任务编码'] in e['原任务ID'].split(';')]
        if not members:raise ValueError('旧任务未追溯：'+t['任务编码'])
        audit.append({**t,'原任务证据是否在职责摘要中':str(t['证据文本']) in str(jobs[t['岗位ID']]['岗位职责摘要（转述）']),'修订提取任务ID':list_ids(e['提取任务ID'] for e in members),'修订典型任务ID':list_ids(e['归纳典型任务ID'] for e in members),'修订建议':'保留片段并合并相似任务' if len(members)==1 else '拆分复合描述后按动作/对象归并；不补造技术要求','子片段数':len(members),'原文证据边界':'原表证据为职责摘要，非招聘原文','研究状态':REVIEW})
    req_review=[{'岗位ID':job,'任职要求摘要':r['任职要求摘要（不含学历经验）'],'补提任务ID':list_ids(e['提取任务ID'] for e in out if e['岗位ID']==job and e['来源摘要字段'].startswith('任职要求')) or '无','处理理由':'只对有明确动作/对象或需显式标记推断的语段补提；SQL/Python/工具列表本身不等于查询、连接、建表任务','原始证据状态':SUMMARY,'研究状态':REVIEW} for job,r in jobs.items()]
    return out,typical,audit,req_review

def build_capabilities(rules,extracted):
    core=[];pending=[];links=[];details=[];domains={};categories={}
    for c in rules['core_capabilities']:
        members=[e for e in extracted if e['归纳典型任务ID'] in c['tasks']]
        if not members:raise ValueError('能力没有任务依据：'+c['id'])
        if any(not re.search(c['support_pattern'],e['摘要证据文本'],re.I) for e in members):raise ValueError('核心能力引用缺明确技术对象或操作：'+c['id'])
        # 叶能力在前，后向归并领域和类别；不加载课程项目内容或旧能力来触发能力。
        domain=c['domain'];category=c['category']
        if domain not in domains:domains[domain]='D'+str(len(domains)+1)
        key=(domain,category)
        if key not in categories:categories[key]=domains[domain]+'-'+str(sum(k[0]==domain for k in categories)+1).zfill(2)
        row={'能力编号':c['id'],'一级能力编号':domains[domain],'一级能力名称':domain,'二级能力编号':categories[key],'二级能力名称':category,'三级能力编号':c['id'],'三级能力名称':c['name'],'能力定义':c['definition'],'可观察的行为表现':c['behavior'],'初步评价指标（研究拟定，非招聘要求）':c['indicator'],'关联典型工作任务编号':list_ids(c['tasks']),'提取任务ID':list_ids(e['提取任务ID'] for e in members),'来源岗位ID':list_ids(e['岗位ID'] for e in members),'证据类型及依据':'岗位摘要中技术对象与操作直接支持；未核验；具体行为与评价表述为研究拟定','摘要证据摘录':list_ids(e['岗位ID']+'：'+e['摘要证据文本'] for e in members),'来源定位':list_ids(e['源文件']+'/'+e['源工作表']+'/'+e['来源摘要字段']+'@'+str(e['源Excel行号']) for e in members),'招聘原文明示证据数':0,'研究者需审查的问题':c['review'],'纳入范围':'摘要直接支持的数据库核心候选；非最终体系','研究状态':REVIEW}
        core.append(row)
        for t in c['tasks']:
            subset=[e for e in members if e['归纳典型任务ID']==t]
            links.append({'关联ID':'TC-'+str(len(links)+1).zfill(3),'典型工作任务ID':t,'三级能力/待审线索ID':c['id'],'关联类型':'摘要直接支持的候选能力关联','纳入核心候选':'是','提取任务ID':list_ids(e['提取任务ID'] for e in subset),'来源岗位ID':list_ids(e['岗位ID'] for e in subset),'判断依据':'摘要明确技术对象/操作：'+list_ids(e['摘要证据文本'] for e in subset)+'；'+c['review'],'研究状态':REVIEW})
        for e in members:details.append({'能力/线索ID':c['id'],'纳入核心候选':'是','关联判断':'技术对象与操作在摘要直接支持；研究归纳待审',**e})
    for c in rules['review_capabilities']:
        members=[e for e in extracted if e['归纳典型任务ID'] in c['tasks'] and (not c['jobs'] or e['岗位ID'] in c['jobs'])]
        if not members:raise ValueError('待审线索无任务来源')
        row={'能力编号':c['id'],'一级能力编号':'未归纳（未纳核心）','一级能力名称':'未归纳（未纳核心）','二级能力编号':'未归纳（未纳核心）','二级能力名称':'未归纳（未纳核心）','三级能力编号':'未授予正式核心三级编码','三级能力名称':c['name'],'能力定义':'待审查的技术细化线索；现有任务不足以确认此技术活动','可观察的行为表现':'暂不作为已有需求指标；需先确认任务技术对象与操作','初步评价指标（研究拟定，非招聘要求）':'暂不制定；未满足核心纳入证据门槛','关联典型工作任务编号':list_ids(c['tasks']),'提取任务ID':list_ids(e['提取任务ID'] for e in members),'来源岗位ID':list_ids(e['岗位ID'] for e in members),'证据类型及依据':'研究推断：'+c['reason'],'招聘原文明示证据数':0,'研究者需审查的问题':c['review'],'纳入范围':'证据不足的待审线索；不纳核心树、不计核心能力数','研究状态':REVIEW}
        pending.append(row)
        for t in c['tasks']:
            subset=[e for e in members if e['归纳典型任务ID']==t]
            links.append({'关联ID':'TC-'+str(len(links)+1).zfill(3),'典型工作任务ID':t,'三级能力/待审线索ID':c['id'],'关联类型':'研究推断；不等于招聘明示要求','纳入核心候选':'否','提取任务ID':list_ids(e['提取任务ID'] for e in subset),'来源岗位ID':list_ids(e['岗位ID'] for e in subset),'判断依据':c['reason'],'研究状态':REVIEW})
        for e in members:details.append({'能力/线索ID':c['id'],'纳入核心候选':'否','关联判断':'研究推断：'+c['reason'],**e})
    hierarchy=[]
    for name,id in domains.items():hierarchy.append({'节点ID':id,'层级':'一级','父编码':None,'能力名称':name,'归纳依据':'从下属已形成的三级候选反向归纳，不声称覆盖全部数据库领域','研究状态':REVIEW})
    for (d,n),id in categories.items():hierarchy.append({'节点ID':id,'层级':'二级','父编码':domains[d],'能力名称':n,'归纳依据':'归并现有叶能力的相对独立技术活动，不从教材章节预设','研究状态':REVIEW})
    for c in core:hierarchy.append({'节点ID':c['能力编号'],'层级':'三级','父编码':c['二级能力编号'],'能力名称':c['三级能力名称'],'归纳依据':c['关联典型工作任务编号'],'研究状态':REVIEW})
    return core,pending,links,details,hierarchy

def compare(selected,core,pending):
    mappings={'N113':('保留并收窄','DB001','只保留数据库/SQL查询摘要，不迁移一般业务分析的推断关系'),
              'N125':('保留并收窄','DB002','一般清洗收窄至明确SQL清洗，不将Python清洗纳入'),
              'N321':('保留并收窄','DB003','只纳明确数据库备份，不纳系统备份'),
              'N322':('保留并收窄','DB004','只纳明确数据库恢复'),
              'N341':('保留并收窄','DB005','保留性能优化目标，不新增索引、执行计划等技术要求'),
              'N112':('转待审线索','R005','接入数据不等于数据库导入'),
              'N123':('转待审线索','R006','数据库更新未证明记录级DML，一般业务数据维护不直接迁移'),
              'N212':('转待审线索','R002','统计监测不等于SQL分组聚合'),
              'N313':('转待审线索','R004','数据库运维未给出可观测的具体检查操作'),
              'N331':('转待审线索','R003','系统权限对象不明，不能称为数据库授权'),
              'N121':('任务保留待细化','','数据字典仅名词列举，创建维护动作及数据库对象未明'),
              'N412':('任务保留待细化','','数据库应用开发虽明确，缺少具体数据库技术活动，不拿通用开发能力充当叶能力')}
    rows=[]
    for old in selected['old34']['rows']:
        action,target,reason=mappings.get(old['新编码'],('排除出核心（保留原始与外围任务）','','通用、旅游业务或其他课程能力，或数据库技术依据不足；不因课程需要强迁移'))
        rows.append({'旧版本':'34项优化建议版','旧能力ID':old['新编码'],'旧能力名称':old['能力名称'],'旧父编码':old['父编码'],'操作':action,'新核心能力ID':target if target.startswith('DB') else None,'待审线索ID':target if target.startswith('R') else None,'判断依据':reason,'旧对应任务ID（照录）':old['对应任务ID'],'迁移旧任务能力关系':'否；旧疑似/推断关系不自动继承，新稿从摘要片段重新追溯','研究状态':REVIEW})
    aux_map={'C113':('合并并收窄','DB001','与C131重叠的查询部分归为数据库查询；仅两条旧编码到同一新叶，不等于迁移旧关系'),
             'C131':('合并并收窄','DB001','泛SQL处理部分不等于查询，只有明示查询/提取进入新叶'),
             'C112':('保留并收窄','DB002','仅SQL清洗'), 'C331':('拆分并收窄','DB003;DB004','备份与恢复分开，要求均为明确数据库对象'),
             'C421':('保留并收窄','DB005','保持性能优化目标，技术手段不扩展'),
             'C321':('任务保留待细化','','数据库应用管理未说明具体技术活动'),
             'C412':('任务保留待细化','','数据库应用开发未细化'),
             'C111':('转待审线索','R005','仅接入部分提供数据库导入推断线索'),
             'C332':('转待审线索','R003','仅权限部分提供推断线索，系统安全不直接迁移')}
    aux=[]
    for old in selected['old26']['rows']:
        action,target,reason=aux_map.get(old['三级能力编码'],('排除出核心（保留原始与外围任务）','','未找到足够细化且明确的数据库技术任务；不以旧触发词继承'))
        aux.append({'旧版本':'26项初版','旧能力ID':old['三级能力编码'],'旧能力名称':old['三级能力'],'操作':action,'新核心能力ID':target if target.startswith('DB') else None,'待审线索ID':target if target.startswith('R') else None,'判断依据':reason,'研究状态':REVIEW})
    course=[{'旧建议ID':r['新编码'],'名称':r['能力名称'],'本稿处理':'保留旧文件，不作为岗位需求依据；不自动迁移到核心','说明':'原表已标无招聘证据；不能以课程建议补齐岗位能力','研究状态':REVIEW} for r in selected['old_course']['rows']]
    reverse=[{'新核心ID':c['能力编号'],'新能力':c['三级能力名称'],'旧34项对应ID':list_ids(r['旧能力ID'] for r in rows if r['新核心能力ID']==c['能力编号']),'旧26项对应ID':list_ids(r['旧能力ID'] for r in aux if c['能力编号'] in str(r['新核心能力ID'] or '').split(';')),'新增判断':'未新增无任务依据核心叶；基于摘要重新提炼，与旧条目相似不表示旧关系获确认','研究状态':REVIEW} for c in core]
    return rows,aux,course,reverse

def workbook(path,sheets):
    w=Workbook();w.remove(w.active);w.properties.creator='候选研究修订脚本';w.properties.created=datetime(2026,10,23);w.properties.modified=datetime(2026,10,23)
    for name,records in sheets.items():
        s=w.create_sheet(name);headers=list(dict.fromkeys(k for r in records for k in r)) or ['说明'];s.append(headers)
        for r in records:
            s.append([r.get(k) for k in headers])
            # 保持原文字符串，不让以等号开头的源文本成为Excel公式。
            for cell in s[s.max_row]:
                if cell.data_type=='f':cell.data_type='s'
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.fill=PatternFill('solid',fgColor='285A50');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='top')
        for col,key in enumerate(headers,1):s.column_dimensions[get_column_letter(col)].width=20 if 'ID' in key or '编号' in key else 42
        for row in s.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    w.save(path)

def run(root=ROOT,output=None,verify=True):
    root=Path(root).resolve();output=Path(output or HERE/'output').resolve()
    if output.is_relative_to(root) and not output.is_relative_to(HERE):raise ValueError('禁止覆盖研究修订目录以外的仓库文件')
    rules=json.loads((HERE/'rules.json').read_text());selected,inventory,missing,files=discover(root)
    snapshot=json.loads((HERE/'input_snapshot.json').read_text())
    if verify:
        for f in snapshot['repository_files']:
            p=root/f['path']
            if p.suffix!='.xlsx':continue
            if not p.is_file() or git_sha(p)!=f['sha']:raise ValueError('输入仓库基线不一致：'+f['path'])
    jobs,descriptions,duplicates=check_jobs(selected)
    extracted,typical,old_audit,req_audit=extract_tasks(selected,rules)
    core,pending,links,details,hier=build_capabilities(rules,extracted)
    comp34,comp26,course,reverse=compare(selected,core,pending)
    logical=[]
    for r in rules['logical_relations']:
        ex=[e for e in extracted if e['归纳典型任务ID'] in r['tasks'] and e['岗位ID'] in r['jobs']]
        logical.append({'关系ID':r['id'],'起点能力ID':r['source'],'终点能力ID':r['target'],'关系类型':r['type'],'方向':r['direction'],'关系定义':r['definition'],'判断依据':r['basis'],'适用条件与边界':r['condition'],'典型任务ID':list_ids(r['tasks']),'提取任务ID':list_ids(e['提取任务ID'] for e in ex),'来源岗位ID':list_ids(r['jobs']),'证据类型':'关系为研究推断；岗位摘要只提供相关活动场景','证据摘录':list_ids(e['摘要证据文本'] for e in ex),'人工审核状态':REVIEW})
    hierarchy_edges=[{'关系ID':'HC-'+str(i).zfill(2),'起点能力ID':n['父编码'],'终点能力ID':n['节点ID'],'关系类型':'层级包含（研究归纳）','方向':'上位→下位','判断依据':n['归纳依据'],'人工审核状态':REVIEW} for i,n in enumerate([n for n in hier if n['父编码']],1)]
    questions=[{'问题类别':'证据边界','涉及ID':'全部50候选岗位','问题':'原始招聘描述未提供、证据摘录均为空；只能使用现有摘要，原质量A/B/C不等于原文已核验','处理要求':'保留状态；本稿不新增逐条招聘原文核验任务'},
               {'问题类别':'复合任务归并','涉及ID':'全部典型任务','问题':'语义分组由AI辅助逐条拟定；技术对象/场景可能仍过宽，须研究者判断归并粒度','处理要求':'查看原片段及岗位对应，不将自动归纳视为专家审核'},
               {'问题类别':'示例不自动纳入','涉及ID':'无能力编码','问题':'多表连接、完整性约束、建表等仅为用户粒度示例；现有任务未明确支撑，不新增核心条目','处理要求':'不得根据课程或教材补齐'},
               {'问题类别':'历史推断关系','涉及ID':'旧109条疑似任务能力关系','问题':'仅保留历史状态，不自动迁移、确认或逐条核验','处理要求':'新关联从摘要片段独立建立；不修改旧文件'}]
    questions.append({'问题类别':'样本独立性','涉及ID':'JOB004;JOB018;JOB023;JOB026;JOB005;JOB022','问题':'两组单位与岗位同时重复；同单位重复不自动认定为同一岗位','处理要求':'保留全部记录，研究者审查独立样本口径；不新增原文核验任务'})
    for t in typical:
        if t['数据库专业相关性']=='直接相关' and t['典型工作任务ID'] not in {x for c in rules['core_capabilities'] for x in c['tasks']}:
            questions.append({'问题类别':'数据库任务尚未形成充分叶能力','涉及ID':t['典型工作任务ID'],'问题':t['典型工作任务']+'；数据库技术相关并不等于具体操作已清楚','处理要求':'结合待审线索和原片段判断；未充分前不补造具体操作'})
    questions += [{'问题类别':'核心候选能力','涉及ID':c['能力编号'],'问题':c['研究者需审查的问题'],'处理要求':'候选、未核验；行为与评价指标为研究拟定，审查后决定是否采用'} for c in core]
    questions += [{'问题类别':'证据不足待审线索','涉及ID':c['能力编号'],'问题':c['证据类型及依据']+'；'+c['研究者需审查的问题'],'处理要求':'不纳核心，不赋正式三级编码，不计直接岗位要求'} for c in pending]
    questions += [{'问题类别':'任务动作未明或补足','涉及ID':e['提取任务ID'],'问题':e['摘要证据文本']+'；名词/工具语段不足以确认完整任务行为','处理要求':'保留为提取候选并显式标推断，不直接升级具体数据库能力'} for e in extracted if '补足' in e['任务提取判断']]
    questions += [{'问题类别':'逻辑关系','涉及ID':r['关系ID'],'问题':r['适用条件与边界'],'处理要求':'待研究者判断；非招聘明示、不设权重，不强制作为教学先修'} for r in logical]
    scopes=[{'项目':'第一阶段','说明':'只检查招聘采集、分类、结构与候选样本质量；旧文件虽曾含任务能力，本稿不改历史阶段归档'},
            {'项目':'第二阶段','说明':'区分原始描述缺失、职责/任职摘要、提取片段、典型任务；原岗位与旧任务ID保留'},
            {'项目':'第三阶段','说明':'从任务到三级技术候选，再归并二级一级；不读课程内容触发能力，不照搬5/14/34'},
            {'项目':'证据','说明':SUMMARY+'；全部结果为AI辅助候选研究结果，不是专家结论'},
            {'项目':'重复','说明':'仅列疑似重复/同单位组，不直接删除或计算独立行业频次'}]
    screening=[]
    for t in typical:
        cid=list_ids(c['能力编号'] for c in core if t['典型工作任务ID'] in c['关联典型工作任务编号'].split(';'))
        rid=list_ids(c['能力编号'] for c in pending if t['典型工作任务ID'] in c['关联典型工作任务编号'].split(';'))
        decision='有直接摘要支持的核心候选活动' if cid else '仅提供待审推断线索，不入核心' if rid else '数据库相关但具体活动未细化，暂不提叶能力' if t['数据库专业相关性']=='直接相关' else '技术实现不明确或非数据库活动；保留任务，不强建能力关联'
        screening.append({**t,'核心候选能力ID':cid or '无','待审线索ID':rid or '无','能力提炼处置':decision})
    job_fields=[r for r in missing if r['文件']==selected['jobs']['path'] and r['工作表']==selected['jobs']['sheet']]
    job_summary=[{'检查项目':'候选岗位数','数量':len(jobs),'口径':'未删除疑似重复，全部候选待核验'}, {'检查项目':'招聘原文证据已验证数','数量':0,'口径':'原始证据摘录缺失，不把摘要当原文'}]
    job_summary += [{'检查项目':r['字段']+'缺失','数量':r['缺失数'],'口径':'分母'+str(r['记录数'])} for r in job_fields]
    for kind in ['招聘链接完全重复','招聘链接规范化重复','单位重复（不等于岗位重复）','岗位名称重复（不等于样本重复）','单位与岗位同时重复']:
        job_summary.append({'检查项目':kind,'数量':len({d['重复组'] for d in duplicates if d['检查类型']==kind}),'口径':'组数，不是独立岗位数；保留全部记录'})
    outputs={
        '01_招聘信息数据检查.xlsx':{'阶段与证据边界':scopes,'检查结果汇总':job_summary,'实际文件工作表盘点':inventory,'字段缺失统计':missing,'岗位原始与证据状态':jobs,'疑似重复与同单位':duplicates},
        '02_典型工作任务修订与归纳表.xlsx':{'原始输入描述索引':descriptions,'任务片段提取':extracted,'典型工作任务候选':typical,'旧79任务逐条对照':old_audit,'任职要求审阅记录':req_audit},
        '03_数据库相关典型工作任务筛选表.xlsx':{'全部典型任务相关性':screening,'直接相关任务':[t for t in screening if t['数据库专业相关性']=='直接相关'],'间接相关任务':[t for t in screening if t['数据库专业相关性']=='间接相关'],'不相关与外围任务':[t for t in screening if t['数据库专业相关性']=='不相关']},
        '04_新版候选三级能力体系表.xlsx':{'核心三级能力候选':core,'自下而上层级':hier,'待审线索不纳核心':pending},
        '05_三级能力定义行为与初步评价指标表.xlsx':{'核心候选定义指标':core,'证据不足线索说明':pending},
        '06_典型工作任务三级能力关联表.xlsx':{'典型任务能力关联':links,'片段岗位证据链':details},
        '07_三级能力逻辑关系表.xlsx':{'候选非层级关系':logical,'候选层级包含关系':hierarchy_edges},
        '08_新旧能力体系对比表.xlsx':{'旧34项逐条对照':comp34,'旧26项辅助对照':comp26,'新核心反向对照':reverse,'旧课程建议单列':course,'旧上位能力保留说明':[{'旧ID':r['新编码'],'旧层级':r['能力层级'],'旧名称':r['能力名称'],'处理':'保留旧编码，禁止上位直接套用；从新叶重新聚合','研究状态':REVIEW} for r in selected['old_all']['rows'] if r['能力层级']!='三级']},
        '09_研究者人工审查问题清单.xlsx':{'审查问题':questions}}
    output.mkdir(parents=True,exist_ok=True)
    for name,sheets in outputs.items():workbook(output/name,sheets)
    statistics={'source_commit':snapshot['main_commit'],'source_jobs':len(jobs),'historical_tasks':len(old_audit),'source_descriptions':len(descriptions),'extracted_task_fragments':len(extracted),'responsibility_fragments':sum(e['来源摘要字段'].startswith('岗位职责') for e in extracted),'requirement_fragments':sum(e['来源摘要字段'].startswith('任职要求') for e in extracted),'typical_tasks':len(typical),'db_relevance_counts':dict(Counter(t['数据库专业相关性'] for t in typical)),'levels':dict(Counter(n['层级'] for n in hier)),'core_direct_summary_candidates':len(core),'core_research_inference_only':0,'pending_inference_clues':len(pending),'verified_original_recruitment_evidence':0,'job_classification':dict(Counter(j['相关性类别'] for j in jobs)),'job_quality_original':dict(Counter(j['证据质量（初评）'] for j in jobs)),'original34_operations':dict(Counter(r['操作'] for r in comp34)),'original26_operations':dict(Counter(r['操作'] for r in comp26)),'core_new_without_old_equivalent':sum(not r['旧34项对应ID'] for r in reverse),'logical_relations':len(logical),'hierarchy_edges':len(hierarchy_edges),'old_suspect_relations':sum(r['疑似误匹配']=='是' for r in selected['review']['rows']),'questions':len(questions),'duplicate_groups':dict(Counter(k for k in {r['重复组']:r['检查类型'] for r in duplicates}.values())),'input_roles':{k:{a:b for a,b in v.items() if a!='rows'} for k,v in selected.items()},'outputs':list(outputs)}
    (output/'processing_summary.json').write_text(json.dumps(statistics,ensure_ascii=False,indent=2)+'\n')
    return statistics

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path);args=p.parse_args()
    print(json.dumps(run(output=args.output),ensure_ascii=False,indent=2))
