"""第三阶段最终规范化候选修订；仅消费既有Excel，不重提任务或实施课程映射。"""
from pathlib import Path
from collections import Counter
import argparse, copy, hashlib, importlib.util, json
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
spec=importlib.util.spec_from_file_location('revision_base',HERE.parent/'build.py')
base=importlib.util.module_from_spec(spec);spec.loader.exec_module(base)
PRINCIPLE='如果拆分后的内容只是同一技术活动的不同方法、步骤或操作细节，优先作为三级能力的行为表现或知识技能点；只有当拆分后的内容具有明显不同的工作目标、技术成果和评价标准，并且有相应岗位任务依据时，才考虑拆成独立三级能力。'
LABELS={'A':'摘要明确支持；非已核验招聘原文','B':'数据库相关、关系型技术解释、待审查','C':'业务任务的合理技术推断、待人工审查'}
STATUS='候选研究结果；待研究者审查；未验证、未定稿'
BOUNDARIES={
'DB001':'基础、范围较广的SQL查询与取数，评价取数对象、条件及结果；不预设JOIN。与DB017并存，不重复解释为独立岗位需求。',
'DB017':'跨关系表关联场景下的查询，评价关联依据和关联结果。当前仅B类：复杂取数并不必然要求JOIN；内、外、自连接仅作方法或行为点。',
'DB005':'优化目标与结果层面，比较查询结果一致性和性能记录；不限定索引方法。',
'DB011':'仅在索引适用时选择调整索引；是DB005的一种技术手段，当前仅B类，不能继承DB005的A类。',
'DB006':'解释分支一：更新记录值。TT019没有确定更新对象，不代表岗位已明确要求UPDATE。',
'DB016':'解释分支二：调整表结构。与DB006互为备选解释，数据库更新也可能指版本升级，不代表两项明确需求。',
'DB003':'备份对象选择、执行、保存产物与检查属于同一技术活动的行为；不继续拆分，不以产物存在直接证明可恢复。',
'DB004':'恢复是独立工作目标与结果。仅在基于备份恢复场景依赖可用备份；备份执行不是必然教学先修。'}

def sha(path):
 b=path.read_bytes();return hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()

def protected():
 for f in json.loads((HERE/'input_snapshot.json').read_text())['repository_files']:
  if sha(ROOT/f['path'])!=f['sha']:raise ValueError('旧成果发生变化：'+f['path'])

def find(sheet,fields,directory):
 matches=[]
 for p in sorted(directory.glob('*.xlsx')):
  w=load_workbook(p,read_only=True)
  try:
   if sheet in w.sheetnames and set(fields)<=set(next(iter(w[sheet].values))):matches.append(p)
  finally:w.close()
 if len(matches)!=1:raise ValueError(f'输入工作表无法唯一定位：{sheet}: {matches}')
 p=matches[0];return p,base.read(p,sheet)

def inputs():
 d=HERE.parent/'round2_abc/output';specs=[('pool','合并三级候选池',['能力编号']),('hier','自下而上全层级',['节点ID']),('baseline_edges','候选层级关系',['关系ID']),('round2_logic','非层级关系仍为待审',['关系ID']),('links','逐条任务能力证据',['关联ID']),('tasks','全部片段及类别',['提取任务ID']),('typical','典型任务类别集合',['典型工作任务ID']),('knowledge','课程知识储备',['储备知识ID']),('reserve','教学情境能力储备',['储备ID']),('questions','研究者审查问题',['问题ID']),('previous','初稿到ABC对照',['初稿编号'])]
 data={};inventory=[]
 for key,sheet,fields in specs:
  p,rows=find(sheet,fields,d);data[key]=rows
  inventory.append({'文件':str(p.relative_to(ROOT)),'工作表':sheet,'实际记录数':len(rows),'字段':';'.join(rows[0]) if rows else '', '缺失必需字段':'无（定位时已验证）', '版本来源':json.loads((HERE/'input_snapshot.json').read_text())['pr7_head']})
 p,rows=find('候选非层级关系',['关系ID'],HERE.parent/'output');data['old_logic']=rows
 inventory.append({'文件':str(p.relative_to(ROOT)),'工作表':'候选非层级关系','实际记录数':len(rows),'字段':';'.join(rows[0]),'版本来源':'PR7第一稿（5项核心能力范围）'})
 p,rows=find('任务片段提取',['提取任务ID'],HERE.parent/'output');data['original_tasks']=rows
 inventory.append({'文件':str(p.relative_to(ROOT)),'工作表':'任务片段提取','实际记录数':len(rows),'字段':';'.join(rows[0]),'版本来源':'PR7第一稿'})
 return data,inventory

def logical(pool):
 names={c['能力编号']:c['三级能力名称'] for c in pool}
 # 所有条目均是研究者待审建议；保留表示保留候选，不表示已确认。
 specs=[
 ('DB001','DB017','技术支撑','基础查询表达与结果核对为跨表查询提供技术基础','两项定义中的查询对象与结果核对要求；不是摘要直接提出的能力关系','任务确需跨表关联，且DB017候选获研究者认可','保留候选','教学先修暂不确立，缺少学习诊断依据'),
 ('DB001','DB005','技术支撑','可执行查询及结果核对为性能对照提供对象和正确性基准','查询与优化的技术逻辑；原三条关系中查询支撑优化条目复核','有明确查询、可比环境与性能记录','保留候选','不据此断言必须先完成某教学项目'),
 ('DB011','DB005','技术支撑','索引选择调整是实现查询性能优化的一种条件性手段','DB005目标与DB011方法定义的技术分析；无招聘摘要明示此关系','索引确与瓶颈相关，且对照支持该方法','保留候选','方向为方法支撑目标；不把DB005的A类传递给DB011'),
 ('DB003','DB004','条件性技术依赖','基于备份的恢复依赖可用备份产物，而非依赖学习者曾执行本次备份','备份产物与恢复输入之间的技术逻辑；原条件性条目细化','明确采用备份恢复；其他恢复路线不适用','保留候选','源指供给备份产物的能力，目标指使用产物恢复的能力，不是教学先修'),
 ('DB003','DB004','组合关系','备份与恢复可共同构成备份恢复验证闭环','JOB032摘要数据库备份恢复共同出现，仅支持组合线索，关系模型是研究推断','同一备份恢复验证工作场景','保留候选','组合不表达因果或学习先后；规范记录一次，不生成反向边'),
 ('DB007','DB016','技术支撑','已有表结构设计与对象关系说明可能支持结构变更影响分析','两项定义的设计成果与变更对象之间的技术逻辑','确实存在关系表结构变更，且可获得原结构说明','待审','TT019更新对象尚不明确，暂不作为默认启用关系'),
 ('DB005','DB011','教学先修','优化目标认识可能用于组织索引学习，但不等于完整性能优化能力先修','仅目标—手段联系不足以确定教学先修','需要后续教学设计及学习诊断依据','不保留','不能从目标—方法联系推导教学先修，亦不与反向支撑边混为循环'),
 ('DB003','DB004','教学先修','备份工作流程先后不等于恢复学习的必然先修','可使用他人提供的备份学习恢复','若未来提出教学先修须有独立学习依据','不保留','技术依赖另列，禁止以流程顺序替代教学依据')]
 return [{'关系编号':f'FR-L{i:03}','源能力编号':s,'源能力名称':names[s],'目标能力编号':t,'目标能力名称':names[t],'关系类型':kind,'关系含义':meaning,'判断依据':basis,'适用条件':condition,'审查建议':decision,'处理理由':reason,'是否属于研究推断':'是','方向说明':'组合为无先后配对，源目标仅作存储顺序' if kind=='组合关系' else '源能力→目标能力；不表示岗位明示或必然教学顺序','人工审查状态':'待研究者审查，未经专家确认','默认启用':'否'} for i,(s,t,kind,meaning,basis,condition,decision,reason) in enumerate(specs,1)]

def run(output=None):
 output=Path(output or HERE/'output').resolve()
 if output.is_relative_to(ROOT) and not output.is_relative_to(HERE):raise ValueError('输出仅允许本轮独立目录或仓库外临时目录')
 protected();d,inventory=inputs();pool=copy.deepcopy(d['pool']);hier=copy.deepcopy(d['hier']);links=copy.deepcopy(d['links']);comparison=[]
 for c in pool:
  old=c['三级能力名称'];id=c['能力编号']
  if id=='DB001':c['三级能力名称']='SQL数据查询与提取'
  if c['二级能力名称']=='业务数据聚合统计':c['二级能力名称']='SQL数据聚合统计'
  c['能力边界']=BOUNDARIES.get(id,'限于定义所列关系型数据库技术对象与活动；行为细节是研究拟定，不是岗位逐项明示要求。')
  c['证据类别及说明']='；'.join(LABELS[t] for t in c['具有的关联证据类别'].split(';'))
  c['证据来源状态']='候选岗位摘要；全部岗位待核验；无已核验招聘原文'
  c['能力候选状态']=STATUS
  c['拆分原则应用']='保留原三级节点；方法、步骤及操作细节列为行为，不新增节点'
  if id=='DB001':c['能力定义（研究拟定）']='在关系型数据库研究范围内，根据取数目标使用SQL执行查询并提取数据；SQL是现有摘要支持与关系型技术解释的范围，不把全部B/C任务升级为SQL明示要求。'
  if id=='DB017':c['可观察行为（研究拟定）']+='；内连接、外连接、自连接仅可在适用场景作方法选择，不新增能力，也不作为招聘明示清单。'
  c['人工审查问题或备注']=c['研究者需审查的问题']+'；'+c['能力边界']
  comparison.append({'三级能力编号':id,'原名称':old,'修订名称':c['三级能力名称'],'是否修改':'是' if old!=c['三级能力名称'] else '否','修订原因':'突出关系型数据库SQL活动并保留逐条证据边界' if id=='DB001' else '研究者认可现有技术粒度，保留名称并规范证据与边界','能力边界':c['能力边界'],'拆分原则应用情况':c['拆分原则应用'],'原二级名称':next(x for x in d['pool'] if x['能力编号']==id)['二级能力名称'],'修订二级名称':c['二级能力名称']})
 byid={c['能力编号']:c for c in pool}
 for n in hier:
  if n['节点ID'] in byid:n['名称']=byid[n['节点ID']]['三级能力名称']
  elif n['名称']=='业务数据聚合统计':n['名称']='SQL数据聚合统计'
 edges=[{'关系编号':f'FR-H{i:03}','源能力编号':n['父编码'],'目标能力编号':n['节点ID'],'关系类型':'候选层级包含','判断依据':'复用第二轮稳定父编码，自下而上归纳；非新增招聘明示要求','人工审查状态':STATUS} for i,n in enumerate([n for n in hier if n['父编码']],1)]
 for r in links:
  r['原能力名称']=r['能力名称'];r['能力名称']=byid[r['能力编号']]['三级能力名称']
  r['原关联证据类别']=r['关联证据类别'];r['关联类别规范含义']=LABELS[r['关联证据类别']]
  r['证据复核结论']='原任务类别与关联类别独立保留；不升级、不新增、不删除'
  r['类别修改理由']='未修改';r['来源状态']='岗位摘要，未核验原始招聘公告'
  r['人工审查意见']='待研究者审查；'+byid[r['能力编号']]['能力边界']
 logic=logical(pool)
 questions=[]
 for id in ['DB001','DB017','DB005','DB011','DB006','DB016','DB007','DB008','DB009','DB010','DB012','DB013','DB014','DB015']:
  c=byid[id];questions.append({'问题编号':f'FR-Q{len(questions)+1:03}','涉及对象':id,'需研究者判断的问题':c['人工审查问题或备注'],'可用依据':c['关联证据ID'],'当前处理':'保留候选，不改变证据类别，不增加招聘核验任务','审查状态':'待研究者审查'})
 questions.append({'问题编号':'FR-Q015','涉及对象':'全部非层级关系','需研究者判断的问题':'审查8条关系判断：5条保留候选、1条待审、2条不保留；保留候选也未经确认，教学先修本轮不确立','可用依据':'05工作簿逐条解释；没有学习诊断或专家审核结论','当前处理':'全部默认不启用，不设置权重','审查状态':'待研究者审查'})
 questions.append({'问题编号':'FR-Q016','涉及对象':'任务片段口径','需研究者判断的问题':'第二轮181片段=原179+同一JOB012摘要补提两维度；是否认可补提的活动维度划分','可用依据':'EX2-JOB012-R01/R02；本轮原样复用，不改第二阶段','当前处理':'179与181分别报告，岗位及典型任务未增加','审查状态':'待研究者审查'})
 output.mkdir(parents=True,exist_ok=True)
 for source,target in [('方法说明.md','06_第三阶段研究方法与划分规则.md'),('修订总结.md','08_第三阶段修订总结与前后对照.md')]:
  (output/target).write_text((HERE/source).read_text())
 base.workbook(output/'01_第三阶段规范化三级候选能力体系.xlsx',{'三级候选能力':pool,'完整能力层级':hier,'实际输入清单':inventory})
 base.workbook(output/'02_三级能力名称与边界修订对照.xlsx',{'名称与边界对照':comparison})
 base.workbook(output/'03_任务能力ABC证据关联复核.xlsx',{'逐条关联复核':links,'证据分类定义':[{'关联证据类别':k,'含义':v} for k,v in LABELS.items()]})
 base.workbook(output/'04_三级能力层级包含关系.xlsx',{'层级包含关系':edges,'层级节点':hier})
 base.workbook(output/'05_三级能力非层级逻辑关系.xlsx',{'逐条关系判断':logic,'保留及待审候选':[r for r in logic if r['审查建议']!='不保留'],'不保留关系':[r for r in logic if r['审查建议']=='不保留']})
 base.workbook(output/'07_第三阶段研究者人工审查清单.xlsx',{'研究者审查问题':questions})
 summary={'能力层级数量':dict(Counter(n['能力层级'] for n in hier)),'任务片段':len(d['tasks']),'原任务片段':len(d['original_tasks']),'典型任务':len(d['typical']),'关联数':len(links),'关联类别':dict(Counter(r['关联证据类别'] for r in links)),'多类去重能力数':{t:sum(t in c['具有的关联证据类别'].split(';') for c in pool) for t in LABELS},'层级关系':len(edges),'非层级判断记录':len(logic),'非层级候选':sum(r['审查建议']!='不保留' for r in logic),'非层级建议':dict(Counter(r['审查建议'] for r in logic)),'已确认非层级关系':0,'研究者问题':len(questions)}
 (output/'processing_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
 return d,pool,hier,links,edges,logic,summary

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--output',type=Path);a=p.parse_args();print(json.dumps(run(a.output)[-1],ensure_ascii=False,indent=2))
