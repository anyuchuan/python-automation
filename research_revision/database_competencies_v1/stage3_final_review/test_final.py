"""真实输入、关系语义及历史文件保护回归检查；不修改研究输入。"""
import unittest,tempfile
from pathlib import Path
from collections import Counter
import build as b
from openpyxl import load_workbook

class FinalReview(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory();cls.out=Path(cls.tmp.name)/'one'
  cls.data,cls.pool,cls.hier,cls.links,cls.edges,cls.logic,cls.summary=b.run(cls.out)
 @classmethod
 def tearDownClass(cls):cls.tmp.cleanup()
 def test_01_counts_and_stable_ids(self):
  self.assertEqual(Counter(n['能力层级'] for n in self.hier),{'一级':3,'二级':9,'三级':17})
  self.assertEqual({c['能力编号'] for c in self.pool},{f'DB{i:03}' for i in range(1,18)})
  self.assertEqual(len(self.pool),17)
 def test_02_required_fields(self):
  for c in self.pool:
   for f in ['三级能力名称','能力定义（研究拟定）','可观察行为（研究拟定）','初步评价指标（研究拟定）','关联典型工作任务编号','证据类别及说明','证据来源状态','能力候选状态','人工审查问题或备注']:self.assertTrue(c[f],(c['能力编号'],f))
 def test_03_stable_parents(self):
  self.assertEqual({n['节点ID']:n['父编码'] for n in self.hier},{n['节点ID']:n['父编码'] for n in self.data['hier']})
  nodes={n['节点ID']:n for n in self.hier}
  for c in self.pool:
   self.assertEqual(nodes[c['能力编号']]['父编码'],c['二级能力编号'])
   self.assertEqual(nodes[c['二级能力编号']]['父编码'],c['一级能力编号'])
 def test_04_endpoints_and_quotes(self):
  tasks={t['提取任务ID']:t for t in self.data['tasks']};typ={t['典型工作任务ID'] for t in self.data['typical']};caps={c['能力编号'] for c in self.pool}
  for r in self.links:
   self.assertIn(r['能力编号'],caps);self.assertIn(r['典型工作任务ID'],typ)
   t=tasks[r['提取任务ID']];self.assertEqual(r['岗位ID'],t['岗位ID']);self.assertEqual(r['摘要证据文本'],t['摘要证据文本']);self.assertIn(r['摘要证据文本'],r['完整摘要上下文'])
 def test_05_two_evidence_dimensions(self):
  tasks={t['提取任务ID']:t for t in self.data['tasks']}
  for r in self.links:
   self.assertIn(r['关联证据类别'],b.LABELS);self.assertEqual(r['原任务类别'],tasks[r['提取任务ID']]['岗位任务关联类别'])
   self.assertEqual(r['关联类别规范含义'],b.LABELS[r['关联证据类别']])
  self.assertTrue(any(r['原任务类别']!=r['关联证据类别'] for r in self.links))
 def test_06_no_evidence_change(self):
  before={r['关联ID']:r for r in self.data['links']}
  self.assertEqual(len(self.links),65);self.assertEqual(set(before),{r['关联ID'] for r in self.links})
  for r in self.links:
   for k,v in before[r['关联ID']].items():
    if k!='能力名称':self.assertEqual(r[k],v,k)
   self.assertEqual(r['原关联证据类别'],r['关联证据类别'])
  self.assertEqual(Counter(r['关联证据类别'] for r in self.links),{'A':11,'B':27,'C':27})
 def test_07_no_inherited_A(self):
  for cap in ['DB006','DB011','DB016','DB017']:
   self.assertEqual({r['关联证据类别'] for r in self.links if r['能力编号']==cap},{'B'})
 def test_08_hierarchy_acyclic_connected(self):
  nodes={n['节点ID']:n for n in self.hier};levels={'一级':1,'二级':2,'三级':3};self.assertEqual(len(self.edges),26)
  self.assertEqual(len({(e['源能力编号'],e['目标能力编号']) for e in self.edges}),26)
  for id,n in nodes.items():
   seen=set()
   while n['父编码']:
    self.assertNotIn(n['节点ID'],seen);seen.add(n['节点ID']);p=nodes[n['父编码']]
    self.assertEqual(levels[n['能力层级']],levels[p['能力层级']]+1);n=p
   self.assertEqual(n['能力层级'],'一级')
 def test_09_logic_semantics(self):
  caps={c['能力编号'] for c in self.pool}
  for r in self.logic:
   self.assertIn(r['源能力编号'],caps);self.assertIn(r['目标能力编号'],caps)
   for f in ['关系含义','判断依据','适用条件','人工审查状态']:self.assertTrue(r[f])
   self.assertEqual(r['是否属于研究推断'],'是');self.assertEqual(r['默认启用'],'否')
  self.assertEqual(Counter(r['审查建议'] for r in self.logic),{'保留候选':5,'待审':1,'不保留':2})
  self.assertTrue(all(r['审查建议']=='不保留' for r in self.logic if r['关系类型']=='教学先修'))
  dependencies=[r for r in self.logic if r['关系类型']=='条件性技术依赖'];self.assertEqual(len(dependencies),1);self.assertIn('备份恢复',dependencies[0]['适用条件'])
 def test_10_method_exact_principle(self):
  self.assertIn(b.PRINCIPLE,(self.out/'06_第三阶段研究方法与划分规则.md').read_text())
 def test_11_all_121_previous_files_unchanged(self):b.protected()
 def test_12_no_task_extraction_or_course_mapping(self):
  self.assertEqual(len(self.data['tasks']),181);self.assertEqual(len(self.data['original_tasks']),179);self.assertEqual(len(self.data['typical']),48)
  self.assertEqual(len(self.data['knowledge']),37);self.assertEqual(len(self.data['reserve']),10)
  for f in self.out.glob('*.xlsx'):
   self.assertNotIn('课程匹配',f.name)
 def test_13_workbooks_and_reproducibility(self):
  second=Path(self.tmp.name)/'two';b.run(second)
  self.assertEqual(len(list(self.out.glob('*.xlsx'))),6)
  for p in self.out.glob('*.xlsx'):
   a=load_workbook(p,read_only=True);z=load_workbook(second/p.name,read_only=True)
   try:
    self.assertEqual(a.sheetnames,z.sheetnames)
    for name in a.sheetnames:self.assertEqual(list(a[name].values),list(z[name].values))
   finally:a.close();z.close()
 def test_14_only_necessary_name_changes(self):
  old={r['能力编号']:r for r in self.data['pool']}
  self.assertEqual([c['能力编号'] for c in self.pool if c['三级能力名称']!=old[c['能力编号']]['三级能力名称']],['DB001'])
 def test_15_protect_output_boundary(self):
  with self.assertRaises(ValueError):b.run(b.ROOT/'phase3/output')

if __name__=='__main__':unittest.main()
