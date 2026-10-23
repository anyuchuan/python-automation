"""PR #7补充修订：A/B/C证据分层候选池。原文件只读，课程储备独立构建。"""
from __future__ import annotations
from collections import Counter,defaultdict
import argparse
import copy
import importlib.util
import json
from pathlib import Path
import re
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
V1=HERE.parent
ROOT=HERE.parents[2]
spec=importlib.util.spec_from_file_location('pr7_first_draft',V1/'build.py')
base=importlib.util.module_from_spec(spec);spec.loader.exec_module(base)
STATUS='候选研究结果；待研究者审查；未验证、未定稿'
LABELS={'A':'明确数据库技术相关','B':'数据库相关、关系型技术解释、待审查','C':'业务数据任务的关系型实现研究推断、待审查','X':'不作本轮数据库岗位能力依据'}

def unique_sheet(directory,sheet,required):
    found=[]
    for path in sorted(Path(directory).rglob('*.xlsx')):
        w=load_workbook(path,read_only=True)
        try:
            if sheet in w.sheetnames:
                h=next(iter(w[sheet].values))
                if set(required)<=set(h):found.append(path)
        finally:w.close()
    if len(found)!=1:raise ValueError(f'无法唯一读取 {sheet}: {found}')
    return found[0],base.read(found[0],sheet)

def load_inputs(root=ROOT):
    selected,inventory,missing,files=base.discover(root)
    extracted_path,extracted=unique_sheet(V1/'output','任务片段提取',['提取任务ID','摘要证据文本'])
    typical_path,typical=unique_sheet(V1/'output','典型工作任务候选',['典型工作任务ID','典型工作任务'])
    pool_path,old_core=unique_sheet(V1/'output','核心三级能力候选',['能力编号','三级能力名称'])
    _,old_pending=unique_sheet(V1/'output','待审线索不纳核心',['能力编号','三级能力名称'])
    course_paths={root/r['文件'] for r in inventory if r['工作表']=='知识点' and '条目ID' in r['字段']}
    if len(course_paths)!=1:raise ValueError('课程储备知识输入不唯一')
    course_path=next(iter(course_paths));knowledge=base.read(course_path,'知识点')
    return {'selected':selected,'inventory':inventory,'missing':missing,'extracted':extracted,'typical':typical,'old_core':old_core,'old_pending':old_pending,'knowledge':knowledge,'source_paths':{'任务片段':extracted_path,'典型任务':typical_path,'旧核心':pool_path,'课程储备知识':course_path}}

def prepare_tasks(data,rules):
    rows=copy.deepcopy(data['extracted']);jobs={r['样本ID']:r for r in data['selected']['jobs']['rows']}
    # 补提原任职摘要中遗漏的复杂取数与清洗语境，不修改旧179片段。
    job='JOB012';field='任职要求摘要（不含学历经验）';span='精通SQL，能完成复杂数据提取清洗与分析'
    if span not in jobs[job][field]:raise ValueError('JOB012补提引文不在源摘要')
    seed=copy.deepcopy(next(r for r in rows if r['岗位ID']==job))
    for id,group,dimension in [('EX2-JOB012-R01','TT004','SQL语境复杂数据提取'),('EX2-JOB012-R02','TT002','复杂数据清洗，SQL作用范围待解释')]:
        row={**seed,'提取任务ID':id,'原任务ID':'无（本轮从任职摘要补提）','来源摘要字段':field,'摘要证据文本':span,'完整摘要上下文':jobs[job][field],'提取的工作任务':span,'提取活动维度（研究判断）':dimension,'归纳典型任务ID':group,'源单元格':'H'+str(seed['源Excel行号']),'任务提取判断':'同一摘要的活动维度补提；清洗是否由SQL实施仍为技术解释，不是新增岗位'}
        row['归纳后的典型工作任务']=next(t['典型工作任务'] for t in data['typical'] if t['典型工作任务ID']==group)
        row['本轮来源']='第二轮补提（同一引文两维度，不能视为两个岗位）';rows.append(row)
    for row in rows:
        id=row['提取任务ID'];group=row['归纳典型任务ID'];classification=rules['task_classification'][group]
        tier='A' if id.startswith('EX2-') else classification['category']
        row['岗位任务关联类别']=tier;row['任务类别说明']=LABELS[tier]
        row['任务分类依据']=classification['reason'] if not id.startswith('EX2-') else '任职摘要明确SQL语境与取数/清洗活动，但具体SQL语法和工具作用边界需分别解释'
        row['本轮来源']=row.get('本轮来源','沿用PR7初稿片段，未改旧文件')
        row['本轮研究状态']=STATUS;row['原证据质量等级说明']='原A/B/C为历史质量初评，与本轮任务关联A/B/C分类无关'
        product=re.findall(r'MySQL|SQL\s*Server|Oracle',jobs[row['岗位ID']]['任职要求摘要（不含学历经验）'],re.I)
        row['同岗数据库产品语境']=base.list_ids(product) or '未明确；不因此排除B类，关系型仅为课程技术解释边界'
        if row['摘要证据文本'] not in jobs[row['岗位ID']][row['来源摘要字段']]:raise ValueError('引文不可追溯：'+id)
    return rows

def make_links(tasks,rules):
    rows=[];seen=set();caps={c['id']:c for c in rules['capabilities']}
    patterns={c['id']:c.get('support_pattern') for c in rules['capabilities']}
    for rule_index,r in enumerate(rules['links'],1):
        matched=[]
        for t in tasks:
            if r['ids'] and t['提取任务ID'] not in r['ids']:continue
            if r['groups'] and t['归纳典型任务ID'] not in r['groups']:continue
            if r['jobs'] and t['岗位ID'] not in r['jobs']:continue
            source_tier=t['岗位任务关联类别'];tier=r['tier']
            # A任务可在更细技术操作上只支持B解释；C关联仅使用C任务，禁止把X强转数据库需求。
            if source_tier=='X':continue
            if tier=='A' and source_tier!='A':continue
            if tier=='B' and source_tier not in {'A','B'}:continue
            if tier=='C' and source_tier!='C':continue
            matched.append(t)
        if not matched:raise ValueError('关联规则无合格任务：'+str(rule_index))
        for t in matched:
            key=(r['cap'],t['提取任务ID'],r['tier'])
            if key in seen:continue
            seen.add(key)
            if r['tier']=='A' and (not patterns.get(r['cap']) or not re.search(patterns[r['cap']],t['摘要证据文本'],re.I)):
                raise ValueError('A关联不能由未明示细化活动支持：'+str(key))
            rows.append({'关联ID':'ABC-R'+str(len(rows)+1).zfill(4),'能力编号':r['cap'],'能力名称':caps[r['cap']]['name'],'关联证据类别':r['tier'],'关联证据性质':LABELS[r['tier']],'原任务类别':t['岗位任务关联类别'],'任务类别与关联类别差异':'同级' if t['岗位任务关联类别']==r['tier'] else '原任务明确技术目标，但本项细化方法仅为B类解释；不升级为A','典型工作任务ID':t['归纳典型任务ID'],'提取任务ID':t['提取任务ID'],'岗位ID':t['岗位ID'],'原任务ID':t['原任务ID'],'摘要证据文本':t['摘要证据文本'],'完整摘要上下文':t['完整摘要上下文'],'任务来源字段':t['来源摘要字段'],'源文件':t['源文件'],'源工作表':t['源工作表'],'源Excel行号':t['源Excel行号'],'源单元格':t.get('源单元格'),'原始招聘链接':t['原始招聘链接'],'原证据质量等级':t['原证据质量等级'],'原核验状态':t['原核验状态'],'数据库产品语境':t['同岗数据库产品语境'],'判断或推断理由':r['reason'],'采用条件与替代解释':r['condition'],'招聘原文明示且已核验':'否，只有候选岗位摘要','关系型技术范围':'关系型为课程范围；A明确任务操作不等于岗位已明确数据库类型，B/C更不等于明示技术要求','规则ID':'RULE-'+str(rule_index).zfill(3),'研究状态':STATUS})
    return rows

def pool_and_hierarchy(links,rules):
    pool=[];domains={};categories={}
    for c in rules['capabilities']:
        ls=[r for r in links if r['能力编号']==c['id']]
        if not ls:raise ValueError('能力无任务依据：'+c['id'])
        tiers=sorted({r['关联证据类别'] for r in ls});highest=tiers[0]
        domain=c['domain'];category=c['category']
        if domain not in domains:domains[domain]='ABC-D'+str(len(domains)+1)
        key=(domain,category)
        if key not in categories:categories[key]=domains[domain]+'-'+str(sum(k[0]==domain for k in categories)+1).zfill(2)
        pool.append({'能力编号':c['id'],'一级能力编号':domains[domain],'一级能力名称':domain,'二级能力编号':categories[key],'二级能力名称':category,'三级能力编号':c['id'],'三级能力名称':c['name'],'能力定义（研究拟定）':c['definition'],'可观察行为（研究拟定）':c['behavior'],'初步评价指标（研究拟定）':c['indicator'],'具有的关联证据类别':';'.join(tiers),'去重计数归类（不升级其他证据）':highest,'A直接支持任务ID':base.list_ids(r['提取任务ID'] for r in ls if r['关联证据类别']=='A') or '无','B解释任务ID':base.list_ids(r['提取任务ID'] for r in ls if r['关联证据类别']=='B') or '无','C推断任务ID':base.list_ids(r['提取任务ID'] for r in ls if r['关联证据类别']=='C') or '无','关联典型工作任务编号':base.list_ids(r['典型工作任务ID'] for r in ls),'来源岗位ID':base.list_ids(r['岗位ID'] for r in ls),'关联证据ID':base.list_ids(r['关联ID'] for r in ls),'关系型范围说明':'关系型数据库为本次课程技术范围限定，不声明所有来源岗位均明示关系型数据库','直接摘要支持边界':'A仅支持同粒度活动目标；B为数据库任务的关系型解释；C为业务数据实现推断，均非已核验原文','研究者需审查的问题':c['review'],'纳入候选结构分析':'是；B/C不因产品名或操作细节缺失而排除','候选粒度与交叉说明':('与DB017为查询目标和条件性细化方法，需审查层级/边界，不能相加为独立岗位需求' if c['id']=='DB001' else '与DB001查询目标存在包含交叉，需审查是否保留为独立三级技术细化' if c['id']=='DB017' else '与DB011是优化目标和索引方法的关系，不能相加为两项明示需求' if c['id']=='DB005' else '是DB005目标的条件性方法，非所有优化都要求索引' if c['id']=='DB011' else '与DB016是更新对象不同的解释分支，不表示同一岗位同时要求两者' if c['id']=='DB006' else '与DB006是更新对象不同的解释分支，不表示两项确定需求' if c['id']=='DB016' else '与同组候选的具体职责分工仍待审查；不是独立需求频次'),'采用或定稿状态':STATUS})
    hierarchy=[]
    for name,id in domains.items():hierarchy.append({'节点ID':id,'能力层级':'一级','父编码':None,'名称':name})
    for (domain,name),id in categories.items():hierarchy.append({'节点ID':id,'能力层级':'二级','父编码':domains[domain],'名称':name})
    for c in pool:hierarchy.append({'节点ID':c['能力编号'],'能力层级':'三级','父编码':c['二级能力编号'],'名称':c['三级能力名称']})
    children=defaultdict(list)
    for n in hierarchy:
        if n['父编码']:children[n['父编码']].append(n['节点ID'])
    def leaves(id):
        return {id} if id not in children else set().union(*(leaves(i) for i in children[id]))
    bycap={c['能力编号']:c for c in pool}
    for n in hierarchy:
        ids=leaves(n['节点ID']);n['下属三级能力ID']=base.list_ids(ids)
        n['下属涉及证据类别']=base.list_ids(tier for id in ids for tier in bycap[id]['具有的关联证据类别'].split(';'))
        n['证据层级边界']='上位类别是从已提炼三级候选归纳；含A不代表上位全部活动都由招聘明确支持'
        n['研究状态']=STATUS
    return pool,hierarchy

def reserve(data,rules):
    knowledge=[];by_id={r['条目ID']:r for r in data['knowledge']}
    for r in data['knowledge']:
        knowledge.append({'储备知识ID':'KR-'+r['条目ID'],**r,'储备来源':'第三阶段课程标准结构化知识表，不是岗位证据','岗位A/B/C类别':'不适用（独立储备）','岗位需求关系':'不生成','教学解释限度':'目标/思政提及不能当作已确认实训支撑' if r['课程条目定位']!='教学项目知识条目' else '课程文本知识储备，不证明学生已掌握或岗位存在需求','研究状态':STATUS})
    teaching=[]
    descriptions={'RS001':'在教学情境区分数据、数据库、DBMS等概念并说明SQL作用','RS002':'在课程情境安装配置MySQL并建立连接，不反推岗位指定产品','RS003':'在教学案例中表达概念模型并转换为关系模式','RS004':'在教学情境分析数据库结构的规范化及1NF、2NF、3NF','RS005':'在教学表结构中讨论并设置主键、唯一键、外键及CHECK约束','RS006':'在教学数据中使用DELETE与TRUNCATE并说明操作影响','RS007':'在教学情境编写子查询与联合查询并检查结果','RS008':'在教学情境创建、查询、修改、删除视图，并解释视图操作限制','RS009':'在教学情境说明ACID并使用BEGIN、COMMIT、ROLLBACK','RS010':'解释权限管理和DCL的基础作用；目标层提及不等于完整授权实训'}
    for id,name,kids,reason in rules['course_reserve']:
        if not set(kids)<=set(by_id):raise ValueError('储备知识ID不存在')
        rs=[by_id[k] for k in kids]
        teaching.append({'储备ID':id,'教学情境能力储备名称':name,'能力说明（研究拟定）':descriptions[id]+'；仅为后续教学设计储备，非岗位需求或已确认培养结果','知识条目ID':';'.join(kids),'课程原文摘录':'\n'.join(r['课程标准原文摘录'] for r in rs),'课程证据ID':base.list_ids(x for r in rs for x in str(r['课程证据ID']).split(';')),'PDF物理页码':base.list_ids(r['PDF页码'] for r in rs),'保留为储备的理由':reason,'岗位A/B/C类别':'不适用（独立储备）','进入岗位候选池或需求频次':'否','审核状态':STATUS})
    return knowledge,teaching

def check_consistency(tasks,links,pool,hierarchy,knowledge,teaching):
    tids={t['提取任务ID'] for t in tasks};pids={c['能力编号'] for c in pool}
    if len(tids)!=len(tasks) or len(pids)!=len(pool):raise ValueError('任务或能力ID重复')
    if len({c['三级能力名称'] for c in pool})!=len(pool):raise ValueError('能力名称重复节点')
    if any(r['能力编号'] not in pids or r['提取任务ID'] not in tids for r in links):raise ValueError('任务能力关系端点失效')
    nids={n['节点ID'] for n in hierarchy}
    if len(nids)!=len(hierarchy) or any(n['父编码'] and n['父编码'] not in nids for n in hierarchy):raise ValueError('层级编码重复或父节点无效')
    if {n['节点ID'] for n in hierarchy if n['能力层级']=='三级'}!=pids:raise ValueError('B/C未完整参与层级')
    if any(r['原任务类别']=='X' for r in links):raise ValueError('排除任务不能强建数据库关联')
    if any(c['岗位A/B/C类别']!='不适用（独立储备）' for c in knowledge+teaching):raise ValueError('课程储备混入岗位分类')
    for c in pool:
        actual=sorted({r['关联证据类别'] for r in links if r['能力编号']==c['能力编号']})
        if actual!=c['具有的关联证据类别'].split(';'):raise ValueError('多类证据集合不一致')

def run(root=ROOT,output=None,verify=True):
    root=Path(root).resolve();output=Path(output or HERE/'output').resolve()
    if output.is_relative_to(root) and not output.is_relative_to(HERE):raise ValueError('不得覆盖初稿、旧研究或网页目录')
    snapshot=json.loads((HERE/'input_snapshot.json').read_text());rules=json.loads((HERE/'rules.json').read_text())
    if verify:
        for f in snapshot['repository_files']:
            # 这里只锁定实际研究输入；测试另验证整个104文件基线未改变。
            if f['path'].endswith('.xlsx') or f['path'] in [str((V1/'rules.json').relative_to(root)),str((V1/'build.py').relative_to(root))]:
                if base.git_sha(root/f['path'])!=f['sha']:raise ValueError('来源文件偏离输入快照：'+f['path'])
    data=load_inputs(root);tasks=prepare_tasks(data,rules);links=make_links(tasks,rules)
    pool,hier=pool_and_hierarchy(links,rules)
    # 重要：岗位候选池已独立构建完成，之后才读取课程储备构建结果；知识表不参与触发候选能力。
    knowledge,teaching=reserve(data,rules);check_consistency(tasks,links,pool,hier,knowledge,teaching)
    row_by_id={c['能力编号']:c for c in pool}
    def category_view(tier):
        return [{**c,'本分表类别':tier,'本分表性质':LABELS[tier],'本类关联ID':base.list_ids(r['关联ID'] for r in links if r['能力编号']==c['能力编号'] and r['关联证据类别']==tier),'分表去重说明':'同一能力可出现在多分表，但合并池仅一个节点；不同关联类别不被最高类别覆盖'} for c in pool if tier in c['具有的关联证据类别'].split(';')]
    typical=[]
    for t in data['typical']:
        ex=[r for r in tasks if r['归纳典型任务ID']==t['典型工作任务ID']]
        typical.append({**t,'本轮任务类别集合':base.list_ids(r['岗位任务关联类别'] for r in ex),'本轮提取片段ID':base.list_ids(r['提取任务ID'] for r in ex),'本轮关联能力ID':base.list_ids(r['能力编号'] for r in links if r['典型工作任务ID']==t['典型工作任务ID']),'本轮边界':'典型任务内可含不同技术明确程度，以片段类别和逐条关联为准，不用单一强标签覆盖','本轮状态':STATUS})
    selection=[]
    for t in tasks:
        rs=[r for r in links if r['提取任务ID']==t['提取任务ID']]
        selection.append({'提取任务ID':t['提取任务ID'],'典型工作任务ID':t['归纳典型任务ID'],'岗位ID':t['岗位ID'],'摘要片段':t['摘要证据文本'],'任务类别':t['岗位任务关联类别'],'关联能力ID':base.list_ids(r['能力编号'] for r in rs) or '无','处理说明':'进入候选关联；逐条保留解释与采用条件' if rs else '未形成具体候选；不把全部业务数据、模型、看板或系统任务都解释成数据库技术','研究状态':STATUS})
    transitions=[]
    rmap={'R001':['DB007'],'R002':['DB012'],'R003':['DB015'],'R004':['DB009','DB010'],'R005':['DB014'],'R006':['DB006','DB016']}
    for old in data['old_core']+data['old_pending']:
        id=old['能力编号'];newids=[id] if id in row_by_id else rmap.get(id,[])
        transitions.append({'初稿编号':id,'初稿名称':old['三级能力名称'],'新版能力ID':base.list_ids(newids),'本轮处理':'保留同一能力节点，加入多类任务关联' if id in row_by_id else '从隔离待审线索转入B/C候选结构；条件解释仍待审查','新版证据类别集合':base.list_ids(tier for nid in newids for tier in row_by_id[nid]['具有的关联证据类别'].split(';')),'说明':'未修改初稿，未将推断改为原文事实','研究状态':STATUS})
    covered={i for r in transitions for i in r['新版能力ID'].split(';') if i}
    for c in pool:
        if c['能力编号'] not in covered:transitions.append({'初稿编号':'无对应独立候选','初稿名称':'无','新版能力ID':c['能力编号'],'本轮处理':'有任务来源的关系型细化备选；非按课程目录补齐','新版证据类别集合':c['具有的关联证据类别'],'说明':c['研究者需审查的问题'],'研究状态':STATUS})
    hierarchy_edges=[{'关系ID':'ABC-H'+str(i).zfill(3),'上位ID':n['父编码'],'下位ID':n['节点ID'],'关系类型':'候选层级包含','下属证据类别':n['下属涉及证据类别'],'依据':'自下而上归纳具体技术活动；不是教材章节或已确认岗位能力关系','研究状态':STATUS} for i,n in enumerate([n for n in hier if n['父编码']],1)]
    original_logical_path,original_logical=unique_sheet(V1/'output','候选非层级关系',['关系ID','起点能力ID'])
    logical=[{**r,'本轮处理':'沿用为研究推断待审，未增加权重、未升级确认','本轮端点证据类别':row_by_id[r['起点能力ID']]['具有的关联证据类别']+' → '+row_by_id[r['终点能力ID']]['具有的关联证据类别']} for r in original_logical]
    questions=[{'问题ID':'ABC-Q001','涉及对象':'全部任务','审查问题':'源质量等级A/B/C和本轮岗位关联类别A/B/C是不同维度；岗位均未核验，不能混淆','处理边界':'保留原质量和状态，不新增逐条招聘原文核验任务'},
               {'问题ID':'ABC-Q002','涉及对象':'全部B/C候选','审查问题':'已经参与候选结构，不因未写关系型或产品名排除；条件用于审查最终采用，不是排除候选资格','处理边界':'只解释可能实现，不称为岗位明示要求'},
               {'问题ID':'ABC-Q003','涉及对象':'EX2-JOB012-R01;EX2-JOB012-R02','审查问题':'同一任职引文拆成取数/清洗维度，不是新增岗位；复杂SQL取数不等于多表连接','处理边界':'多表路线保持B，不把复杂等同JOIN'},
               {'问题ID':'ABC-Q004','涉及对象':'DB006;DB016','审查问题':'数据库更新可能是记录、结构或版本；两个候选是条件分支，不可作为两项明示需求相加','处理边界':'候选池允许并存，最终采用需审查实际对象'},
               {'问题ID':'ABC-Q005','涉及对象':'课程储备','审查问题':'课程原文仅进入储备，不用于触发ABC或岗位频次；目标/思政提及不作已确认实训依据','处理边界':'不更新旧课程支撑矩阵或评价'}]
    questions.append({'问题ID':'ABC-Q006','涉及对象':'DB001;DB017;DB005;DB011','审查问题':'查询目标与多表方法、优化目标与索引方法存在包含交叉；候选池允许细化备选，不宣称它们是互相独立岗位需求','处理边界':'研究者需审查最终层级和粒度，不因参考领域而保留冗余；不按节点数计算需求频次'})
    for c in pool:questions.append({'问题ID':'ABC-Q'+str(len(questions)+1).zfill(3),'涉及对象':c['能力编号'],'审查问题':c['研究者需审查的问题'],'处理边界':'候选，未验证、未定稿；行为评价是研究拟定'})
    files={
        '01_岗位任务ABC重分类与证据表.xlsx':{'全部片段及类别':tasks,'典型任务类别集合':typical,'全部片段纳入审阅':selection},
        '02_A类明确支持候选能力.xlsx':{'A类能力去重视图':category_view('A'),'A类逐条证据':[r for r in links if r['关联证据类别']=='A']},
        '03_B类数据库相关备选能力.xlsx':{'B类能力去重视图':category_view('B'),'B类技术解释证据':[r for r in links if r['关联证据类别']=='B']},
        '04_C类合理推断待审能力.xlsx':{'C类能力去重视图':category_view('C'),'C类推断证据':[r for r in links if r['关联证据类别']=='C']},
        '05_ABC合并候选三级能力体系.xlsx':{'合并三级候选池':pool,'自下而上全层级':hier,'候选层级关系':hierarchy_edges},
        '06_任务能力多类证据关联.xlsx':{'逐条任务能力证据':links},
        '07_课程知识与教学能力独立储备.xlsx':{'课程知识储备':knowledge,'教学情境能力储备':teaching},
        '08_前后对照与研究审查.xlsx':{'初稿到ABC对照':transitions,'研究者审查问题':questions,'非层级关系仍为待审':logical}}
    output.mkdir(parents=True,exist_ok=True)
    for name,sheets in files.items():base.workbook(output/name,sheets)
    views={tier:len(category_view(tier)) for tier in ['A','B','C']}
    summary={'source_pr':7,'source_pr_head':snapshot['pr7_head'],'source_main':snapshot['main_commit'],'source_jobs':len(data['selected']['jobs']['rows']),'v1_fragments':len(data['extracted']),'supplement_fragments':2,'all_fragments':len(tasks),'typical_tasks':len(typical),'fragment_class_counts':dict(Counter(t['岗位任务关联类别'] for t in tasks)),'unique_pool_capabilities':len(pool),'overlapping_category_view_counts':views,'overlapping_views_must_not_be_summed':True,'exclusive_count_by_strongest_existing_link':dict(Counter(c['去重计数归类（不升级其他证据）'] for c in pool)),'pool_hierarchy_counts':dict(Counter(n['能力层级'] for n in hier)),'task_capability_link_count':len(links),'link_class_counts':dict(Counter(r['关联证据类别'] for r in links)),'course_knowledge_reserve':len(knowledge),'teaching_reserve':len(teaching),'verified_recruitment_originals':0,'review_questions':len(questions),'source_files':{k:str(p.relative_to(root)) for k,p in data['source_paths'].items()},'outputs':list(files)}
    (output/'processing_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
    return summary

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path);args=p.parse_args();print(json.dumps(run(output=args.output),ensure_ascii=False,indent=2))
