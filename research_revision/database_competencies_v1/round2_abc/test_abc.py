"""候选池分类、证据追溯、分表交叉和储备隔离检查；不是招聘或研究结论验证。"""
import copy
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from collections import Counter
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent
spec=importlib.util.spec_from_file_location('abc_round2',HERE/'build.py');b=importlib.util.module_from_spec(spec);spec.loader.exec_module(b)

class ABCEvidenceTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory();cls.out=Path(cls.tmp.name);cls.summary=b.run(output=cls.out)
  cls.rules=json.loads((HERE/'rules.json').read_text());cls.data=b.load_inputs();cls.tasks=b.prepare_tasks(cls.data,cls.rules)
  cls.links=b.make_links(cls.tasks,cls.rules);cls.pool,cls.hier=b.pool_and_hierarchy(cls.links,cls.rules)
  cls.knowledge,cls.teaching=b.reserve(cls.data,cls.rules)
 @classmethod
 def tearDownClass(cls):cls.tmp.cleanup()
 def test_all_104_preexisting_files_unchanged(self):
  snapshot=json.loads((HERE/'input_snapshot.json').read_text())
  self.assertEqual(len(snapshot['repository_files']),104)
  for f in snapshot['repository_files']:self.assertEqual(b.base.git_sha(b.ROOT/f['path']),f['sha'],f['path'])
 def test_task_classification_and_supplement_quotes(self):
  self.assertEqual(len(self.tasks),181);self.assertEqual(Counter(t['岗位任务关联类别'] for t in self.tasks),{'A':16,'B':12,'C':100,'X':53})
  jobs={r['样本ID']:r for r in self.data['selected']['jobs']['rows']}
  for t in self.tasks:
   self.assertIn(t['摘要证据文本'],jobs[t['岗位ID']][t['来源摘要字段']]);self.assertEqual(t['招聘原文已核验'],'否')
  self.assertEqual(len({t['岗位ID'] for t in self.tasks}),50)
 def test_b_included_without_product_name(self):
  bids={t['提取任务ID'] for t in self.tasks if t['岗位任务关联类别']=='B'}
  self.assertTrue(bids<={r['提取任务ID'] for r in self.links})
  self.assertTrue(any(t['岗位任务关联类别']=='B' and t['同岗数据库产品语境'].startswith('未明确') for t in self.tasks))
  leaf={n['节点ID'] for n in self.hier if n['能力层级']=='三级'}
  self.assertTrue({r['能力编号'] for r in self.links if r['关联证据类别']=='B'}<=leaf)
 def test_c_not_all_forced_and_x_excluded(self):
  linked={r['提取任务ID'] for r in self.links};ctasks={t['提取任务ID'] for t in self.tasks if t['岗位任务关联类别']=='C'}
  self.assertTrue(ctasks-linked);self.assertTrue(ctasks&linked)
  self.assertFalse(any(r['原任务类别']=='X' for r in self.links))
  self.assertTrue(all(r['判断或推断理由'] and r['采用条件与替代解释'] for r in self.links))
 def test_unique_pool_keeps_multi_evidence_not_duplicate_nodes(self):
  self.assertEqual(len(self.pool),17);self.assertEqual(len({r['能力编号'] for r in self.pool}),17)
  self.assertEqual(len({r['三级能力名称'] for r in self.pool}),17)
  for id in ['DB001','DB002','DB003']:self.assertEqual(next(c for c in self.pool if c['能力编号']==id)['具有的关联证据类别'],'A;B;C')
  views=self.summary['overlapping_category_view_counts'];self.assertEqual(views,{'A':5,'B':13,'C':8})
  self.assertGreater(sum(views.values()),17)
  self.assertEqual(sum(self.summary['exclusive_count_by_strongest_existing_link'].values()),17)
 def test_link_tier_not_upgraded_from_goal_to_method(self):
  for r in self.links:
   if r['关联证据类别']=='A':self.assertEqual(r['原任务类别'],'A')
   if r['关联证据类别']=='B':self.assertIn(r['原任务类别'],['A','B'])
   if r['关联证据类别']=='C':self.assertEqual(r['原任务类别'],'C')
  for id in ['DB011','DB017']:
   rs=[r for r in self.links if r['能力编号']==id];self.assertTrue(rs);self.assertTrue(all(r['关联证据类别']=='B' for r in rs))
 def test_invalid_a_detail_claim_rejected(self):
  altered=copy.deepcopy(self.rules);altered['links'][0]={'cap':'DB017','tier':'A','groups':['TT020'],'jobs':None,'ids':None,'reason':'不得从性能目标声称JOIN明示','condition':'无'}
  with self.assertRaisesRegex(ValueError,'不能由未明示细化'):b.make_links(self.tasks,altered)
 def test_hierarchy_all_abc_and_valid_edges(self):
  ids={n['节点ID'] for n in self.hier};self.assertEqual(len(ids),len(self.hier))
  self.assertEqual(Counter(n['能力层级'] for n in self.hier),{'一级':3,'二级':9,'三级':17})
  for n in self.hier:
   if n['父编码']:self.assertIn(n['父编码'],ids)
  b.check_consistency(self.tasks,self.links,self.pool,self.hier,self.knowledge,self.teaching)
 def test_course_reserve_independent_of_pool(self):
  self.assertEqual(len(self.knowledge),37);self.assertEqual(len(self.teaching),10)
  self.assertTrue(all(r['岗位A/B/C类别']=='不适用（独立储备）' for r in self.knowledge+self.teaching))
  altered=copy.deepcopy(self.data);altered['knowledge']=[]
  tasks=b.prepare_tasks(altered,self.rules);links=b.make_links(tasks,self.rules)
  self.assertEqual(b.pool_and_hierarchy(links,self.rules)[0],self.pool)
  self.assertTrue(all(r['进入岗位候选池或需求频次']=='否' for r in self.teaching))
 def test_course_quotes_preserved_and_page_evidence_not_claimed_practice(self):
  byid={r['条目ID']:r for r in self.data['knowledge']}
  for r in self.knowledge:
   self.assertEqual(r['课程标准原文摘录'],byid[r['条目ID']]['课程标准原文摘录'])
   self.assertEqual(r['PDF页码'],byid[r['条目ID']]['PDF页码'])
   if r['课程条目定位']!='教学项目知识条目':self.assertIn('不能当作已确认实训',r['教学解释限度'])
 def test_workbooks_open_and_class_views_have_correct_rows(self):
  self.assertEqual(len(list(self.out.glob('*.xlsx'))),8)
  for p in self.out.glob('*.xlsx'):
   w=load_workbook(p,read_only=True);self.assertTrue(w.sheetnames)
   for sheet in w:self.assertGreaterEqual(sum(1 for _ in sheet.values),1)
   w.close()
  p=self.out/'03_B类数据库相关备选能力.xlsx';rows=b.base.read(p,'B类能力去重视图');self.assertEqual(len(rows),13)
 def test_reproducible_content_and_old_output_protection(self):
  second=self.out/'second';self.assertEqual(b.run(output=second),self.summary)
  for name in self.summary['outputs']:
   w=load_workbook(self.out/name,read_only=True);sheets=w.sheetnames;w.close()
   for sheet in sheets:self.assertEqual(b.base.read(self.out/name,sheet),b.base.read(second/name,sheet))
  with self.assertRaisesRegex(ValueError,'不得覆盖'):b.run(output=b.V1/'output')

if __name__=='__main__':unittest.main()
