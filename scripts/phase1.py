"""旅游岗位候选样本第一阶段：离线、保守、可追溯的任务和能力归纳。"""
import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SOURCE_NAME = '旅游行业招聘岗位样本库_50条候选_待核验.xlsx'
DUTY = '岗位职责摘要（转述）'
REQ = '任职要求摘要（不含学历经验）'
REQUIRED = ['样本ID', '相关性类别', '招聘单位', '岗位名称', DUTY, REQ,
            '原始招聘链接', '证据质量（初评）', '核验状态', '证据摘录']
# 每行：三级编码、一级、二级、三级名称、可观察指标、任务触发词。
# 编码体系是研究者对本批摘要的归纳，不是课程标准或官方职业标准。
RULES = [
 ('C111','数据资源处理','数据获取与质量','采集接入业务数据','能按业务来源采集、整理或接入数据','采集|收集|接入|整理'),
 ('C112','数据资源处理','数据获取与质量','清洗与标准化数据','能识别并处理缺失、异常及格式不一致问题','清洗|标准化|异常与缺失'),
 ('C113','数据资源处理','数据获取与质量','提取查询业务数据','能围绕业务问题提取或查询所需数据','提取|查询'),
 ('C121','数据资源处理','数据组织与治理','管理指标口径与数据质量','能记录指标定义、口径并检查数据质量','口径|数据质量|数据字典|数据治理|核对|数据链路'),
 ('C122','数据资源处理','数据组织与治理','维护业务数据资源','能维护、更新和同步业务或会员数据','数据资源管理|数据管理|数据库更新|同步'),
 ('C123','数据资源处理','数据组织与治理','组织会员标签与画像','能依据业务对象组织标签与画像','标签|画像'),
 ('C211','业务数据分析','指标与统计分析','构建和监测业务指标','能构建或监测客流、营收、转化等业务指标','指标|客流监测|数据监测|跟踪.*(转化|增长|留存)|统计|经营.*监控|数据监控'),
 ('C212','业务数据分析','指标与统计分析','分析旅游业务数据','能结合经营场景解释数据和业务变化','分析|诊断|归因|评估'),
 ('C231','业务数据分析','模型与决策支持','构建分析与预测模型','能构建用于分析或预测的模型','建模|模型|预测'),
 ('C232','业务数据分析','模型与决策支持','形成运营改进建议','能依据分析结果提出运营、成本或收益改进建议','策略|决策|成本优化|运营.*优化|经营优化|改进|复盘|收益管理'),
 ('C233','业务数据分析','模型与决策支持','实施业务增长改进','能围绕产品转化与用户增长推进业务运营改进','提升.*转化|用户增长'),
 ('C221','业务数据分析','分析成果表达','制作报表与可视化','能将数据整理为报表、看板或可视化成果','报表|报告|看板|可视化|统计汇总'),
 ('C311','信息系统保障','系统运维与配置','维护与配置业务系统','能维护、配置旅游业务系统并记录运行情况','运维|维护|配置|巡检|系统监控'),
 ('C312','信息系统保障','系统运维与配置','排查并处理系统故障','能定位系统或网络问题并处理故障','故障|排除|排查'),
 ('C331','信息系统保障','数据安全与恢复','备份恢复业务数据','能实施备份并按需要恢复数据','备份|恢复'),
 ('C332','信息系统保障','数据安全与恢复','实施权限与安全管理','能管理权限并检查安全事件','权限|安全|审计'),
 ('C321','信息系统保障','数据库管理','管理数据库应用系统','能开展数据库或数据库应用系统管理','数据库日常运维|数据库应用系统管理|数据库运维'),
 ('C411','业务系统建设','需求与应用实现','分析业务系统需求','能整理业务需求和功能约束','系统需求|产品规划|功能管理|产品设计'),
 ('C412','业务系统建设','需求与应用实现','开发数据库业务应用','能围绕酒店或旅游业务实现系统功能','系统开发|开发数据库|数据库应用系统.*开发|开发数据报表'),
 ('C431','业务系统建设','测试与交付','测试升级与上线系统','能支持配置测试、系统升级和上线交付','测试|上线|升级|实施'),
 ('C421','业务系统建设','数据库管理与优化','优化数据库查询性能','能识别并优化业务系统数据库查询性能','查询性能优化'),
 ('C511','业务协作与研究','调研与协作','开展旅游业务调研','能组织业务、市场或游客调研','调研|调查'),
 ('C512','业务协作与研究','调研与协作','协调业务与数据支持','能沟通跨部门需求并提供业务支持','协作|协调|跨部门|项目支持|项目实施'),
 ('C531','业务协作与研究','服务改进','跟进游客服务整改','能跟进服务问题并记录整改进展','服务整改'),
 ('C521','业务协作与研究','文档与流程','维护技术文档与流程','能记录系统文档或整理处理流程','文档|流程'),
 ('C131','数据资源处理','查询工具应用','使用SQL处理和查询数据','能使用SQL完成岗位相关数据处理或查询','SQL'),
]
# 只依据任职摘要中的正向表述标记；优先和否定按局部短句判定。
REQ_PATTERNS = {
 'C111': '数据接入|数据整理', 'C112': '清洗', 'C113': '提取|查询',
 'C121': '数据治理|指标管理|数据指标|数据质量|数据字典',
 'C122': '数据管理', 'C123': '标签|画像', 'C211': '指标体系|指标管理',
 'C212': '分析|统计分析', 'C231': '模型|建模', 'C232': '运营决策',
 'C221': '可视化|BI工具|Power BI|Tableau|帆软|报表',
 'C311': '运维|系统配置', 'C312': '故障', 'C331': '备份|恢复',
 'C332': '安全|权限', 'C321': '数据库.*管理|数据库运维',
 'C411': '需求分析|产品设计', 'C412': '数据库应用系统管理开发|开发数据库',
 'C131': r'(?<!My)SQL(?! Server)',
}
NOTE = '候选样本；未联网逐条核验；不得解释为正式核验样本或全国代表性结论'


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def norm(v):
    return re.sub(r'[\s·•（）()\-—_]', '', unicodedata.normalize('NFKC', str(v or ''))).casefold()


def canonical_url(value):
    """仅去跟踪参数，保留所有可能区分岗位的参数，包括 id/jobUnionId。"""
    if blank(value):
        return ''
    p = urlsplit(str(value).strip())
    query = [(k, v) for k, v in parse_qsl(p.query, keep_blank_values=True)
             if not k.lower().startswith('utm_') and k.lower() not in {'spm', 'from', 'highlighttype'}]
    return urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path.rstrip('/'), urlencode(sorted(query)), ''))


def read_source(path):
    wb = load_workbook(path, data_only=False)
    if '招聘岗位主库' not in wb.sheetnames:
        raise ValueError('缺少工作表：招聘岗位主库')
    values = list(wb['招聘岗位主库'].values)
    headers = list(values[0])
    if len(headers) != len(set(headers)) or any(blank(x) for x in headers):
        raise ValueError('主库字段名为空或重复')
    missing = set(REQUIRED) - set(headers)
    if missing:
        raise ValueError('缺少必需字段：' + ','.join(sorted(missing)))
    jobs = []
    for rowno, row in enumerate(values[1:], 2):
        if all(blank(v) for v in row):
            continue
        item = dict(zip(headers, row))
        item['源行号'] = rowno
        jobs.append(item)
    ids = [j['样本ID'] for j in jobs]
    if any(blank(i) for i in ids) or len(ids) != len(set(ids)):
        raise ValueError('岗位ID为空或重复；请先人工解决，禁止静默改写')
    return wb, headers, jobs


def duplicates(jobs):
    rows = []
    checks = [('标准化招聘链接', lambda j: canonical_url(j['原始招聘链接'])),
              ('同单位同岗位', lambda j: (norm(j['招聘单位']), norm(j['岗位名称']))),
              ('重复单位（不等于重复招聘）', lambda j: norm(j['招聘单位'])),
              ('重复岗位名称（不等于重复招聘）', lambda j: norm(j['岗位名称'])),
              ('单位简称线索+同岗位', lambda j: (norm(j['招聘单位']).replace('文化旅游', '文旅'), norm(j['岗位名称'])))]
    for kind, fn in checks:
        groups = defaultdict(list)
        for j in jobs:
            key = fn(j)
            if not key or (isinstance(key, tuple) and not all(key)):
                continue
            groups[key].append(j)
        for key, group in groups.items():
            if len(group) < 2:
                continue
            # 简称组只保留严格单位匹配未覆盖的组。
            if kind == '单位简称线索+同岗位' and len({norm(j['招聘单位']) for j in group}) == 1:
                continue
            gid = f'D{len(rows)+1:03d}'
            conflict = len({j['相关性类别'] for j in group}) > 1
            for j in group:
                rows.append({'重复组': gid, '检查类型': kind, '匹配键': str(key), '组内记录数': len(group),
                             '岗位ID': j['样本ID'], '招聘单位': j['招聘单位'], '岗位名称': j['岗位名称'],
                             '原始招聘链接': j['原始招聘链接'], '相关性类别': j['相关性类别'],
                             '分类冲突': '是' if conflict else '否', '处理建议': '保留；人工确认是否同一招聘事件'})
    return rows


def classification(j):
    text = str(j[DUTY] or '') + '；' + str(j[REQ] or '')
    if re.search('SQL|数据库|清洗|建模|数据治理|数据提取|指标口径', text, re.I):
        basis = '摘要涉及数据库、查询、数据处理或治理任务，与课程主题有直接联系（待原文复核）'
        suggested = '直接相关岗位'
    elif re.search('分析|统计|报表|系统|运维|数据', text):
        basis = '摘要涉及数据使用或信息系统业务；是否涉及数据库核心操作尚需核验'
        suggested = '关联岗位'
    else:
        basis = '摘要不足，相关性待研究者复核'
        suggested = '待复核'
    return {'岗位ID': j['样本ID'], '原分类（保留）': j['相关性类别'], '复核建议分类': suggested,
            '分类依据': basis, '触发证据': text, '是否与原分类不同': '是' if suggested != j['相关性类别'] else '否',
            '分类说明': '建议是研究判断，不覆盖原分类；直接相关不等于正式核验或已满足课程标准'}


def extract_tasks(jobs):
    tasks = []
    for j in jobs:
        # 按摘要分号、句号、中文逗号分段；保留枚举内部的顿号与英文标识。
        chunks = [s.strip() for s in re.split('[；;。\n，]', str(j[DUTY] or '')) if s.strip()]
        for n, evidence in enumerate(chunks, 1):
            tasks.append({'任务编码': f"T-{j['样本ID']}-{n:02d}", '岗位ID': j['样本ID'],
                          '典型工作任务': evidence, '提取方式': '职责摘要分句；保留原有粒度，复合任务待原文细分',
                          '证据文本': evidence, '证据类型': '摘要证据', '证据字段': DUTY,
                          '证据位置': f"招聘岗位主库!G{j['源行号']}",
                          '任务性质': '职责摘要明示任务（未核验）',
                          '原始招聘链接': j['原始招聘链接'], '相关性类别': j['相关性类别'],
                          '证据质量（初评）': j['证据质量（初评）'], '核验状态': '待核验'})
    return tasks


def requirement_evidence(text, code):
    pattern = REQ_PATTERNS.get(code)
    if not pattern:
        return []
    result = []
    for segment in re.split('[；;。\n，]', str(text or '')):
        if re.search('未明确|未单独|未.*写明|待.*核验|需核对|不要求|无需', segment):
            continue
        # 顿号分割避免 SQL/Python 优先影响前面的 Excel；保留斜线工具组。
        for part in re.split('[、,]', segment):
            if re.search(pattern, part, re.I):
                result.append((part, '优先要求（摘要明示）' if '优先' in part else '明确要求（摘要明示）'))
    return result


def map_abilities(jobs, tasks):
    byid = {j['样本ID']: j for j in jobs}
    relations = []
    for task in tasks:
        j = byid[task['岗位ID']]
        matched = []
        for code, first, second, leaf, indicator, pattern in RULES:
            hit = re.search(pattern, task['证据文本'], re.I)
            # SQL 是工具要求，只有与数据处理/分析任务相关时才附加；连接仍为研究推断。
            sql_requirement = code == 'C131' and re.search('分析|提取|查询|数据处理', task['证据文本']) and requirement_evidence(j[REQ], code)
            if not hit and not sql_requirement:
                continue
            matched.append(code)
            reqs = requirement_evidence(j[REQ], code)
            evidence, nature = reqs[0] if reqs else ('', '研究推断（由职责任务归纳）')
            relations.append({'岗位ID': j['样本ID'], '任务编码': task['任务编码'], '三级能力编码': code,
                              '一级能力': first, '二级能力': second, '三级能力': leaf,
                              '任务证据': task['证据文本'], '任职要求证据': '；'.join(x[0] for x in reqs),
                              '能力要求性质': nature, '任务到能力关系性质': '研究推断',
                              '归纳依据': f'规则{code}；任务触发词：{hit.group() if hit else "数据任务+任职摘要SQL"}；能力表述与层级由研究者归纳',
                              '证据类型': '摘要证据', '原始招聘链接': j['原始招聘链接'],
                              '相关性类别': j['相关性类别'], '证据质量（初评）': j['证据质量（初评）'], '核验状态': '待核验'})
        task['能力映射状态'] = '已映射（研究归纳）' if matched else '待人工归纳（摘要过于概括）'
    indicators = []
    for code, first, second, leaf, indicator, pattern in RULES:
        linked = [r for r in relations if r['三级能力编码'] == code]
        if not linked:
            continue
        indicators.append({'一级编码': code[:2], '一级能力': first, '二级编码': code[:3], '二级能力': second,
                           '三级能力编码': code, '三级能力': leaf, '可观察能力指标（研究拟定）': indicator,
                           '归纳属性': '研究归纳，待专家复核；非课程标准', '任务触发规则': pattern,
                           '支撑岗位数（含疑似重复）': len({r['岗位ID'] for r in linked}),
                           '支撑任务数': len({r['任务编码'] for r in linked}),
                           '摘要明示要求关系数': sum(r['能力要求性质'].startswith('明确') for r in linked),
                           '摘要优先要求关系数': sum(r['能力要求性质'].startswith('优先') for r in linked),
                           '研究推断关系数': sum(r['能力要求性质'].startswith('研究') for r in linked),
                           '支撑岗位ID': ';'.join(sorted({r['岗位ID'] for r in linked})),
                           '支撑任务编码': ';'.join(sorted({r['任务编码'] for r in linked})),
                           '结论状态': '候选指标；全部支撑岗位待核验'})
    return indicators, relations


def write_book(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows, headers in sheets:
        ws = wb.create_sheet(title[:31])
        if headers is None:
            headers = list(rows[0]) if rows else ['说明']
        ws.append(headers)
        for row in rows:
            ws.append([row.get(k) for k in headers])
        # 显式字符串，防止原始证据被解释为 Excel 公式。
        for row in ws:
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = 's'
                c.alignment = Alignment(vertical='top', wrap_text=True)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='245A81')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cells in ws.columns:
            ws.column_dimensions[cells[0].column_letter].width = min(55, max(14, len(str(cells[0].value)) * 2))
    wb.save(path)


def run(source, output):
    source, output = Path(source), Path(output)
    wb, headers, jobs = read_source(source)
    output.mkdir(parents=True, exist_ok=True)
    inventory, fields = [], []
    for ws in wb:
        values = list(ws.values)
        nonempty = [r for r in values[1:] if not all(blank(v) for v in r)]
        inventory.append({'工作表': ws.title, '实际字段数': len(values[0]), '有效记录数': len(nonempty),
                          '物理总行数（含表头）': ws.max_row, '字段': '；'.join(str(h) for h in values[0])})
        for i, h in enumerate(values[0]):
            count = sum(blank(r[i]) for r in nonempty)
            fields.append({'工作表': ws.title, '字段': h, '记录数': len(nonempty), '缺失数': count,
                           '缺失比例': count / len(nonempty) if nonempty else 0})
    dup = duplicates(jobs)
    tasks = extract_tasks(jobs)
    indicators, relations = map_abilities(jobs, tasks)
    enriched = [dict(j, **{'本轮核验状态': '待核验', '本轮证据类型': '摘要证据',
                         '本轮说明': NOTE, '原文摘录提示': '原证据摘录仅透传；本轮任务仅使用职责摘要'}) for j in jobs]
    issues = []
    for j in jobs:
        for h in headers:
            if blank(j[h]):
                issues.append({'岗位ID': j['样本ID'], '问题类型': '字段缺失', '字段': h, '说明': '原值保留为空，未补写'})
        if j['证据质量（初评）'] == 'A':
            issues.append({'岗位ID': j['样本ID'], '问题类型': '等级解释', '字段': '证据质量（初评）',
                           '说明': '原A仅代表原表所称官方详情；无逐字原文与本轮核验，不能升级为已核验'})
    notes = [{'项目': '研究边界', '说明': NOTE},
             {'项目': '明示定义', '说明': '明确/优先仅指任职摘要出现正向表述；不是招聘原文核验结论。职责到能力及三级层级均为研究归纳。'},
             {'项目': '缺失定义', '说明': 'None、空字符串、全空白为缺失；占位文本不计为空但应人工审查。'},
             {'项目': '分类定义', '说明': '原直接/关联分类保留；复核建议依据摘要中数据库、查询、处理、治理及业务系统任务，不依据虚构课程标准。'},
             {'项目': '原证据等级', '说明': 'A/B/C完全透传。原表A为官方详情线索，B为平台详情或历史样本，C为转载聚合等；均待逐项核验。'},
             {'项目': '重复解释', '说明': '重复单位或岗位名仅是线索；同单位同岗位及简称线索也不自动合并，可能属于不同批次、日期或城市。'},
             {'项目': '原重复来源追踪', '说明': '原表已有20条历史来源追踪，不计入当前50条主库；单独透传。'},
             {'项目': '未匹配任务', '说明': '保留在任务表和关系表的未映射任务页，不使用通用能力强行填充。'}]
    sheets = [('汇总说明', notes, None), ('工作表盘点', inventory, None), ('字段缺失统计', fields, None),
              ('岗位原始与本轮状态', enriched, headers + ['源行号', '本轮核验状态', '本轮证据类型', '本轮说明', '原文摘录提示']),
              ('逐条问题', issues, None), ('疑似重复记录', dup, ['重复组','检查类型','匹配键','组内记录数','岗位ID','招聘单位','岗位名称','原始招聘链接','相关性类别','分类冲突','处理建议']),
              ('分类依据与复核', [classification(j) for j in jobs], None)]
    for name in wb.sheetnames:
        if name != '招聘岗位主库':
            values = list(wb[name].values)
            sheets.append(('原表_' + name, [dict(zip(values[0], r)) for r in values[1:]], list(values[0])))
    write_book(output / '岗位数据质量检查.xlsx', sheets)
    task_headers = ['任务编码','岗位ID','典型工作任务','提取方式','证据文本','证据类型','证据字段','证据位置','任务性质','原始招聘链接','相关性类别','证据质量（初评）','核验状态','能力映射状态']
    write_book(output / '岗位任务编码表.xlsx', [('任务编码', tasks, task_headers), ('方法说明', notes, None)])
    aggregate = []
    for level, key, idkey in [('一级', '一级能力', '一级编码'), ('二级', '二级能力', '二级编码')]:
        for code in sorted({i[idkey] for i in indicators}):
            subset = [i for i in indicators if i[idkey] == code]
            aggregate.append({'层级': level,'编码': code,'能力名称': subset[0][key],
                              '下属三级能力编码': ';'.join(i['三级能力编码'] for i in subset),
                              '聚合依据': '按任务处理对象与业务目的归并；研究归纳，待专家审核'})
    write_book(output / '三级职业能力指标体系.xlsx', [('三级能力指标', indicators, None), ('一级二级聚合', aggregate, None), ('方法说明', notes, None)])
    write_book(output / '岗位任务能力关系表.xlsx', [('关系明细', relations, None),
               ('未映射任务', [t for t in tasks if t['能力映射状态'].startswith('待')], task_headers), ('方法说明', notes, None)])
    report = {'源文件': source.name, '源文件SHA256': hashlib.sha256(source.read_bytes()).hexdigest(),
              '主库岗位数': len(jobs), '分类': dict(Counter(j['相关性类别'] for j in jobs)),
              '本轮待核验岗位数': len(jobs), '任务数': len(tasks), '一级能力数': len({i['一级编码'] for i in indicators}),
              '二级能力数': len({i['二级编码'] for i in indicators}), '三级能力数': len(indicators),
              '关系数': len(relations), '未映射任务数': sum(t['能力映射状态'].startswith('待') for t in tasks),
              '重复组数': {kind: len({r['重复组'] for r in dup if r['检查类型'] == kind}) for kind in ['标准化招聘链接','同单位同岗位','重复单位（不等于重复招聘）','重复岗位名称（不等于重复招聘）','单位简称线索+同岗位']},
              '主库字段缺失': {h: sum(blank(j[h]) for j in jobs) for h in headers}, '工作表': inventory,
              '限制': NOTE}
    (output / '处理报告.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', default=SOURCE_NAME)
    parser.add_argument('--output', default='output')
    args = parser.parse_args()
    run(args.input, args.output)
