"""只从原始PDF+人工定位清单提取；保持教学内容、活动、产物与评价来源分层。"""
from __future__ import annotations
import argparse,hashlib,json,re,shutil
from pathlib import Path
from collections import Counter,defaultdict
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from datetime import datetime
HERE=Path(__file__).resolve().parent;ROOT=HERE.parent
PDF_SHA='8e7b5c9669b5ab0345d74b25e924085f7f9185c6'
STATUS='待研究者审查；非已确认教学效果'
HOUR_RULE='学时同源、不重复累计：8项目实践32与3实训32为同一批实践学时；课程64=理论32+实践32，禁止相加为96。'
COMPLEMENT='岗位能力体系与课程内容体系具有不同的形成逻辑和功能定位，二者相互关联、相互补充，也各有局限。'

def compact(s):return re.sub(r'\s+','',s)
def gitsha(path):
 b=path.read_bytes();return hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()
def protect():
 snapshot=json.loads((HERE/'input_snapshot.json').read_text())
 for f in snapshot['repository_files']:
  if gitsha(ROOT/f['path'])!=f['sha']:raise ValueError('历史文件发生变化：'+f['path'])
 return snapshot

def locate_pdf():
 matches=[p for p in ROOT.rglob('课标0913.pdf') if not any(x in p.parts for x in ['.git','work','.venv']) and gitsha(p)==PDF_SHA]
 if len(matches)!=1:raise ValueError('原始PDF缺失、版本变化或不唯一，不允许用旧摘要替代：'+str(matches))
 return matches[0]

def write_book(path,sheets):
 w=Workbook();w.remove(w.active);w.properties.created=w.properties.modified=datetime(2026,10,24)
 for title,rows in sheets.items():
  s=w.create_sheet(title[:31]);keys=list(dict.fromkeys(k for r in rows for k in r))
  if not keys:keys=['说明'];rows=[{'说明':'无记录'}]
  s.append(keys)
  for row in rows:
   s.append([json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(dict,list)) else row.get(k) for k in keys])
   for cell in s[s.max_row]:
    if isinstance(cell.value,str):cell.data_type='s'
  for c in s[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='23566D')
  s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
  for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(65,max(16,len(str(col[0].value))*2))
 w.save(path)

class Evidence:
 def __init__(self,doc,path):
  self.doc=doc;self.path=path;self.raw=[p.get_text() for p in doc];self.normal=[compact(s) for s in self.raw];self.records={};self.links=[]
 def quote(self,page,needle):
  text=self.raw[page-1];target=compact(needle);normal=self.normal[page-1];start=normal.find(target)
  if not target or start<0:raise ValueError(f'原文不存在 p{page}: {needle}')
  positions=[i for i,c in enumerate(text) if not c.isspace()];return text[positions[start]:positions[start+len(target)-1]+1]
 def add(self,page,quote,entity,section,nature='原文明示',note=''):
  raw=self.quote(page,quote);normal=compact(raw);key=(page,normal)
  if key not in self.records:
   # 依页码及原文内容生成稳定证据号；增补其他证据不改变已有编号。
   id='CE'+str(page).zfill(2)+'-'+hashlib.sha256(normal.encode()).hexdigest()[:10]
   hidden=page==13 and self.normal[12].find(normal)>=self.normal[12].index('6"8')
   self.records[key]={'证据ID':id,'文档ID':'DOC01','来源文件名':self.path.name,'仓库实际路径':str(self.path.relative_to(ROOT)),'PDF实际页码':page,'原文短摘录':raw,'提取方式':'PyMuPDF原生文本层；13页渲染浏览，异常页专项核对','可信状态':'文本层存在但页面不可见，尚未可视确认' if hidden else '可见正文文本层摘录；已作页面浏览，待研究者审查','是否存在歧义':'是，隐藏文本层' if hidden else ('是，详见备注' if note else '暂未发现明显提取歧义'),'来源证据层':'隐藏文本候选' if hidden else '可见PDF原文','研究者审查状态':'待审查','摘录说明':note}
  rec=self.records[key];self.links.append({'实体ID':entity,'证据ID':rec['证据ID'],'所属章节或项目':section,'使用性质':nature,'说明':note,'研究者审查状态':'待审查'})
  return rec['证据ID']
 def fields(self,ids):
  ids=list(dict.fromkeys(ids));rs=[r for r in self.records.values() if r['证据ID'] in ids]
  return {'证据ID':';'.join(ids),'来源文件名':self.path.name,'PDF实际页码':';'.join(str(p) for p in sorted({r['PDF实际页码'] for r in rs})),'原文摘录':'\n——\n'.join(r['原文短摘录'] for r in rs),'来源状态':';'.join(sorted({r['可信状态'] for r in rs})),'研究者审查状态':'待审查'}

def build(output=None):
 snapshot=protect();pdf=locate_pdf();doc=fitz.open(pdf)
 if len(doc)!=13:raise ValueError('不是完整13页PDF')
 e=Evidence(doc,pdf);manifest=json.loads((HERE/'extraction_manifest.json').read_text())
 entities=[];relations=[]
 def node(id,type,name):entities.append({'实体ID':id,'实体类型':type,'名称':name,'文档ID':'DOC01','审查状态':STATUS})
 def rel(a,b,kind,basis='原文结构',status='待研究者审查'):relations.append({'源实体ID':a,'目标实体ID':b,'关系类型':kind,'依据性质':basis,'审查状态':status})
 node('DOC01','课程标准文档','《数据库原理及应用》课程标准')
 info=[]
 basics=[('课程名称',1,'《数据库原理及应用》','《数据库原理及应用》'),('适用专业与年级',1,'2025级大数据技术专业','面向2025 级大数据技术专业'),('课程性质',1,'专业核心课程；岗课赛证融通背景下理实一体化课程','一门专业核心课程，属于‘岗课赛证’融通背景下的理实一体化课程。'),('课程作用',1,None,'其作用在于让学生理解关系型数据库的核心原理与架构，掌握数据库设计、创建、操作及数据库安全管理的核心技能'),('前导课程',1,'《Python程序设计基础》','前导课程：《Python 程序设计基础》'),('后续课程',1,'《NoSQL非关系型数据库》','后续课程：本课程为《NoSQL 非关系型数据库》课程提供理论与实操支撑。'),('课程教学理念',1,None,'课程坚持“理论引领、实践驱动、能力导向”的教育理念。')]
 for i,(name,pg,value,quote) in enumerate(basics,1):
  id=f'CB{i:03}';node(id,'课程基本信息',name);rel('DOC01',id,'载明基本信息');ce=e.add(pg,quote,id,name)
  info.append({'条目ID':id,'字段':name,'内容':value or compact(quote),'内容性质':'原文摘录/空白规范化，不是岗位需求','字段完整性':'已定位',**e.fields([ce])})
 # 教学理念跨页全文与作用全文分别追加来源，不将其降格为项目活动。
 for id,start,end,pgs in [('CB004','（二）课程作用','（三）前导、后续课程',[1]),('CB007','（一）课程教学理念','（二）课程设计思路',[1,2])]:
  ids=[];parts=[]
  for pg in pgs:
   raw=e.raw[pg-1];a=raw.find(start);z=raw.find(end)
   part=raw[a+len(start):] if a>=0 else raw
   if z>=0:part=raw[(a+len(start) if a>=0 else 0):z]
   if part.strip():ids.append(e.add(pg,part,id,'课程定位/理念'));parts.append(part)
  row=next(r for r in info if r['条目ID']==id);row['内容']='\n'.join(parts);row.update(e.fields(ids))
 # 设计思路5条，知识4、能力6、素质4，逐项保留，不复制旧目标库。
 objectives=[]
 def goal(id,name,pg,quote):
  node(id,'课程总体目标/设计原则',name);rel('DOC01',id,'载明总体目标或设计原则');ce=e.add(pg,quote,id,name)
  objectives.append({'条目ID':id,'目标类别':name,'层次':'课程总体，不等同项目实训或考核','内容':e.quote(pg,quote),**e.fields([ce])})
 raw=e.raw[1];design=raw[raw.index('1.以应用场景'):raw.index('0、课程')]
 for i,part in enumerate(re.split(r'(?=\d\.)',design),0):
  if part.strip():goal(f'CD{len([g for g in objectives if g["条目ID"].startswith("CD")])+1:03}','课程设计思路',2,part)
 # 分页目标片段按编号而非关键词猜测。
 for kind,prefix,segments in [('知识目标','GK',[(2,'（一）知识目标','（二）能力目标')]),('能力目标','GA',[(2,'（二）能力目标',None),(3,None,'（三）素质目标')]),('素质目标','GQ',[(3,'（三）素质目标','3、课程内容')])]:
  count=0
  for pg,start,end in segments:
   raw=e.raw[pg-1];part=raw[raw.index(start)+len(start):] if start else raw
   if end:part=part[:part.index(end)]
   for m in re.finditer(r'\d\.[\s\S]*?(?=\n\d\.|\Z)',part):
    count+=1;goal(f'{prefix}{count:03}',kind,pg,m.group())
 goal('CI001','课程设计思路总述',2,design)
 # 项目从可见正文标题独立解析；P02/P05原文用实训，统一列实践并另存原词。
 project_headers=[]
 for pg,text in enumerate(e.raw,1):
  for m in re.finditer(r'项目([一二三四五六七八])\s+([^\n]+?)[(（]共\s*(\d+)\s*学时，理论\s*(\d+)\s*学时，(实践|实训)\s*(\d+)\s*学时[)）]',text):project_headers.append((pg,m.start(),m.end(),m))
 if len(project_headers)!=8:raise ValueError('实际项目标题不能完整解析为8项')
 projects=[];project_blocks={}
 for i,(pg,start,end,m) in enumerate(project_headers):
  id=f'P{i+1:02}';node(id,'教学项目',compact(m.group(2)));rel('DOC01',id,'包含教学项目')
  lastpg=project_headers[i+1][0] if i<7 else 13
  pages=[]
  for n in range(pg,lastpg+1):
   text=e.raw[n-1];a=start if n==pg else 0;z=project_headers[i+1][1] if i<7 and n==lastpg else len(text)
   if i==7 and n==13:z=text.index('6"8')
   if z>a and text[a:z].strip():pages.append((n,text[a:z]))
  full='\n'.join(text for n,text in pages);project_blocks[id]=pages
  ces=[e.add(n,text,id,'教学项目正文（含内容/方式/要求/重难点/思政）') for n,text in pages]
  def section(a,z):
   s=full.find(a)
   if s<0:return '未可靠定位，待审查'
   tail=full[s+len(a):];j=tail.find(z) if z else -1;return tail[:j].strip() if j>=0 else tail.strip()
  projects.append({'项目ID':id,'项目原文名称':compact(m.group(2)),'总学时':int(m.group(3)),'理论学时':int(m.group(4)),'实践学时':int(m.group(6)),'原文学时术语':m.group(5),'教学内容':section('1.教学内容','2.教学方式与要求'),'教学方式':section('（1）教学方式','（2）教学要求'),'教学要求':section('（2）教学要求','3.教学重难点'),'教学重点':section('[教学重点]','[教学难点]'),'教学难点':section('[教学难点]','4.课程思政'),'课程思政':section('4.课程思政',None),'字段规范化说明':'项目名称仅去除排版空白，保留原文用词；实践列兼容原文“实训”',**e.fields(ces)})
 knowledge=[];occurrences=[];byname={}
 for r in manifest['knowledge_occurrences']:
  name=r['name'];id=r['knowledge_id']
  if name not in byname:
   node(id,'规范化知识技能点',name);byname[name]={'知识点ID':id,'规范化名称':name,'课程功能标签':r['tag'],'功能标签依据':'按技术对象组织，非岗位证据等级或重要性排序','多标签理由':'同一技术作用于定义/操作与查询，或兼有概念/管理功能' if ';' in r['tag'] else '单一主要技术功能','原文说明':e.quote(r['page'],r['quote']),'内容性质':'原文明示内容；名称及功能标签为研究规范化','证据ID集合':[],'对应项目集合':[],'页码集合':[]}
  ces=[e.add(r['page'],r['quote'],id,r['project']+'知识技能点')];row=byname[name];row['证据ID集合']+=ces;row['对应项目集合'].append(r['project']);row['页码集合'].append(r['page'])
  oid=r['occurrence_id'];node(oid,'知识点项目出现记录',r['project']+' '+name);rel(r['project'],oid,'出现知识技能点');rel(oid,id,'规范化为同一知识点')
  for ce in ces:e.links.append({'实体ID':oid,'证据ID':ce,'所属章节或项目':r['project'],'使用性质':'原文出现记录','说明':'保留跨项目出现，不复制规范化实体','研究者审查状态':'待审查'})
  occurrences.append({'出现记录ID':oid,'知识点ID':id,'项目ID':r['project'],'规范化名称':name,'项目语境原文':e.quote(r['page'],r['quote']),'知识点去重说明':'同一知识点跨项目共用CK，出现记录分别保留',**e.fields(ces)})
 for row in byname.values():
  knowledge.append({'知识点ID':row['知识点ID'],'规范化名称':row['规范化名称'],'原文说明':row['原文说明'],'课程功能标签':row['课程功能标签'],'功能标签依据':row['功能标签依据'],'多标签理由':row['多标签理由'],'对应项目ID':';'.join(dict.fromkeys(row['对应项目集合'])),'内容性质':row['内容性质'],**e.fields(row['证据ID集合'])})
 other=[]
 for i,r in enumerate(manifest['other_level_knowledge'],1):
  id=f'CL{i:03}';node(id,'非项目技术实训层内容',r['name']);rel('DOC01',id,'目标或思政提及');ce=e.add(r['page'],r['quote'],id,r['level'],note='目标或思政提及，不自动认定独立技术实训')
  other.append({'条目ID':id,'名称':r['name'],'所属层次':r['level'],'允许解释':'课程文本有提及','禁止解释':'不等于项目安排独立权限/DCL/加密审计实训',**e.fields([ce])})
 activities=[]
 for r in manifest['activities']:
  id=r['id'];node(id,'教学活动',r['name']);rel(r['project'],id,'安排教学活动');ce=e.add(r['page'],r['quote'],id,r['project'])
  activities.append({'教学活动ID':id,'项目ID':r['project'],'活动主体':r['actor'],'教学方式/活动类别':r['kind'],'技术活动或教学方式名称':r['name'],'明确独立操作':r['independent'],'活动性质':'教师教学方式' if r['actor']=='教师' else '学生应完成活动，非实际完成记录','是否正式考核':'否，原文未据此规定正式评分','正式评分规则':'未提供',**e.fields([ce])})
 outcomes=[];activity_outcome=[]
 for r in manifest['outcomes']:
  id=r['id'];node(id,'学习成果/潜在产物',r['name']);ce=e.add(r['page'],r['quote'],id,r['project'],r['nature'],note='成果名称为技术结果归纳；潜在产物不是原文提交要求' if '推断' in r['nature'] else '')
  ces=[ce]+[e.add(x['page'],x['quote'],id,r['project'],r['nature'],note='原句跨第4—5页，分别保留摘录，不能默补熟字') for x in r.get('additional_evidence',[])]
  outcomes.append({'学习成果ID':id,'项目ID':r['project'],'学习成果或产物':r['name'],'成果证据性质':r['nature'],'正式提交物要求':r['submission'],'是否学生已完成':'无法判断，不采集学生表现','是否属于研究推断':'是' if '推断' in r['nature'] else '名称为归纳，活动或成果原文明示',**e.fields(ces)})
  for ca in r['activity_ids']:
   activity_outcome.append({'教学活动ID':ca,'学习成果ID':id,'关联性质':'研究组织关联；须区分明示技术结果与推断潜在产物','审查状态':'待审查'});rel(ca,id,'可对应成果或潜在产物','研究组织关联，非已完成事实')
 checks=[];activity_checks=[]
 for r in manifest['checks']:
  id=r['id'];node(id,'评价/检查/完成要求证据',r['kind']);ce=e.add(r['page'],r['quote'],id,r['project'] or '课程考核章节')
  checks.append({'评价检查ID':id,'项目ID':r['project'] or '课程总体','类别':r['kind'],'证据边界':r['limit'],'正式评分细则':'未提供或无法确认','成绩比例':'未提供或无法确认','学生成绩':'未采集，不适用','是否构成正式量化评价':'否',**e.fields([ce])})
  for ca in r['activity_ids']:
   activity_checks.append({'教学活动ID':ca,'评价检查ID':id,'关联性质':'同项目相关活动的研究组织关联，非评分绑定','审查状态':'待审查'});rel(ca,id,'相关检查或完成要求','研究组织关联，非量规')
  if not r['activity_ids']:rel('DOC01',id,'考核章节标题候选','仅隐藏文本层，无法确认具体规则')
 # 末页独立实训是文本层候选；学时数值和同源规则还拥有本轮研究者确认来源。
 practice=[];practice_operations=[];practice_project=[]
 specs=[('PT01','实训一：数据定义',10,'实训二：数据操作',['使用SQL 语言，创建数据库；','使用CREATE TABLE 语句，创建多张表；','使用ALTER TABLE 语句，修改表结构；'],['P03']),('PT02','实训二：数据操作',12,'实训三：数据查询',['使用INERT INTO 语句，插入样本数据；','使用UPDATE 语句，修改数据；','使用DELETE FROM 语句，删除数据；'],['P04']),('PT03','实训三：数据查询',10,'总计学时：',['使用SELECT 语句对数据库完成相应查询，包括：列查询、条件查询、多表连接和分类查询与汇总。'],['P05'])]
 raw=e.raw[12]
 for id,name,hours,end,ops,pids in specs:
  node(id,'独立实训文本层候选',name);rel('DOC01',id,'文本层列出独立实训','隐藏文本层候选＋研究者确认学时口径')
  part=raw[raw.index(name):raw.index(end)]
  text_hours=re.findall(r'(?<![A-Za-z0-9])(10|12)(?![A-Za-z0-9])',part)
  if len(text_hours)!=1 or int(text_hours[0])!=hours:raise ValueError('实训文本层学时与研究者确认值不一致：'+id)
  ce=e.add(13,part,id,'末尾实训表',note='页面渲染不可见；任务细节不能作为已确认正向实训证据')
  practice.append({'实训ID':id,'实训原文名称':name,'学时':hours,'文本层提取学时':int(text_hours[0]),'研究者确认学时':hours,'实训内容原文':part,'学时来源':'PDF文本层数值＋本轮研究者明确确认；页面尚未可视确认','内容来源状态':'仅隐藏文本层候选，待研究者确认可读版本','研究者确认来源':'2026-10-24本轮任务指令（四）3实训名称与10/12/10学时；（三）同批实践32学时','疑似拼写问题':'INERT INTO（原文保留）' if id=='PT02' else '无明显拼写异常','规范化建议（非原文）':'疑似应为INSERT INTO，待研究者确认' if id=='PT02' else '无','学时累计组':'PRACTICE32同源组，不与项目实践再累计',**e.fields([ce])})
  for j,op in enumerate(ops,1):
   opid=id+'-OP'+str(j);node(opid,'实训操作文本层候选',compact(op));ce=e.add(13,op,opid,'末尾实训表',note='隐藏文本，尚未可视确认；不升级为已确认正向证据');rel(id,opid,'文本层列出操作','隐藏文本层候选')
   practice_operations.append({'操作ID':opid,'实训ID':id,'原文技术操作':e.quote(13,op),'规范化建议':'INERT INTO疑似INSERT INTO' if 'INERT' in op else '不改原文','内容状态':'仅文本层候选',**e.fields([ce])})
  for pid in pids:
   practice_project.append({'实训ID':id,'项目ID':pid,'关联理由':'技术对象与活动内容相似，研究组织线索','关联状态':'内容相关、学时对应未确认','精确学时分配':'未确认，不拆分或重分配','原文是否明确逐项对应':'否，PDF未见逐项对应','审查状态':'待研究者审查'});rel(id,pid,'内容相似线索','研究推断，非精确学时对应')
 hours=[{'视角':'8项目','总学时':sum(p['总学时'] for p in projects),'理论学时':sum(p['理论学时'] for p in projects),'实践学时':sum(p['实践学时'] for p in projects),'来源':'PDF可见项目标题','口径规则':HOUR_RULE},{'视角':'3独立实训','总学时':None,'理论学时':None,'实践学时':sum(p['学时'] for p in practice),'来源':'PDF隐藏文本＋本轮研究者明确确认','口径规则':HOUR_RULE},{'视角':'课程去重总量','总学时':64,'理论学时':32,'实践学时':32,'来源':'PDF8项目合计＋研究者同源确认','口径规则':HOUR_RULE}]
 questions=[('PDF第13页','文本层实训与考核标题在页面渲染中不可见，不能作为已确认正向实践/考核证据','保留文本层与研究者确认学时分层；请求可读版本或研究者确认具体任务原文'),('PT01—PT03与P01—P08','精确逐项学时对应未明示','仅内容相关线索，不分摊学时'),('PT02','INERT INTO疑似拼写错误','原文保留，规范建议INSERT INTO另列'),('课程考核','末尾只有不可见标题，未提供可读评分细则或比例','不编造评分或成绩'),('目标与项目','权限管理/DCL在总体目标，P08技术实践正文主要备份恢复及事务','不由目标或思政提及认定已安排独立权限配置实训'),('P03','替代键/唯一键术语对照需审查','保留两种原文措辞，不擅自认为术语完全同义'),('P05','分组聚会疑似分组聚合笔误','原文保留，不作为新知识点'),('P07','索引分类与作疑似缺字','原文保留，规范化名称依上下文标注'),('P04','DELETE可回滚与TRUNCATE不可逆表述缺少存储引擎/事务条件','原文保留；教学解释条件待审，不添加教材论断作课程原文'),('学习成果','多数仅要求完成操作，未规定交付文件形式','潜在SQL脚本/备份文件/对比记录均标研究推断，非提交要求'),('知识规范化','CREATE TABLE在P03库表与P07建表索引语境、WHERE在操作和查询语境出现','共用规范化实体，出现记录分别保留；审查边界是否适当'),('课程定位/页眉','部分标题编码异常，P02教室需求疑似教师需求','保留提取原文，按可见章节及上下文组织，不静默替换')]
 issues=[{'问题ID':f'CQ{i:03}','对象':obj,'问题':q,'当前处理与所需审查':handling,'审查状态':'待研究者审查'} for i,(obj,q,handling) in enumerate(questions,1)]
 output=Path(output or HERE/'output').resolve()
 if output.is_relative_to(ROOT) and not output.is_relative_to(HERE):raise ValueError('禁止覆盖历史成果目录')
 output.mkdir(parents=True,exist_ok=True)
 write_book(output/'01_课程标准基本信息与目标.xlsx',{'课程基本信息':info,'课程目标与设计原则':objectives})
 write_book(output/'02_八个教学项目结构化表.xlsx',{'教学项目':projects})
 write_book(output/'03_课程知识技能点规范化表.xlsx',{'规范化知识技能点':knowledge,'项目出现记录':occurrences,'目标及思政单列提及':other})
 write_book(output/'04_教学活动与学习成果表.xlsx',{'教学活动':activities,'学习成果与潜在产物':outcomes,'活动成果多对多关联':activity_outcome})
 write_book(output/'05_三项独立实训任务表.xlsx',{'独立实训任务':practice,'实训技术操作':practice_operations,'项目内容相关线索':practice_project})
 write_book(output/'06_课程评价与检查证据表.xlsx',{'检查反馈与完成要求':checks,'活动检查多对多关联':activity_checks})
 records=list(e.records.values());evidence_links=[dict(r,关联ID=f'EL{i:04}') for i,r in enumerate(e.links,1)]
 write_book(output/'07_课程教学要素证据索引.xlsx',{'课程文档':[{'文档ID':'DOC01','文件名':pdf.name,'实际路径':str(pdf.relative_to(ROOT)),'页数':len(doc),'GitBlobSHA':gitsha(pdf),'完整原件':'已读取完整13页，末页文本可见性存在异常'}],'原文证据':records,'实体证据多对多关联':evidence_links,'教学要素实体':entities,'教学要素关系':relations,'数据字典':dictionary()})
 write_book(output/'08_课程学时双视角核对.xlsx',{'双视角学时核对':hours,'八项目学时分配':[{k:p[k] for k in ['项目ID','项目原文名称','总学时','理论学时','实践学时','证据ID','PDF实际页码']} for p in projects],'三实训学时分配':practice,'对应关系限制':practice_project})
 write_book(output/'10_课程标准人工审查问题清单.xlsx',{'人工审查问题':issues})
 summary={'输入PDF':str(pdf.relative_to(ROOT)),'页数':13,'PDFGitSHA':gitsha(pdf),'main提交':snapshot['main_commit'],'项目':len(projects),'知识技能点':len(knowledge),'项目出现记录':len(occurrences),'目标思政单列提及':len(other),'教学活动':len(activities),'教师活动':sum(a['活动主体']=='教师' for a in activities),'学生活动':sum(a['活动主体']=='学生' for a in activities),'学习成果或潜在产物':len(outcomes),'其中推断潜在产物':sum('推断' in r['成果证据性质'] for r in outcomes),'评价检查证据':len(checks),'独立实训':len(practice),'实训操作':len(practice_operations),'原文证据':len(records),'原文证据页码覆盖':sorted({r['PDF实际页码'] for r in records}),'隐藏文本证据':sum(r['来源证据层']=='隐藏文本候选' for r in records),'实体':len(entities),'关系':len(relations),'学时':hours,'人工问题':len(issues),'确认的能力课程支持关系':0,'学生评价结果':0}
 (output/'课程教学要素数据库.json').write_text(json.dumps({'summary':summary,'document':{'id':'DOC01','path':str(pdf.relative_to(ROOT)),'pages':13},'basic_info':info,'objectives':objectives,'projects':projects,'knowledge':knowledge,'occurrences':occurrences,'other_level_mentions':other,'activities':activities,'outcomes':outcomes,'checks':checks,'practice':practice,'practice_operations':practice_operations,'practice_project':practice_project,'activity_outcome':activity_outcome,'activity_check':activity_checks,'evidence':records,'evidence_links':evidence_links,'entities':entities,'relations':relations,'issues':issues},ensure_ascii=False,indent=2)+'\n')
 (output/'processing_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
 # 全13页原生文本留档，说明可见性与文字层不等同，不作为第二个原始资料来源。
 (output/'PDF逐页文本_含隐藏文本警示.txt').write_text('注意：第13页末尾实训/考核为不可见文字层，不能视作已确认可见正文。\n'+ '\n'.join(f'=== PDF实际第{i+1}页 ===\n{s}' for i,s in enumerate(e.raw)))
 shutil.copyfile(HERE/'研究方法.md',output/'09_第四阶段第一轮研究方法说明.md')
 (output/'11_第四阶段第一轮执行总结.md').write_text(report(summary))
 doc.close();return json.loads((output/'课程教学要素数据库.json').read_text())

def dictionary():
 items=[('CKxxx','规范化知识点实体；固定ID来自人工清单，不因多项目出现而复制'),('KOxxx','知识点在具体项目的出现记录，独立于CK'),('CAxxx/COxxx/CExx-hash','活动/成果/证据；证据号为页码＋摘录哈希，页码相同但原文不同可区分'),('来源证据层','可见PDF原文与隐藏文本候选分开；课程目标不等于项目实践'),('成果证据性质','明确产物/可直接确认技术结果/研究推断潜在产物分开，均非学生实际表现'),('课程功能标签','课程内容组织标签，非岗位A/B/C等级或重要性排名'),('PRACTICE32同源组',HOUR_RULE),('关联状态','内容相关不等于精确学时对应；研究组织关联不等于原文明示评分绑定'),('评分规则','没有可读原文则标未提供或无法确认，不凭教学常识补写')]
 return [{'字段或编号':k,'定义及使用规则':v} for k,v in items]

def report(s):
 validation=HERE/'验证记录.json'
 result=json.loads(validation.read_text()) if validation.exists() else {}
 verified=result.get('build_git_blob')==gitsha(Path(__file__)) and result.get('manifest_git_blob')==gitsha(HERE/'extraction_manifest.json')
 test_status=f"实际运行{result['tests']}项：通过{result['passed']}，失败{result['failures']}，错误{result['errors']}；命令和完整输出见验证结果.txt。" if verified else '当前版本验证记录未登记，以实际执行日志为准。'
 return f'''# 第四阶段第一轮执行总结

基线main：{s['main提交']}；PR #7已合并。本轮不修改任何第三阶段能力、关联或网页。完整读取原始{s['输入PDF']}，13页，Git blob {s['PDFGitSHA']}，与main原件相同，没有复制/修改PDF。

## 提取数量及统计口径

8个教学项目；{s['知识技能点']}个规范化知识技能点，{s['项目出现记录']}条项目出现记录，另将{s['目标思政单列提及']}条目标/思政层提及独立存放，不等同项目技术实训。知识点是按技术对象和操作归纳的候选组织实体，数量由原文与规范化规则决定，不以历史37条为目标。

{s['教学活动']}项教学活动（教师{s['教师活动']}、学生{s['学生活动']}），为合并同一主体、目的和原文段落后的活动记录，不是每个教学动作词的计数。{s['学习成果或潜在产物']}项成果/潜在产物，其中{s['其中推断潜在产物']}项为研究推断；正式提交形式多数未规定。{s['评价检查证据']}条检查/反馈/作业/实验完成要求或考核标题证据，不是17项评分任务；无可确认正式评分规则、成绩比例或学生成绩。

3项独立实训、{s['实训操作']}条操作，来源是末页隐藏文字层，尚未可视确认。研究者本轮确认三项名称与10/12/10学时；任务具体技术内容仍不能作为已确认正向实训证据。

{s['原文证据']}条去重原文摘录，证据覆盖PDF实际第1—13页；其中{s['隐藏文本证据']}条隐藏文本候选。每条具有来源路径、页码、原文、可信状态及实体关联。原文摘录保持原字符和换行；只允许匹配定位时去排版空白。

## 学时核对

8项目64=理论32+实践32；3实训10+12+10=32。{HOUR_RULE}同源属于研究者明确确认，不冒称PDF已逐项说明对应关系。项目—实训只保留技术内容相似线索，学时对应未确认。

## 异常与研究边界

全部13页已渲染浏览；第13页专项放大核对：实训表及考核标题在文本层可提取但渲染不可见。保留隐藏文本原文，不能称完整可见实训依据。INERT INTO原样保存，INSERT INTO仅为待审规范化建议。另记录分组聚会、索引分类与作、教室需求等疑似字词问题；DELETE/TRUNCATE回滚描述的条件需审查。权限/DCL主要见目标、权限加密审计见思政，不能升级为独立实操。多数操作要求未规定提交SQL脚本、备份文件等，推断产物单列。

## 交付、验证与下一轮

output/01—11为11项要求成果（9Excel＋2Markdown）；附数据库JSON、计数JSON及逐页文本。上级提供build.py、extraction_manifest.json、README、依赖、输入快照、测试与实际验证日志。知识、项目出现、活动、成果、检查和证据均独立建表，可多对多关联。后续可引用CK/KO/CA/CO/PT/CE编号独立设计双向映射，当前支持矩阵记录0，未开展第五阶段。

测试命令：`python -m unittest discover -s stage4_course_extraction -p test_extraction.py -v`。{test_status}

本次仅完成第四阶段第一轮课程标准结构化提取，不代表岗位能力—课程内容映射已经完成，不代表学生能力达成度评价已经完成。
'''

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path);args=parser.parse_args();print(json.dumps(build(args.output)['summary'],ensure_ascii=False,indent=2))
