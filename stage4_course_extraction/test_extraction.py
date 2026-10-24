"""以原PDF实际摘录、数据边界和旧文件哈希为约束的回归测试。"""
import unittest,tempfile,json,re
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
import fitz
import build as b

class ExtractionTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory();cls.out=Path(cls.tmp.name)/'one';cls.d=b.build(cls.out)
 @classmethod
 def tearDownClass(cls):cls.tmp.cleanup()
 def test_01_original_pdf_complete(self):
  p=b.locate_pdf();self.assertEqual(b.gitsha(p),b.PDF_SHA)
  with fitz.open(p) as d:self.assertEqual(len(d),13);self.assertTrue(all(len(p.get_text())>0 for p in d))
 def test_02_eight_projects_unique_actual_titles(self):
  ps=self.d['projects'];self.assertEqual(len(ps),8);self.assertEqual([p['项目ID'] for p in ps],[f'P{i:02}' for i in range(1,9)])
  self.assertEqual([p['项目原文名称'] for p in ps],['数据库认知与基础管理','关系数据库设计与规范化','数据库表的创建与管理','数据库操作','数据库查询','数据库视图','数据库索引','数据库安全'])
  for p in ps:
   for k in ['教学内容','教学方式','教学要求','教学重点','教学难点','证据ID']:self.assertTrue(p[k]);self.assertNotIn('未可靠定位',p[k])
 def test_03_project_hours_from_pdf(self):
  ps=self.d['projects'];self.assertEqual([p['总学时'] for p in ps],[4,8,12,12,16,4,4,4])
  self.assertEqual(sum(p['总学时'] for p in ps),64);self.assertEqual(sum(p['理论学时'] for p in ps),32);self.assertEqual(sum(p['实践学时'] for p in ps),32)
  self.assertTrue(all(p['总学时']==p['理论学时']+p['实践学时'] for p in ps))
 def test_04_three_practice_hours_have_separate_sources(self):
  ps=self.d['practice'];self.assertEqual(len(ps),3);self.assertEqual([p['学时'] for p in ps],[10,12,10]);self.assertEqual(sum(p['学时'] for p in ps),32)
  for p in ps:
   self.assertEqual(p['文本层提取学时'],p['研究者确认学时']);self.assertIn('研究者',p['学时来源']);self.assertIn('不可见',p['来源状态'])
 def test_05_no_double_count(self):
  rows=self.d['summary']['学时'];course=next(r for r in rows if r['视角']=='课程去重总量');self.assertEqual(course['总学时'],64);self.assertEqual(course['实践学时'],32)
  self.assertIsNone(next(r for r in rows if r['视角']=='3独立实训')['总学时'])
  self.assertTrue(all('不重复' in r['口径规则'] for r in rows))
 def test_06_knowledge_ids_and_occurrence_separation(self):
  ks={r['知识点ID']:r for r in self.d['knowledge']};os=self.d['occurrences'];self.assertEqual(len(ks),len(self.d['knowledge']))
  self.assertTrue(all(re.fullmatch(r'CK\d{3}',k) for k in ks));self.assertTrue(all(re.fullmatch(r'KO\d{3}',r['出现记录ID']) for r in os))
  self.assertEqual(len({r['出现记录ID'] for r in os}),len(os));self.assertTrue(all(r['知识点ID'] in ks for r in os))
  create=next(r['知识点ID'] for r in ks.values() if r['规范化名称']=='CREATE TABLE创建数据表');self.assertEqual({r['项目ID'] for r in os if r['知识点ID']==create},{'P03'})
 def test_07_project_and_entity_endpoints(self):
  ps={p['项目ID'] for p in self.d['projects']};ids={r['实体ID'] for r in self.d['entities']};self.assertEqual(len(ids),len(self.d['entities']))
  for table in ['occurrences','activities','outcomes','practice_project']:
   for r in self.d[table]:self.assertIn(r['项目ID'],ps)
  for r in self.d['relations']:self.assertIn(r['源实体ID'],ids);self.assertIn(r['目标实体ID'],ids)
 def test_08_evidence_verbatim_and_page_traceable(self):
  ids=set()
  with fitz.open(b.locate_pdf()) as pdf:
   for r in self.d['evidence']:
    self.assertIn(r['PDF实际页码'],range(1,14));self.assertNotIn(r['证据ID'],ids);ids.add(r['证据ID'])
    self.assertIn(r['原文短摘录'],pdf[r['PDF实际页码']-1].get_text());self.assertEqual(r['来源文件名'],'课标0913.pdf')
  self.assertEqual({r['PDF实际页码'] for r in self.d['evidence']},set(range(1,14)))
 def test_09_entity_evidence_associations(self):
  ids={r['实体ID'] for r in self.d['entities']};es={r['证据ID'] for r in self.d['evidence']}
  for r in self.d['evidence_links']:self.assertIn(r['实体ID'],ids);self.assertIn(r['证据ID'],es)
  linked={r['实体ID'] for r in self.d['evidence_links']}
  self.assertTrue((ids-{'DOC01'})<=linked)
 def test_10_teacher_demo_not_student_independent(self):
  acts=self.d['activities']
  for r in acts:
   if r['活动主体']=='教师':self.assertEqual(r['明确独立操作'],'不适用（教师活动）')
   if r['明确独立操作']=='是，原文明示':self.assertIn('独立',r['原文摘录'])
  self.assertTrue(any('同步' in r['明确独立操作'] for r in acts if r['活动主体']=='学生'))
 def test_11_practice_not_formal_scoring(self):
  for r in self.d['activities']:self.assertEqual(r['是否正式考核'],'否，原文未据此规定正式评分')
  for r in self.d['checks']:
   self.assertEqual(r['是否构成正式量化评价'],'否');self.assertEqual(r['成绩比例'],'未提供或无法确认');self.assertEqual(r['学生成绩'],'未采集，不适用')
 def test_12_outcomes_inference_separate(self):
  rs=self.d['outcomes'];self.assertTrue(any('研究推断' in r['成果证据性质'] for r in rs))
  for r in rs:
   self.assertEqual(r['是否学生已完成'],'无法判断，不采集学生表现')
   if '推断' in r['成果证据性质']:self.assertIn(r['是否属于研究推断'],'是');self.assertNotEqual(r['正式提交物要求'],'已明确正式提交')
  backup=next(r for r in rs if '备份文件' in r['学习成果或产物']);self.assertIn('推断',backup['成果证据性质'])
 def test_13_multimodal_hidden_evidence_not_confirmed(self):
  es={r['证据ID']:r for r in self.d['evidence']}
  for r in self.d['practice']+self.d['practice_operations']:
   for id in r['证据ID'].split(';'):self.assertEqual(es[id]['来源证据层'],'隐藏文本候选');self.assertIn('尚未可视确认',es[id]['可信状态'])
  final=self.d['checks'][-1];self.assertIn('不可见',final['来源状态']);self.assertIn('无法确认',final['正式评分细则'])
 def test_14_typo_preserved(self):
  pt=self.d['practice'][1];self.assertIn('INERT INTO',pt['原文摘录']);self.assertIn('INSERT INTO',pt['规范化建议（非原文）'])
 def test_15_goals_vs_project_practice(self):
  counts=Counter(r['目标类别'] for r in self.d['objectives']);self.assertEqual(counts['知识目标'],4);self.assertEqual(counts['能力目标'],6);self.assertEqual(counts['素质目标'],4)
  for r in self.d['other_level_mentions']:self.assertIn('不等于',r['禁止解释'])
  self.assertFalse(any('权限配置' in r['技术活动或教学方式名称'] for r in self.d['activities']))
 def test_16_no_unconfirmed_hour_mapping(self):
  for r in self.d['practice_project']:self.assertEqual(r['关联状态'],'内容相关、学时对应未确认');self.assertIn('未确认',r['精确学时分配'])
 def test_17_old_138_files_including_stage3_and_app_unchanged(self):
  snap=b.protect();self.assertEqual(len(snap['repository_files']),138)
 def test_18_no_ability_support_matrix_or_student_results(self):
  self.assertEqual(self.d['summary']['确认的能力课程支持关系'],0);self.assertEqual(self.d['summary']['学生评价结果'],0)
  self.assertFalse(any(re.fullmatch(r'DB\d{3}',r['实体ID']) for r in self.d['entities']))
  self.assertFalse(any('能力支持' in r['关系类型'] for r in self.d['relations']))
 def test_19_crosspage_evidence_not_silently_joined(self):
  row=next(r for r in self.d['outcomes'] if r['学习成果ID']=='CO003');self.assertEqual(row['PDF实际页码'],'4;5')
 def test_20_multi_to_multi_activity_outputs_checks(self):
  acts={a['教学活动ID'] for a in self.d['activities']};os={r['学习成果ID'] for r in self.d['outcomes']};ecs={r['评价检查ID'] for r in self.d['checks']}
  for r in self.d['activity_outcome']:self.assertIn(r['教学活动ID'],acts);self.assertIn(r['学习成果ID'],os)
  for r in self.d['activity_check']:self.assertIn(r['教学活动ID'],acts);self.assertIn(r['评价检查ID'],ecs)
 def test_21_method_has_complement_and_hours(self):
  text=(self.out/'09_第四阶段第一轮研究方法说明.md').read_text();self.assertIn(b.COMPLEMENT,text);self.assertIn(b.HOUR_RULE,text);self.assertIn('不自动判定为多余',text);self.assertIn('不自动判定为课程教学缺陷',text)
 def test_22_workbooks_open_and_repeatable(self):
  second=Path(self.tmp.name)/'two';b.build(second);files=list(self.out.glob('*.xlsx'));self.assertEqual(len(files),9)
  for p in files:
   a=load_workbook(p,read_only=True);z=load_workbook(second/p.name,read_only=True)
   try:
    self.assertEqual(a.sheetnames,z.sheetnames)
    for name in a.sheetnames:self.assertEqual(list(a[name].values),list(z[name].values))
   finally:a.close();z.close()
  self.assertEqual(json.loads((second/'课程教学要素数据库.json').read_text()),self.d)
 def test_23_invalid_evidence_rejected(self):
  with fitz.open(b.locate_pdf()) as doc:
   e=b.Evidence(doc,b.locate_pdf())
   with self.assertRaises(ValueError):e.add(1,'不存在的评分比例80%', 'bad','bad')

 def test_24_review_interfaces_and_contexts(self):
  d=self.d
  self.assertEqual(len(d['knowledge']),60)
  self.assertEqual(len(d['occurrences']),61)
  ko={r['出现记录ID']:r for r in d['occurrences']}
  self.assertEqual(ko['KO056']['知识点ID'],'CK060')
  self.assertIn('操作',ko['KO027']['应用语境标签'])
  self.assertIn('SELECT',ko['KO034']['应用语境标签'])
  self.assertEqual({r['技术行为标签'] for r in d['technical_behavior_interfaces']},{'备份相关技术行为','恢复相关技术行为'})
  self.assertTrue(all(r['支持等级']=='本轮不判定' for r in d['technical_behavior_interfaces']))
  self.assertEqual({r['学习成果ID'] for r in d['outcomes'] if '推断' in r['成果证据性质']},{'CO005','CO007','CO009','CO012','CO014'})
  self.assertEqual(len(d['issues']),12)
  self.assertTrue(all(r['是否仍需研究者确认']=='是' for r in d['issues']))
  w=load_workbook(self.out/'03_课程知识技能点规范化表.xlsx',read_only=True)
  self.assertEqual(sum(1 for _ in w['规范化知识技能点'].values)-1,len(d['knowledge']));w.close()

if __name__=='__main__':unittest.main()
