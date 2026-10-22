import contextlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from scripts.phase1 import (SOURCE_NAME, DUTY, REQ, REQUIRED, canonical_url,
                            duplicates, read_source, requirement_evidence, run, write_book)

ROOT = Path(__file__).resolve().parents[1]


class Phase1Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.temp.name)
        with contextlib.redirect_stdout(io.StringIO()):
            cls.report = run(ROOT / SOURCE_NAME, cls.out)
        cls.jobs = read_source(ROOT / SOURCE_NAME)[2]

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def rows(self, filename, sheet):
        wb = load_workbook(self.out / filename)
        values = list(wb[sheet].values)
        wb.close()
        return [dict(zip(values[0], r)) for r in values[1:]]

    def test_actual_input_and_missingness(self):
        self.assertEqual(self.report['主库岗位数'], 50)
        self.assertEqual(self.report['主库字段缺失']['证据摘录'], 50)
        self.assertEqual(self.report['主库字段缺失']['岗位发布日期'], 44)
        self.assertEqual(self.report['主库字段缺失']['城市/地区'], 3)
        self.assertEqual(len(self.report['工作表']), 4)

    def test_original_records_and_status_preserved(self):
        rows = self.rows('岗位数据质量检查.xlsx', '岗位原始与本轮状态')
        self.assertEqual(len(rows), len(self.jobs))
        for before, after in zip(self.jobs, rows):
            for field, value in before.items():
                self.assertEqual(after[field], value)
            self.assertEqual(after['本轮核验状态'], '待核验')

    def test_task_traceability_and_every_job_covered(self):
        tasks = self.rows('岗位任务编码表.xlsx', '任务编码')
        jobs = {j['样本ID']: j for j in self.jobs}
        self.assertEqual({t['岗位ID'] for t in tasks}, set(jobs))
        self.assertEqual(len({t['任务编码'] for t in tasks}), len(tasks))
        for t in tasks:
            self.assertIn(t['证据文本'], jobs[t['岗位ID']][DUTY])
            self.assertEqual(t['证据类型'], '摘要证据')
            self.assertEqual(t['核验状态'], '待核验')

    def test_relations_have_real_foreign_keys_and_evidence(self):
        tasks = {t['任务编码']: t for t in self.rows('岗位任务编码表.xlsx', '任务编码')}
        abilities = {a['三级能力编码'] for a in self.rows('三级职业能力指标体系.xlsx', '三级能力指标')}
        jobs = {j['样本ID']: j for j in self.jobs}
        for r in self.rows('岗位任务能力关系表.xlsx', '关系明细'):
            self.assertIn(r['三级能力编码'], abilities)
            self.assertEqual(r['岗位ID'], tasks[r['任务编码']]['岗位ID'])
            self.assertEqual(r['任务证据'], tasks[r['任务编码']]['证据文本'])
            self.assertEqual(r['任务到能力关系性质'], '研究推断')
            if not r['能力要求性质'].startswith('研究'):
                self.assertTrue(r['任职要求证据'])
                for e in r['任职要求证据'].split('；'):
                    self.assertIn(e, jobs[r['岗位ID']][REQ])

    def test_hierarchy_has_unique_parent_names(self):
        rows = self.rows('三级职业能力指标体系.xlsx', '三级能力指标')
        for codekey, namekey in [('一级编码','一级能力'), ('二级编码','二级能力')]:
            for code in {r[codekey] for r in rows}:
                self.assertEqual(len({r[namekey] for r in rows if r[codekey] == code}), 1)
        for r in rows:
            self.assertTrue(r['三级能力编码'].startswith(r['二级编码']))
            self.assertTrue(r['二级编码'].startswith(r['一级编码']))

    def test_duplicates_do_not_remove_records(self):
        dup = duplicates(self.jobs)
        strict = {r['岗位ID'] for r in dup if r['检查类型'] == '同单位同岗位'}
        self.assertTrue({'JOB004', 'JOB018', 'JOB023', 'JOB026', 'JOB005', 'JOB022'} <= strict)
        aliases = {r['岗位ID'] for r in dup if r['检查类型'] == '单位简称线索+同岗位'}
        self.assertEqual(aliases, {'JOB005','JOB022','JOB027'})
        self.assertEqual(len(self.jobs), 50)

    def test_url_identity_and_tracking(self):
        self.assertEqual(canonical_url('https://a.test/job?id=1&utm_source=x#t'), 'https://a.test/job?id=1')
        self.assertNotEqual(canonical_url('https://a.test/job?id=1'), canonical_url('https://a.test/job?id=2'))
        self.assertNotEqual(canonical_url('https://a.test/job?jobUnionId=1'), canonical_url('https://a.test/job?jobUnionId=2'))
        self.assertEqual(canonical_url(''), '')

    def test_negated_sql_is_not_explicit(self):
        for s in ['未明确SQL', '未单独写明SQL', '无需SQL', 'SQL待原文核验']:
            self.assertEqual(requirement_evidence(s, 'C131'), [])
        self.assertTrue(requirement_evidence('熟练SQL', 'C131'))
        self.assertEqual(requirement_evidence('SQL/Python优先', 'C131')[0][1], '优先要求（摘要明示）')

    def test_priority_is_local_to_clause(self):
        self.assertEqual(requirement_evidence('熟练SQL，Python/R优先', 'C131')[0][1], '明确要求（摘要明示）')

    def test_invalid_ids_fail_before_outputs(self):
        wb = Workbook()
        ws = wb.active
        ws.title = '招聘岗位主库'
        ws.append(REQUIRED)
        ws.append(['same'] + ['x'] * (len(REQUIRED)-1))
        ws.append(['same'] + ['x'] * (len(REQUIRED)-1))
        path = self.out / 'invalid.xlsx'
        wb.save(path)
        with self.assertRaisesRegex(ValueError, 'ID为空或重复'):
            read_source(path)

    def test_missing_schema_fails(self):
        wb = Workbook()
        wb.active.title = '招聘岗位主库'
        wb.active.append(['样本ID'])
        path = self.out / 'missing.xlsx'
        wb.save(path)
        with self.assertRaisesRegex(ValueError, '缺少必需字段'):
            read_source(path)

    def test_excel_evidence_is_text_not_formula(self):
        path = self.out / 'literal.xlsx'
        write_book(path, [('证据', [{'证据': '=1+1'}], None)])
        wb = load_workbook(path)
        self.assertEqual(wb.active['A2'].value, '=1+1')
        self.assertEqual(wb.active['A2'].data_type, 's')
        wb.close()

    def test_repeat_run_has_same_semantic_results(self):
        with contextlib.redirect_stdout(io.StringIO()):
            result = run(ROOT / SOURCE_NAME, self.out / 'rerun')
        self.assertEqual(result, self.report)
        for name in ['岗位数据质量检查.xlsx','岗位任务编码表.xlsx','三级职业能力指标体系.xlsx','岗位任务能力关系表.xlsx']:
            a, b = load_workbook(self.out/name), load_workbook(self.out/'rerun'/name)
            self.assertEqual(a.sheetnames, b.sheetnames)
            for s in a.sheetnames:
                self.assertEqual(list(a[s].values), list(b[s].values))
            a.close()
            b.close()


if __name__ == '__main__':
    unittest.main()
