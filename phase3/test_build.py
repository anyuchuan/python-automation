"""第三阶段的引用真实性、边界、图谱完整性与可复现性测试。"""
import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

import fitz
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('p3build',ROOT/'phase3/build.py')
b=importlib.util.module_from_spec(spec);spec.loader.exec_module(b)


class Phase3Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.out=Path(cls.temp.name)/'out'
        cls.stats=b.run(ROOT,cls.out)
        cls.snapshot=json.loads((ROOT/'phase3/input_snapshot.json').read_text())
        cls.before={e['path']:b.sha(ROOT/e['path']) for e in cls.snapshot['files']}
        cls.data=b.load_inputs(ROOT,cls.snapshot)
        cls.pdf,_=b.discover_pdf(ROOT)
        cls.corpus=b.Corpus(cls.pdf)
        cls.specs=json.loads((ROOT/'phase3/course_specs.json').read_text())
        cls.catalog=b.parse_course(cls.corpus,cls.specs)
        cls.plan=json.loads((ROOT/'phase3/matching_plan.json').read_text())
        cls.matrix,cls.summary=b.match_course(cls.data,cls.plan,cls.catalog,cls.corpus)
        cls.graph,cls.suggestions=b.build_graph(cls.data,cls.catalog,cls.matrix,cls.corpus)

    @classmethod
    def tearDownClass(cls):
        cls.corpus.doc.close();cls.temp.cleanup()

    def test_actual_inputs_and_pr2_merged(self):
        self.assertTrue(self.snapshot['pr2_merged'])
        self.assertEqual([self.stats[k] for k in ['岗位数','任务数','原关系数','三级候选能力数','疑似原关系数']],[50,79,188,34,109])
        self.assertEqual(self.stats['输入快照差异数'],0)
        with self.assertRaisesRegex(ValueError,'PR #2未确认合并'):
            b.load_inputs(ROOT,dict(self.snapshot,pr2_merged=False))

    def test_pdf_discovery_uses_content_not_name(self):
        root=Path(self.temp.name)/'renamed';root.mkdir()
        renamed=root/'unrelated_name.PDF';renamed.write_bytes(self.pdf.read_bytes())
        p,files=b.discover_pdf(root)
        self.assertEqual(p,renamed)
        self.assertEqual(len(files),1)

    def test_ambiguous_pdf_needs_review_not_guess(self):
        root=Path(self.temp.name)/'ambiguous';root.mkdir()
        for name in ['first.pdf','second.pdf']:
            (root/name).write_bytes(self.pdf.read_bytes())
        with self.assertRaisesRegex(ValueError,'不能唯一识别'):
            b.discover_pdf(root)

    def test_all_seven_workbooks_open_with_no_formulas(self):
        files=list(self.out.glob('*.xlsx'))
        self.assertEqual(len(files),7)
        for f in files:
            w=load_workbook(f,data_only=False)
            self.assertTrue(w.sheetnames)
            for s in w:
                self.assertEqual(s.freeze_panes,'A2')
                for row in s:
                    self.assertFalse(any(c.data_type=='f' for c in row))
            w.close()

    def test_all_prior_stage_files_and_pdf_preserved(self):
        self.assertEqual(self.before,{p:b.sha(ROOT/p) for p in self.before})
        self.assertEqual(len(self.before),26)
        for p in [ROOT/'output',ROOT/'phase2/output',ROOT/'scripts',ROOT/'tests/sub']:
            with self.assertRaisesRegex(ValueError,'禁止覆盖'):
                b.run(ROOT,p)

    def test_project_hours_are_inclusive_and_candidate_table_separate(self):
        goals,projects,knowledge,practice,requirements,assessment,hours=self.catalog
        self.assertEqual(sum(p['总学时'] for p in projects),64)
        self.assertEqual(sum(p['理论学时'] for p in projects),32)
        self.assertEqual(sum(p['内含实践学时'] for p in projects),32)
        self.assertEqual(hours[1]['总学时'],32)
        self.assertEqual([p['单独分配学时'] for p in practice if p['条目ID'].startswith('RP')],[10,12,10])
        self.assertTrue(all(p['单独分配学时'] is None for p in practice if p['条目ID'].startswith('EP')))
        self.assertEqual(len(projects),8)

    def test_historical_candidate_evidence_not_promoted(self):
        self.assertTrue(all(not m['实践任务ID'].startswith('RP') for m in self.matrix))
        self.assertTrue(all(a['考核比例'] is None for a in self.catalog[5]))
        self.assertTrue(all(a['正式考核方案可确认']=='否' for a in self.catalog[5]))

    def test_every_course_quote_resolves_to_pdf_page_and_coordinates(self):
        b.validate(self.corpus,self.data,self.catalog,self.matrix,self.summary,self.graph)
        for e in self.corpus.evidence:
            self.assertTrue(1<=e['PDF页码']<=len(self.corpus.pages))
            self.assertEqual(e['PDF_SHA256'],b.sha(self.pdf))
            self.assertTrue(e['定位矩形'])
            self.assertTrue(e['原文摘录'])
        cross=next(p for p in self.catalog[3] if p['条目ID']=='EP02')
        self.assertEqual(cross['PDF页码'],'4;5')

    def test_no_added_knowledge_without_pdf_quote(self):
        self.assertEqual(len(self.catalog[2]),37)
        for k in self.catalog[2]:
            self.assertTrue(k['课程证据ID'])
            self.assertTrue(k['课程标准原文摘录'])
        bad=copy.deepcopy(self.specs);bad['knowledge'][0]['quote']='原课标绝不存在的虚构课程知识ABC123'
        c=b.Corpus(self.pdf)
        try:
            with self.assertRaisesRegex(ValueError,'原文引文未找到'):
                b.parse_course(c,bad)
        finally:
            c.doc.close()

    def test_34_ability_partition_and_explicit_denominators(self):
        self.assertEqual(len(self.summary),34)
        counts=self.stats['能力匹配状态']
        self.assertEqual(counts,{'未见明确支撑':14,'直接支撑':1,'部分支撑':18,'待人工审核':1})
        status,recruit,known=b.coverage_tables(self.data,self.catalog,self.matrix,self.summary,self.suggestions)
        self.assertEqual(sum(r['能力数'] for r in status[:4]),34)
        self.assertTrue(all(r['统计分母']==34 for r in status))
        self.assertEqual([r['统计分母'] for r in recruit],[50,50,188])
        self.assertEqual(self.stats['已核验招聘原文明示数'],0)

    def test_matching_is_many_to_many_and_ids_valid(self):
        bycap={};byk={}
        for m in self.matrix:
            bycap.setdefault(m['能力ID'],set()).add(m['知识点ID'])
            if m['知识点ID']:
                byk.setdefault(m['知识点ID'],set()).add(m['能力ID'])
            self.assertIn(m['匹配状态'],b.STATUSES)
            self.assertEqual(m['人工审核状态'],b.PENDING)
        self.assertGreater(len(bycap['N113']),1)
        self.assertTrue(any(len(v)>1 for v in byk.values()))
        bad=copy.deepcopy(self.plan);bad['abilities'][2]['refs'][0]['knowledge']='NONEXISTENT'
        with self.assertRaisesRegex(ValueError,'知识点ID无效'):
            b.match_course(self.data,bad,self.catalog,self.corpus)

    def test_unseen_does_not_make_positive_support_edge(self):
        states={s['能力ID']:s['最终候选匹配状态'] for s in self.summary}
        for m in self.matrix:
            if m['匹配状态']=='未见明确支撑':
                self.assertFalse(m['课程证据ID'])
                self.assertFalse(m['课程标准原文摘录'])
                self.assertFalse(m['知识点ID'])
        for e in self.graph.edges:
            if e['关系类型'].startswith('能力候选'):
                self.assertNotEqual(states[e['起点ID']],'未见明确支撑')
        self.assertEqual(states['N221'],'未见明确支撑') # E-R不是预测/分析模型
        self.assertEqual(states['N232'],'未见明确支撑') # GUI不是看板
        self.assertEqual(states['N332'],'待人工审核') # 思政审计不等于审计训练

    def test_109_suspect_relations_never_promoted_to_confirmed(self):
        self.assertEqual(len(self.data['suspect']),109)
        for rid in self.data['suspect']:
            edges=[e for e in self.graph.edges+self.graph.quarantine if rid in e['源记录ID'].split(';')]
            self.assertTrue(edges)
            self.assertTrue(all(e['已人工确认']=='否' for e in edges))
            self.assertTrue(all('疑似' in e['审核或置信状态'] for e in edges))
        self.assertTrue(all(e['可用于已确认招聘频次']=='否' for e in self.graph.edges+self.graph.quarantine))

    def test_graph_unique_nodes_valid_endpoints_and_mapping(self):
        self.assertEqual(len(self.graph.ids),len(self.graph.nodes))
        self.assertEqual(len(self.graph.nodes),314)
        ids=[e['关系ID'] for e in self.graph.edges+self.graph.quarantine]
        self.assertEqual(len(ids),len(set(ids)))
        for e in self.graph.edges+self.graph.quarantine:
            self.assertIn(e['起点ID'],self.graph.ids);self.assertIn(e['终点ID'],self.graph.ids)
        maps={(e['起点ID'],e['终点ID']) for e in self.graph.edges if e['关系类型']=='旧编码建议映射新编码'}
        self.assertEqual(maps,{(m['原编码'],m['新编码']) for m in self.data['mapping'] if m['原编码']})

    def test_isolated_edges_keep_invalid_migrations_out_of_default_graph(self):
        self.assertTrue(self.graph.quarantine)
        self.assertTrue(any('不建议迁移' in e['匹配或关系理由'] for e in self.graph.quarantine))
        self.assertTrue(any(e['终点ID'].startswith('RP') for e in self.graph.quarantine))
        self.assertFalse(any(e['终点ID'].startswith('RP') for e in self.graph.edges))

    def test_course_suggestions_are_separate_from_34_and_have_no_job_evidence(self):
        self.assertEqual({s['课程建议能力ID'] for s in self.suggestions},{'S111','S112','S113'})
        self.assertTrue(all(s['招聘需求频次']==0 for s in self.suggestions))
        self.assertFalse({s['课程建议能力ID'] for s in self.suggestions}&{s['能力ID'] for s in self.summary})
        self.assertTrue(all(s['课程标准原文摘录'] for s in self.suggestions))
        for e in self.graph.edges:
            if e['终点ID'].startswith('S'):
                self.assertFalse(e['起点ID'].startswith('T-JOB'))

    def test_manual_review_is_not_simulated(self):
        issues=b.make_issues(self.data,self.catalog,self.matrix,self.summary,self.corpus,self.graph)
        self.assertTrue(all(not r['审核人'] and not r['审核结论'] for r in issues))
        self.assertTrue(all(r['审核状态']==b.PENDING for r in issues))
        self.assertFalse(any(r['问题类别'] in {'PDF可见性','招聘原文不足','关系疑似误匹配'} for r in issues))
        self.assertTrue(all(r['执行范围']=='当前课程能力图谱复核' for r in issues))

    def test_work_behavior_boundaries_and_four_downgrades(self):
        index={s['能力ID']:s for s in self.summary}
        for r in self.summary+self.matrix:
            self.assertTrue(r['课程支撑的具体技术行为'])
            self.assertTrue(r['课程尚未明确覆盖的工作行为'])
        for code in ['N313','N321','N322','N341']:
            self.assertEqual(index[code]['最终候选匹配状态'],'部分支撑')
        self.assertEqual(index['N113']['最终候选匹配状态'],'直接支撑')
        self.assertIn('仅限',index['N113']['课程支撑判断'])
        self.assertEqual(index['N331']['最终候选匹配状态'],'部分支撑')
        self.assertEqual(index['N332']['最终候选匹配状态'],'待人工审核')
        compare=b.revision_comparison(self.summary)
        self.assertEqual(len(compare),34)
        self.assertEqual({r['能力ID'] for r in compare if r['等级是否变化']=='是'},{'N313','N321','N322','N341'})

    def test_34_by_8_project_review_complete(self):
        wide,cells=b.project_review(self.catalog,self.matrix,self.summary)
        self.assertEqual(len(wide),34)
        self.assertEqual(len(cells),272)
        self.assertEqual(len({(c['能力ID'],c['课程项目ID']) for c in cells}),272)
        for cell in cells:
            self.assertTrue(cell['项目审阅范围证据ID'])
            if cell['项目层匹配等级'] in {'直接支撑','部分支撑'}:
                self.assertTrue(cell['正向课程证据ID'])
            else:
                self.assertFalse(cell['正向课程证据ID'])
        self.assertTrue(all(c['项目层匹配等级']=='待人工审核' for c in cells if c['能力ID'] in {'N331','N512'}))

    def test_unseen_14_recommendations_not_standard_facts(self):
        from collections import Counter
        rows=[s for s in self.summary if s['最终候选匹配状态']=='未见明确支撑']
        counts=Counter(s['未见支撑建议分类'] for s in rows)
        self.assertEqual(counts,{'本课程潜在改进方向':4,'其他专业课程可能承担':5,'旅游业务背景能力':4,'证据不足无法判断':1})
        self.assertTrue(all('研究建议' in r['分类性质'] for r in rows))
        self.assertEqual(len(rows),14)

    def test_current_scope_does_not_require_recruitment_checks(self):
        report=(self.out/'第三阶段研究报告.md').read_text()
        self.assertIn('不开展或新增',report)
        self.assertIn('进入第四阶段不以招聘原文或109条关系核验为条件',report)
        w=load_workbook(self.out/'人工审核清单_第三阶段.xlsx')
        self.assertIn('历史问题归档不再执行',w.sheetnames)
        self.assertIn('课程能力关系审核索引',w.sheetnames)
        self.assertNotIn('图谱全部待审关系索引',w.sheetnames)
        w.close()

    def test_project_and_graph_level_not_copy_whole_course_grade(self):
        rows=[m for m in self.matrix if m['能力ID']=='N113' and m['课程项目ID']=='P06']
        self.assertTrue(rows)
        self.assertTrue(all(m['匹配状态']=='部分支撑' and m['能力整体支撑等级']=='直接支撑' for m in rows))
        edges=[e for e in self.graph.edges if e['起点ID']=='N113' and e['终点ID']=='P06']
        self.assertTrue(edges)
        self.assertTrue(all('部分支撑' in e['审核或置信状态'] and '；直接支撑' not in e['审核或置信状态'] for e in edges))

    def test_semantic_results_repeat_without_overwriting_inputs(self):
        other=Path(self.temp.name)/'again'
        stats=b.run(ROOT,other)
        self.assertEqual(stats,self.stats)
        for f in self.out.glob('*.xlsx'):
            x,y=load_workbook(f),load_workbook(other/f.name)
            self.assertEqual(x.sheetnames,y.sheetnames)
            for s in x.sheetnames:
                self.assertEqual(list(x[s].values),list(y[s].values))
            x.close();y.close()
        self.assertEqual((self.out/'第三阶段研究报告.md').read_text(),(other/'第三阶段研究报告.md').read_text())
        self.assertEqual(self.before,{p:b.sha(ROOT/p) for p in self.before})


if __name__=='__main__':
    unittest.main()
