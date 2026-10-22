"""第三阶段：仅依据实际课程PDF构建可追溯的匹配矩阵与候选图谱数据。"""
import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import fitz
import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT=Path(__file__).resolve().parents[1]
HERE=Path(__file__).resolve().parent
PENDING='待人工审核'
STATUSES={'直接支撑','部分支撑','未见明确支撑','待人工审核'}
JOBFILE='旅游行业招聘岗位样本库_50条候选_待核验.xlsx'
SHEETS=[(JOBFILE,'招聘岗位主库','jobs'),('output/岗位任务编码表.xlsx','任务编码','tasks'),
 ('output/岗位任务能力关系表.xlsx','关系明细','old_rel'),('output/三级职业能力指标体系.xlsx','三级能力指标','old_cap'),
 ('phase2/output/三级职业能力体系_优化建议版.xlsx','优化岗位能力全层级','caps'),
 ('phase2/output/三级职业能力体系_优化建议版.xlsx','课程建议能力单列','suggestions'),
 ('phase2/output/职业能力体系_优化对照表.xlsx','新旧编码映射','mapping'),
 ('phase2/output/职业能力体系_优化对照表.xlsx','原体系逐项审核','old_nodes'),
 ('phase2/output/人工审核问题清单.xlsx','人工审核问题','inherited'),
 ('phase2/output/任务能力关系_复核表.xlsx','原关系逐条复核','reviewed_rel'),
 ('phase2/output/任务能力关系_复核表.xlsx','新旧关系迁移判定','proposals')]


def norm(text):
    return re.sub(r'\s+','',str(text or ''))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_rows(path,sheet,required=()):
    wb=load_workbook(path,data_only=False)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f'缺少工作表：{path}/{sheet}')
        values=list(wb[sheet].values)
        head=list(values[0])
        if len(head)!=len(set(head)) or not set(required)<=set(head):
            raise ValueError(f'字段缺失或重复：{path}/{sheet}')
        return [dict(zip(head,r)) for r in values[1:] if any(v is not None for v in r)]
    finally:
        wb.close()


def keyed(rows,key):
    vals=[r[key] for r in rows]
    if any(not v for v in vals) or len(set(vals))!=len(vals):
        raise ValueError('ID为空或重复：'+key)
    return {r[key]:r for r in rows}


def inventory(root,paths,commit):
    result,missing=[],[]
    for rel in sorted(set(paths)):
        wb=load_workbook(root/rel,data_only=False)
        for ws in wb:
            vals=list(ws.values);rows=[r for r in vals[1:] if any(v is not None for v in r)]
            result.append({'文件':rel,'工作表':ws.title,'字段数':len(vals[0]),'有效记录数':len(rows),
                           '实际字段':'；'.join(str(x) for x in vals[0]),'版本来源main提交':commit,'文件SHA256':sha(root/rel)})
            for i,h in enumerate(vals[0]):
                n=sum(v[i] is None or isinstance(v[i],str) and not v[i].strip() for v in rows)
                missing.append({'文件':rel,'工作表':ws.title,'字段':h,'记录分母':len(rows),'空值数':n,
                                '空值说明':'含有意留空的审核人、补证及未匹配来源；空值不自动视为数据错误'})
        wb.close()
    return result,missing


class Corpus:
    def __init__(self,path):
        self.path=Path(path)
        self.doc=fitz.open(path)
        self.pages=[];self.full='';self.evidence=[]
        for index,page in enumerate(self.doc):
            raw=page.get_text(sort=True)
            offsets=[i for i,c in enumerate(raw) if not c.isspace()]
            compact=norm(raw)
            chars=[c for b in page.get_text('rawdict',sort=True)['blocks'] for l in b.get('lines',[]) for s in l['spans'] for c in s['chars'] if not c['c'].isspace()]
            traces=page.get_texttrace();last=max((t['seqno'] for t in traces),default=-1)
            masks=[]
            for drawing in page.get_drawings():
                fill=drawing.get('fill')
                if fill and min(fill)>.98 and drawing.get('fill_opacity',0)>.98 and drawing['seqno']>last and drawing['rect'].get_area()>1000:
                    masks.append(tuple(drawing['rect']))
            self.pages.append({'page':index+1,'raw':raw,'norm':compact,'offsets':offsets,'start':len(self.full),
                               'end':len(self.full)+len(compact),'chars':chars,'charseq':''.join(c['c'] for c in chars),'masks':masks})
            self.full+=compact

    def _add(self,page,quote,boxes,clip=None):
        p=self.pages[page-1]
        covered=sum(any((fitz.Rect(b)&fitz.Rect(m)).get_area()>=fitz.Rect(b).get_area()*.5 for m in p['masks']) for b in boxes)
        hidden=bool(boxes and covered/len(boxes)>.2)
        if not boxes:
            raise ValueError('引文无法定位坐标：'+quote[:40])
        record={'证据ID':f'CE-{len(self.evidence)+1:04d}','课程PDF':self.path.name,'PDF_SHA256':sha(self.path),
                'PDF页码':page,'原文摘录':quote,'校验用去空白原文':norm(quote),
                '定位矩形':json.dumps(list(fitz.Rect(clip)) if clip else list(fitz.Rect(min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes)))),
                '课程证据状态':'遮挡文本层候选，待人工核验' if hidden else '可见文本层摘录，未人工核验',
                '是否遮挡':'是' if hidden else '否','原文处理':'保留原文字词与原提取换行；仅匹配校验时去除空白','人工审核状态':PENDING}
        self.evidence.append(record)
        return record['证据ID']

    def span(self,start,end):
        ids=[]
        for p in self.pages:
            a=max(start,p['start'])-p['start'];b=min(end,p['end'])-p['start']
            if a>=b:
                continue
            token=p['norm'][a:b]
            loc=p['charseq'].find(token)
            if loc<0:
                raise ValueError(f'页{p["page"]}阅读顺序不可靠，请使用单元格坐标引用')
            boxes=[c['bbox'] for c in p['chars'][loc:loc+len(token)]]
            quote=p['raw'][p['offsets'][a]:p['offsets'][b-1]+1]
            ids.append(self._add(p['page'],quote,boxes))
        if not ids:
            raise ValueError('空引文')
        return ids

    def cite(self,quote,pages=None):
        token=norm(quote)
        for p in self.pages:
            if pages and p['page'] not in pages:
                continue
            loc=p['norm'].find(token)
            if loc>=0:
                return self.span(p['start']+loc,p['start']+loc+len(token))
        if pages:
            first=self.pages[min(pages)-1]['start'];last=self.pages[max(pages)-1]['end']
            loc=self.full[first:last].find(token)
            if loc>=0:
                return self.span(first+loc,first+loc+len(token))
        raise ValueError('原文引文未找到：'+quote)

    def cell(self,page,bbox):
        raw=self.doc[page-1].get_textbox(fitz.Rect(bbox))
        p=self.pages[page-1]
        boxes=[c['bbox'] for c in p['chars'] if fitz.Rect(bbox).contains(fitz.Point((c['bbox'][0]+c['bbox'][2])/2,(c['bbox'][1]+c['bbox'][3])/2))]
        return self._add(page,raw,boxes,bbox)

    def fields(self,ids):
        records=[self.evidence[int(i.split('-')[1])-1] for i in dict.fromkeys(ids)]
        return {'课程证据ID':';'.join(r['证据ID'] for r in records),
                '课程标准原文摘录':'\n\n'.join(f"[{r['证据ID']}，PDF第{r['PDF页码']}页]\n{r['原文摘录']}" for r in records),
                'PDF页码':';'.join(str(p) for p in sorted({r['PDF页码'] for r in records})),
                '课程证据状态':'遮挡候选证据，待核验' if any(r['是否遮挡']=='是' for r in records) else '可见文本摘录，待人工核验'}


def discover_pdf(root,explicit=None):
    files=[Path(explicit).resolve()] if explicit else sorted(p for p in root.rglob('*') if p.is_file() and p.suffix.lower()=='.pdf' and not any(x in p.relative_to(root).parts for x in ['work','.git','.venv']))
    candidates=[]
    for p in files:
        with fitz.open(p) as d:
            head=''.join(norm(x.get_text()) for x in list(d)[:3])
            if '数据库原理及应用' in head:
                candidates.append(p)
    if len(candidates)!=1:
        raise ValueError('课程PDF不能唯一识别，请人工确认：'+str([str(x) for x in candidates]))
    return candidates[0],files


def parse_course(corpus,specs):
    goals=[]
    markers=[('G-K','知识目标','（一）知识目标','（二）能力目标'),
             ('G-A','能力目标','（二）能力目标','（三）素质目标'),
             ('G-Q','素质目标','（三）素质目标','3、课程内容')]
    for prefix,kind,begin,end in markers:
        a=corpus.full.index(begin)+len(begin);b=corpus.full.index(end,a)
        part=corpus.full[a:b];marks=list(re.finditer(r'([1-9])\.',part))
        for i,m in enumerate(marks):
            start=a+m.end();stop=a+(marks[i+1].start() if i+1<len(marks) else len(part))
            ids=corpus.span(start,stop)
            goals.append({'条目ID':f'{prefix}{int(m.group(1)):02d}','目标类别':kind,'目标原文':corpus.full[start:stop],**corpus.fields(ids),'人工审核状态':PENDING})
    pattern=r'项目([一二三四五六七八])(.+?)\(共(\d+)学时，理论(\d+)学时，(?:实践|实训)(\d+)学时\)'
    heads=list(re.finditer(pattern,corpus.full));projects=[];requirements=[]
    if len(heads)!=8:
        raise ValueError('实际教学项目标题无法可靠解析，请人工核验')
    for i,m in enumerate(heads):
        pid=f'P{i+1:02d}';end=heads[i+1].start() if i+1<len(heads) else corpus.full.index('6"8实<=1内容',m.end())
        body=corpus.full[m.end():end]
        content=re.search(r'1\.教学内容(.*?)2\.教学方式与要求',body)
        if not content:
            raise ValueError('教学内容边界不明：'+pid)
        ids=corpus.span(m.start(),m.end())+corpus.span(m.end()+content.start(1),m.end()+content.end(1))
        projects.append({'条目ID':pid,'教学项目名称':m.group(2),'总学时':int(m.group(3)),
                         '理论学时':int(m.group(4)),'内含实践学时':int(m.group(5)),
                         '统计口径':'64学时教学项目口径；实践已包含在项目总学时内',
                         '项目跨页范围':';'.join(str(p['page']) for p in corpus.pages if p['start']<end and p['end']>m.start()),'教学内容原文':content.group(1),**corpus.fields(ids),'人工审核状态':PENDING})
        a=body.index('（2）教学要求')+len('（2）教学要求')
        tail=body[a:]
        if tail.startswith('学生要求：'):
            a+=len('学生要求：');tail=body[a:]
        endings=[v for marker in ['教师要求：','教室需求：','3.教学重难点'] if (v:=tail.find(marker))>=0]
        if not endings:
            raise ValueError('教学要求边界不明：'+pid)
        ids=corpus.span(m.end()+a,m.end()+a+min(endings))
        requirements.append({'条目ID':f'CR{i+1:02d}','课程项目ID':pid,'条目类别':'项目学生教学要求',**corpus.fields(ids),'人工审核状态':PENDING})
    knowledge=[]
    for item in specs['knowledge']:
        ids=corpus.cite(item['quote'],[item['page']])
        knowledge.append({'条目ID':item['id'],'课程项目ID':item['project'],'知识点名称':item['name'],
                          '课程条目定位':item['scope'],'条目命名说明':'依据原文归纳命名，不新增原文外内容',**corpus.fields(ids),'人工审核状态':PENDING})
    practice=[]
    for item in specs['embedded_practice']:
        pages=item['page'] if isinstance(item['page'],list) else [item['page']]
        ids=corpus.cite(item['quote'],pages)
        practice.append({'条目ID':item['id'],'课程项目ID':item['project'],'实训任务名称':item['name'],
                         '实训类型':'教学项目内课堂或课后实操要求','知识点ID':';'.join(item['knowledge']),
                         '单独分配学时':None,'学时口径':'已列在教学项目内，未单独分配；不可再叠加学时',
                         **corpus.fields(ids),'人工审核状态':PENDING})
    tables=[]
    with pdfplumber.open(corpus.path) as doc:
        for page in doc.pages:
            for table in page.find_tables():
                cells=table.extract()
                if cells and norm(''.join(cells[0]))=='实训项目名称内容与任务学时':
                    tables.append((page.page_number,table,cells))
    if len(tables)!=1:
        raise ValueError('实践任务表不能唯一解析，须人工核验')
    page,table,cells=tables[0]
    total=int(cells[-1][-1]);hours=[]
    for i,row in enumerate(cells[1:-1],1):
        ids=[corpus.cell(page,b) for b in table.rows[i].cells if b]
        hours.append(int(row[-1]))
        practice.append({'条目ID':f'RP{i:02d}','课程项目ID':'','实训任务名称':row[0],
                         '实训类型':'单列实践表候选（被遮挡文本层）','知识点ID':'',
                         '单独分配学时':int(row[-1]),'学时口径':'文本层32学时表；可见性及与64学时关系均待核验，不加总',
                         **corpus.fields(ids),'人工审核状态':'待人工核验（遮挡内容）'})
    if sum(hours)!=total:
        raise ValueError('实践表文本层合计不一致')
    assessment=[]
    ids=corpus.cite('课堂即时反馈与操作检查教师巡查学生安装与连接过程中的问题，及时纠错，确保学生能够独立完成环境部署工作。',[3])
    assessment.append({'条目ID':'AS01','条目类别':'教学方式中的即时操作检查','考核比例':None,'正式考核方案可确认':'否',
                       '说明':'可见教学反馈要求，不等同正式考核组成或权重',**corpus.fields(ids),'人工审核状态':PENDING})
    ids=corpus.cite('（一）课程考核组成与评价',[page])
    assessment.append({'条目ID':'AS02','条目类别':'考核章节标题候选（遮挡）','考核比例':None,'正式考核方案可确认':'否',
                       '说明':'仅提取标题且被遮挡，未找到组成、比例、评分标准等明细',**corpus.fields(ids),'人工审核状态':PENDING})
    hours_info=[{'口径ID':'H64','口径':'可见教学项目','条目数':len(projects),'总学时':sum(p['总学时'] for p in projects),
                 '理论学时':sum(p['理论学时'] for p in projects),'内含实践学时':sum(p['内含实践学时'] for p in projects),
                 '核验结论':'8项目64=32理论+32内含实践；分项内含不重复计数','人工审核状态':PENDING},
                {'口径ID':'H32-CAND','口径':'被遮挡文本层的单列实践表','条目数':3,'总学时':total,'理论学时':None,'内含实践学时':None,
                 '核验结论':'10+12+10=32仅文本层可提取；不能确认其可见有效性及与64学时关系，禁止相加为96','人工审核状态':'待人工核验（遮挡及口径）'}]
    return goals,projects,knowledge,practice,requirements,assessment,hours_info


def load_inputs(root,snapshot):
    if not snapshot.get('pr2_merged'):
        raise ValueError('PR #2未确认合并，禁止使用第一阶段替代34项能力库')
    required={
        'jobs':['样本ID','岗位名称','岗位职责摘要（转述）','证据摘录'],
        'tasks':['任务编码','岗位ID','证据文本'],
        'old_rel':['岗位ID','任务编码','三级能力编码'],
        'old_cap':['三级能力编码'],
        'caps':['新编码','父编码','能力层级','能力名称','可观察的工作行为','来源岗位ID','对应任务ID','证据类型'],
        'suggestions':['新编码','父编码','能力名称','能力层级'],
        'mapping':['原编码','新编码','优化动作'],
        'old_nodes':['原编码','原能力名称'],
        'inherited':['问题ID','问题类型','原关系ID'],
        'reviewed_rel':['原关系ID','岗位ID','任务编码','三级能力编码','疑似误匹配','复核说明'],
        'proposals':['原关系ID','任务ID','岗位ID','新三级能力编码','关系建议','证据类型']}
    data={}
    for path,sheet,key in SHEETS:
        if not (root/path).exists():
            raise ValueError('未找到第二阶段或前置输入：'+path+'；禁止回退第一阶段能力')
        data[key]=read_rows(root/path,sheet,required[key])
    jobs=keyed(data['jobs'],'样本ID');tasks=keyed(data['tasks'],'任务编码')
    caps=keyed(data['caps'],'新编码')
    data['leaf']=[c for c in data['caps'] if c['能力层级']=='三级']
    for t in data['tasks']:
        if t['岗位ID'] not in jobs or t['证据文本'] not in jobs[t['岗位ID']]['岗位职责摘要（转述）']:
            raise ValueError('原任务无法追溯')
    reviewed=keyed(data['reviewed_rel'],'原关系ID')
    original_keys={(r['岗位ID'],r['任务编码'],r['三级能力编码']) for r in data['old_rel']}
    reviewed_keys={(r['岗位ID'],r['任务编码'],r['三级能力编码']) for r in data['reviewed_rel']}
    if original_keys!=reviewed_keys:
        raise ValueError('第二阶段原关系与第一阶段不一致')
    for p in data['proposals']:
        if p['任务ID'] not in tasks or p['岗位ID'] not in jobs or p['新三级能力编码'] not in caps or p['原关系ID'] not in reviewed:
            raise ValueError('第二阶段迁移关系端点或原关系ID不一致')
        if tasks[p['任务ID']]['岗位ID']!=p['岗位ID']:
            raise ValueError('任务与岗位外键错误')
    map_codes={r['新编码'] for r in data['mapping']}
    if not set(caps)<=map_codes:
        raise ValueError('新编码缺少旧→新映射')
    inherited_suspect={r['原关系ID'] for r in data['inherited'] if r['问题类型']=='关系疑似误匹配'}
    suspect={r['原关系ID'] for r in data['reviewed_rel'] if r['疑似误匹配']=='是'}
    if inherited_suspect!=suspect:
        raise ValueError('第二阶段109条疑似关系清单与复核表不一致')
    data['suspect']=suspect
    return data


def match_course(data,plan,catalog,corpus):
    goals,projects,knowledge,practice,requirements,assessment,hours_info=catalog
    ks=keyed(knowledge,'条目ID');es=keyed(practice,'条目ID');gs=keyed(goals,'条目ID')
    caps=keyed(data['leaf'],'新编码')
    if set(caps)!={p['ability'] for p in plan['abilities']}:
        raise ValueError('实际候选能力与逐项匹配方案不一致，禁止替代或静默漏项')
    matrix=[];summaries=[]
    for item in plan['abilities']:
        cap=caps[item['ability']]
        if item['status'] not in STATUSES:
            raise ValueError('无效匹配状态')
        refs=item['refs'] or [{'knowledge':'','practice':''}]
        supplemental=[]
        if item.get('supplementary_quote'):
            s=item['supplementary_quote'];supplemental=corpus.cite(s['quote'],[s['page']])
        all_ids=[]
        for ref in refs:
            kid,pid=ref['knowledge'],ref['practice'];goalid=item.get('goal','');ids=[]
            project=''
            if kid:
                if kid not in ks:
                    raise ValueError('知识点ID无效：'+kid)
                ids+=ks[kid]['课程证据ID'].split(';');project=ks[kid]['课程项目ID']
            if pid:
                if pid not in es or es[pid]['实训类型']!='教学项目内课堂或课后实操要求':
                    raise ValueError('不得用被遮挡实践表确认能力实训支撑')
                if kid not in es[pid]['知识点ID'].split(';') or project!=es[pid]['课程项目ID']:
                    raise ValueError('知识点、项目、实训引用不一致')
                ids+=es[pid]['课程证据ID'].split(';')
            if goalid:
                ids+=gs[goalid]['课程证据ID'].split(';')
            ids+=supplemental;ids=list(dict.fromkeys(ids));all_ids+=ids
            fields=corpus.fields(ids) if ids else {'课程证据ID':'','课程标准原文摘录':'','PDF页码':'','课程证据状态':'无正向引文；全PDF可见内容查阅后未见明确支撑'}
            if item['status'] in {'直接支撑','部分支撑'} and (not ids or any(corpus.evidence[int(x.split('-')[1])-1]['是否遮挡']=='是' for x in ids)):
                raise ValueError('直接或部分支撑不能使用无引文或遮挡证据')
            matrix.append({'匹配ID':f'M-{len(matrix)+1:04d}','能力ID':cap['新编码'],'能力名称':cap['能力名称'],
                '可观察工作行为':cap['可观察的工作行为'],'课程项目ID':project,'知识点ID':kid,
                '实践任务ID':pid,'课程目标ID':goalid,'匹配状态':item['status'],**fields,
                '查阅范围':f'PDF物理页1—{len(corpus.pages)}；遮挡文本层仅作候选，不作为正向确认',
                '匹配理由':item['reason'],'课程支撑边界或改进建议':item['gap'],
                '招聘证据类型':cap['证据类型'],'招聘证据状态':'候选岗位全部待核验；无招聘原文明示',
                '原课程适配类别（不等同培养）':cap['课程适配类别'],
                '关系性质':'逐项行为比较形成的研究推断，不是培养效果验证','人工审核状态':PENDING})
        supports=[p for p in data['proposals'] if p['新三级能力编码']==cap['新编码'] and p['关系建议'] in {'摘要支持（待审）','待人工判断（仅推断）'}]
        suspected={p['原关系ID'] for p in supports if p['原关系ID'] in data['suspect']}
        summaries.append({'能力ID':cap['新编码'],'能力名称':cap['能力名称'],'最终候选匹配状态':item['status'],
            '知识点ID':';'.join(dict.fromkeys(r['knowledge'] for r in refs if r['knowledge'])),
            '实践任务ID':';'.join(dict.fromkeys(r['practice'] for r in refs if r['practice'])),
            '课程证据ID':';'.join(dict.fromkeys(all_ids)),
            '课程支撑判断':item['reason'],'课程标准支撑不足或边界':item['gap'],
            '岗位需求证据不足':'全部来源岗位未核验；摘要支持不是招聘原文明示',
            '来源岗位ID':cap['来源岗位ID'],'对应任务ID':cap['对应任务ID'],
            '摘要支持候选岗位数（非独立需求频次）':cap['摘要支持岗位数（含疑似重复）'],
            '涉及疑似原关系ID':';'.join(sorted(suspected)),'涉及疑似原关系数':len(suspected),
            '已核验招聘原文岗位数':0,'统计分母说明':'能力状态比例分母为实际三级候选能力数；不按岗位频次加权',
            '人工审核状态':PENDING})
    return matrix,summaries


class Graph:
    def __init__(self):
        self.nodes=[];self.edges=[];self.quarantine=[];self.ids=set();self.edge_index={}

    def node(self,id,kind,name,source,attrs,fields=None):
        if id in self.ids:
            raise ValueError('节点ID重复：'+id)
        self.ids.add(id)
        self.nodes.append({'节点ID':id,'节点类型':kind,'节点名称':name,'源文件或条目':source,
            '来源属性JSON':json.dumps(attrs,ensure_ascii=False,default=str),
            '课程证据ID':(fields or {}).get('课程证据ID',''),'PDF页码':(fields or {}).get('PDF页码',''),
            '证据或核验状态':(fields or {}).get('课程证据状态','岗位待核验或研究建议待审'),
            '人工审核状态':PENDING})

    def edge(self,start,end,kind,source,record='',ev='',recruit='研究推断',state=PENDING,reason='',quarantine=False):
        dest=self.quarantine if quarantine else self.edges
        key=(quarantine,start,end,kind)
        if key in self.edge_index:
            old=self.edge_index[key]
            for field,value in [('源记录ID',record),('课程证据ID',ev),('来源文件',source),('招聘证据类型',recruit),('审核或置信状态',state)]:
                old[field]=';'.join(dict.fromkeys(v for v in (str(old[field])+';'+str(value)).split(';') if v))
            return
        row={'关系ID':f'{"QEDGE" if quarantine else "EDGE"}-{len(dest)+1:05d}',
             '起点ID':start,'终点ID':end,'关系类型':kind,'来源文件':source,'源记录ID':record,
             '课程证据ID':ev,'招聘证据类型':recruit,'审核或置信状态':state,
             '匹配或关系理由':reason,'关系集':'隔离（不建议迁移或遮挡候选）' if quarantine else '候选图谱关系，非确认事实全集',
             '已人工确认':'否','可用于已确认招聘频次':'否'}
        dest.append(row);self.edge_index[key]=row


def build_graph(data,catalog,matrix,corpus):
    g=Graph();goals,projects,knowledge,practice,requirements,assessment,hours_info=catalog
    g.node('COURSE-PDF','课程标准文件',corpus.path.name,corpus.path.name,{'SHA256':sha(corpus.path),'物理页数':len(corpus.pages),'学时口径':hours_info})
    for j in data['jobs']:
        g.node(j['样本ID'],'岗位候选样本',j['岗位名称'],JOBFILE,j)
    for t in data['tasks']:
        g.node(t['任务编码'],'典型工作任务（摘要）',t['典型工作任务'],'output/岗位任务编码表.xlsx',t)
        g.edge(t['岗位ID'],t['任务编码'],'岗位摘要包含任务','output/岗位任务编码表.xlsx',t['任务编码'],recruit='岗位摘要支持',state='岗位待核验；摘要不等于招聘原文')
    for c in data['caps']+data['suggestions']:
        g.node(c['新编码'],'职业能力建议' if c in data['caps'] else '课程建议能力（无招聘证据）',c['能力名称'],'phase2/output/三级职业能力体系_优化建议版.xlsx',c)
        if c['父编码']:
            g.edge(c['父编码'],c['新编码'],'建议能力层级包含','phase2/output/三级职业能力体系_优化建议版.xlsx',c['新编码'])
    for c in data['old_nodes']:
        g.node(c['原编码'],'原职业能力（归档）',c['原能力名称'],'phase2/output/职业能力体系_优化对照表.xlsx',c)
        if len(c['原编码'])>2:
            g.edge(c['原编码'][:-1],c['原编码'],'原能力层级归档','phase2/output/职业能力体系_优化对照表.xlsx',c['原编码'])
    for c in data['mapping']:
        if c['原编码']:
            g.edge(c['原编码'],c['新编码'],'旧编码建议映射新编码','phase2/output/职业能力体系_优化对照表.xlsx',c['原编码']+'→'+c['新编码'],reason=c['优化动作'])
    for r in data['reviewed_rel']:
        g.edge(r['任务编码'],r['三级能力编码'],'原任务能力研究关系（归档）','phase2/output/任务能力关系_复核表.xlsx',r['原关系ID'],
               recruit='岗位摘要支持;研究推断',state='待人工审核（疑似）' if r['原关系ID'] in data['suspect'] else PENDING,
               reason=r['复核说明'])
    for r in data['proposals']:
        reject=r['关系建议'].startswith('不建议迁移')
        state='待人工审核（含疑似原关系）' if r['原关系ID'] in data['suspect'] else PENDING
        g.edge(r['任务ID'],r['新三级能力编码'],'任务候选关联新能力','phase2/output/任务能力关系_复核表.xlsx',r['原关系ID'],
               recruit=r['证据类型'],state=state,reason=r['关系建议'],quarantine=reject)
    for kind,rows in [('课程目标',goals),('教学项目',projects),('知识条目',knowledge),('实训任务',practice),('课程能力要求',requirements),('考核相关条目',assessment)]:
        for item in rows:
            id=item['条目ID'];name=next((item[k] for k in ['教学项目名称','知识点名称','实训任务名称','目标原文','条目类别'] if k in item),kind)
            g.node(id,kind,name,corpus.path.name,item,item)
            hidden='遮挡' in item['课程证据状态']
            g.edge('COURSE-PDF',id,'课程PDF结构条目',corpus.path.name,id,item['课程证据ID'],recruit='不适用（课程证据）',
                   state='遮挡候选，待人工核验' if hidden else '课程可见原文结构化，未人工核验',quarantine=hidden)
            if item.get('课程项目ID'):
                g.edge(item['课程项目ID'],id,'教学项目包含条目',corpus.path.name,id,item['课程证据ID'],recruit='不适用（课程证据）',state='结构化归类，待人工核验',quarantine=hidden)
            if kind=='实训任务' and item.get('知识点ID'):
                for k in item['知识点ID'].split(';'):
                    g.edge(k,id,'知识条目对应实操要求',corpus.path.name,id,item['课程证据ID'],recruit='不适用（课程证据）',state='依课程动作归类，待人工核验')
    for m in matrix:
        if m['匹配状态']=='未见明确支撑':
            continue # 缺少支撑不是一条正向支撑边，不创造虚构知识节点。
        for source,kind in [(m['课程项目ID'],'能力候选对应教学项目'),(m['知识点ID'],'能力候选由知识条目支撑'),(m['实践任务ID'],'能力候选由实操训练'),(m['课程目标ID'],'能力候选对应课程目标')]:
            if source:
                g.edge(m['能力ID'],source,kind,'phase3/matching_plan.json',m['匹配ID'],m['课程证据ID'],
                       state=PENDING+'；'+m['匹配状态'],reason=m['匹配理由'])
    course_suggestions=[]
    ks=keyed(knowledge,'条目ID')
    for sid,keys in [('S111',['K02-01','K02-02','K02-03']),('S112',['K03-01']),('S113',['K04-01','K04-02','K04-03'])]:
        spec=next(s for s in data['suggestions'] if s['新编码']==sid)
        ids=list(dict.fromkeys(e for k in keys for e in ks[k]['课程证据ID'].split(';')))
        course_suggestions.append({'课程建议能力ID':sid,'能力名称':spec['能力名称'],'课程知识点ID':';'.join(keys),
            **corpus.fields(ids),'招聘证据状态':'无具体招聘证据，不属于34项岗位能力分母','招聘需求频次':0,
            '结论':'课程标准有相应内容，不得反推招聘岗位要求该能力','人工审核状态':PENDING})
        for key in keys:
            g.edge(sid,key,'无招聘证据课程建议对应课程内容','phase2/output/三级职业能力体系_优化建议版.xlsx',sid,ks[key]['课程证据ID'],
                   recruit='无招聘证据的课程建议',state=PENDING+'；不计岗位需求频次')
    # 将证据内容展开至边表，避免只有ID而缺少可直接审核的原文或摘要。
    task_index=keyed(data['tasks'],'任务编码');job_index=keyed(data['jobs'],'样本ID')
    for edge in g.edges+g.quarantine:
        tid=edge['起点ID'] if edge['起点ID'] in task_index else edge['终点ID'] if edge['终点ID'] in task_index else ''
        jobid=task_index[tid]['岗位ID'] if tid else ''
        edge['来源岗位ID']=jobid
        edge['来源任务ID']=tid
        edge['岗位摘要证据']=task_index[tid]['证据文本'] if tid else ''
        edge['原始招聘链接']=job_index[jobid]['原始招聘链接'] if jobid else ''
        fields=corpus.fields(list(filter(None,edge['课程证据ID'].split(';')))) if edge['课程证据ID'] else {}
        edge['课程PDF页码']=fields.get('PDF页码','')
        edge['课程标准原文摘录']=fields.get('课程标准原文摘录','')
        edge['课程证据可见性']=fields.get('课程证据状态','不适用（非课程原文关系）')
    return g,course_suggestions


def validate(corpus,data,catalog,matrix,summary,g):
    if len({n['节点ID'] for n in g.nodes})!=len(g.nodes):
        raise ValueError('图谱节点ID不唯一')
    ev={e['证据ID']:e for e in corpus.evidence}
    for edge in g.edges+g.quarantine:
        if edge['起点ID'] not in g.ids or edge['终点ID'] not in g.ids:
            raise ValueError('关系端点不存在：'+str(edge))
        for eid in filter(None,edge['课程证据ID'].split(';')):
            if eid not in ev:
                raise ValueError('关系课程证据ID失效')
        if edge['已人工确认']!='否':
            raise ValueError('不得伪造人工确认')
    edgeids=[e['关系ID'] for e in g.edges+g.quarantine]
    if len(set(edgeids))!=len(edgeids):
        raise ValueError('关系ID重复')
    leaf={a['新编码'] for a in data['leaf']}
    if {m['能力ID'] for m in matrix}!=leaf or {s['能力ID'] for s in summary}!=leaf:
        raise ValueError('34项能力匹配覆盖不完整')
    ks={r['条目ID'] for r in catalog[2]};ps={r['条目ID'] for r in catalog[1]};pr={r['条目ID'] for r in catalog[3]}
    for m in matrix:
        if m['知识点ID'] and m['知识点ID'] not in ks or m['课程项目ID'] and m['课程项目ID'] not in ps or m['实践任务ID'] and m['实践任务ID'] not in pr:
            raise ValueError('匹配矩阵ID失效')
        if m['匹配状态'] in {'直接支撑','部分支撑'}:
            if not m['课程证据ID']:
                raise ValueError('正向匹配没有引文')
            if any(ev[e]['是否遮挡']=='是' for e in m['课程证据ID'].split(';')):
                raise ValueError('遮挡证据不可用作确认支撑')
    for e in corpus.evidence:
        page=e['PDF页码']
        if not 1<=page<=len(corpus.pages):
            raise ValueError('PDF页码无效')
        if e['定位矩形']:
            raw=corpus.doc[page-1].get_textbox(fitz.Rect(json.loads(e['定位矩形'])))
            # 全段坐标可能含其他行；逐页文本是主校验，单元格用矩形补充。
            if norm(e['原文摘录']) not in corpus.pages[page-1]['norm'] and norm(e['原文摘录']) not in norm(raw):
                raise ValueError('原文摘录无法回溯到PDF页或单元格：'+e['证据ID'])


def write_book(path,sheets):
    wb=Workbook();wb.remove(wb.active)
    for title,rows in sheets:
        ws=wb.create_sheet(title[:31]);headers=list(rows[0]) if rows else ['说明']
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):
                    if len(cell.value)>32767:
                        raise ValueError('Excel单元格文本过长，应拆分至证据表')
                    cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245A81')
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(55,max(16,len(str(col[0].value))*2))
    wb.save(path);wb.close()


def make_issues(data,catalog,matrix,summaries,corpus,g):
    issues=[]
    def add(kind,priority,question,ref='',page='',evidence=''):
        issues.append({'第三阶段问题ID':f'P3-Q{len(issues)+1:04d}','问题类别':kind,'优先级':priority,
            '关联条目或匹配ID':ref,'PDF页码':page,'课程证据ID':evidence,
            '待人工核验事项':question,'审核人':'','审核结论':'','补证链接或位置':'','审核状态':PENDING})
    add('PDF可见性','高','第13页文本层实训表及考核标题被后绘制白色矩形覆盖；核验课程发布原稿，不能据此确认可见有效内容。','RP01;RP02;RP03;AS02',13,
        ';'.join(e['证据ID'] for e in corpus.evidence if e['是否遮挡']=='是'))
    add('学时口径','高','可见8项目64学时内含32实践；单列表32仅遮挡文本层候选。确认是否同一实践分配，禁止直接相加为96或重复累计。','H64;H32-CAND',13)
    add('文本笔误','高','遮挡文本层保留INERT INTO原字；不能静默纠正为INSERT INTO并声称课程原文已核验。','RP02',13)
    add('正式考核方案不足','高','未提取到可确认的考核组成、比例或评价细则；课堂即时操作检查不是完整考核方案，需补完整有效版本。','AS01;AS02','3;13')
    add('目标与实施落点','高','目标提及DCL与权限，思政提及加密审计，但具体技术项目和实训落点不明；不能据关键词推定完整培养。','KG-01;KG-02;KI-01','2;3;13')
    add('技术表述核验','中','第8页课程思政中DELETE与TRUNCATE的可回滚/不可逆表述应由教师核对适用条件；本阶段保留原文，不据常识改写标准。','P04',8)
    add('文本乱码','中','部分章节标题文本层出现乱码；项目标题与内容可解析，但章节名和末页结构须对照原始发布文件核验。','COURSE-PDF','1;2;3;13')
    for project in catalog[1]:
        if ';' in project['项目跨页范围']:
            add('跨页结构','中','该项目跨PDF物理页；核对内容、学生要求、思政及下一项目边界，已分证据保存页码。',project['条目ID'],project['项目跨页范围'],project['课程证据ID'])
    for s in summaries:
        add('能力课程匹配','高' if s['最终候选匹配状态'] in {'待人工审核','未见明确支撑'} else '中',
            s['课程支撑判断']+'；'+s['课程标准支撑不足或边界'],s['能力ID'],evidence=s['课程证据ID'])
    add('招聘原文不足','高','50条岗位全部待核验；摘要不能当招聘原文。第二阶段50条原文核验问题及109条疑似关系继续继承，不能被课程证据替代。','50岗位;109疑似原关系')
    add('图谱候选关系','高','图谱含研究映射与候选语义边；未人工确认的路径不能作为已核验需求—培养成效链。隔离关系不进入默认候选边集。','候选边及隔离边')
    add('统计解释','高','34项能力状态不按招聘频次加权；含部分支撑的文本关联比例不是培养达成率。岗位证据不足与课程标准支撑不足分别报告。','能力状态统计')
    return issues


def coverage_tables(data,catalog,matrix,summaries,course_suggestions):
    denominator=len(summaries);counts=Counter(s['最终候选匹配状态'] for s in summaries)
    stats=[{'匹配状态':status,'能力数':counts[status],'统计分母':denominator,'分母定义':'实际第二阶段三级候选能力数，不包含3项课程建议',
            '比例':counts[status]/denominator,'解释':'课程标准文本层研究匹配，不是学生培养达成率；全部待人工审核'}
           for status in ['直接支撑','部分支撑','未见明确支撑','待人工审核']]
    stats.append({'匹配状态':'直接或部分文本支撑（并集）','能力数':counts['直接支撑']+counts['部分支撑'],
                  '统计分母':denominator,'分母定义':'实际三级候选能力数','比例':(counts['直接支撑']+counts['部分支撑'])/denominator,
                  '解释':'只表示存在部分以上课程文本对应，不表示能力完整覆盖或实际培养效果'})
    recruit=[{'不足类型':'招聘原文待核验','数量':len(data['jobs']),'统计分母':len(data['jobs']),'分母定义':'岗位候选样本数','比例':1.0,
              '含义':'岗位真实性与具体要求未逐条核验，不能由课程PDF补足'},
             {'不足类型':'招聘原文明示已核验','数量':0,'统计分母':len(data['jobs']),'分母定义':'岗位候选样本数','比例':0.0,
              '含义':'本轮无新增招聘核验记录'},
             {'不足类型':'原任务能力关系疑似误匹配','数量':len(data['suspect']),'统计分母':len(data['old_rel']),
              '分母定义':'第一阶段原关系数','比例':len(data['suspect'])/len(data['old_rel']),
              '含义':'疑似并不等于确认错误，全部继续待审'}]
    known=[]
    for k in catalog[2]:
        abilities={m['能力ID'] for m in matrix if m['知识点ID']==k['条目ID'] and m['匹配状态'] in {'直接支撑','部分支撑'}}
        known.append({'知识点ID':k['条目ID'],'知识点名称':k['知识点名称'],'课程条目定位':k['课程条目定位'],
                      '关联候选能力ID':';'.join(sorted(abilities)),'关联能力去重数':len(abilities),
                      '解释':'无34项能力关联不等于课程内容无用；可能关联3项单列课程建议或基础内容'})
    return stats,recruit,known


def render_report(stats,summary,catalog,corpus):
    status=stats['能力匹配状态'];den=stats['三级候选能力数']
    table='\n'.join(f"| {s} | {status.get(s,0)} | {status.get(s,0)}/{den}（{status.get(s,0)/den:.2%}） |" for s in ['直接支撑','部分支撑','未见明确支撑','待人工审核'])
    cases='\n'.join(f"- {s['能力ID']} {s['能力名称']}：{s['最终候选匹配状态']}。{s['课程支撑判断']}" for s in summary if s['能力ID'] in {'N113','N123','N221','N232','N331','N332','N341'})
    return f'''# 第三阶段研究报告

## 一、研究目标与边界

面向高职大数据技术专业《数据库原理及应用》单门课程，形成“岗位候选样本—典型工作任务—职业能力建议—教学项目—知识条目—实操要求”的可追溯数据。本阶段只提交矩阵及结构化候选图谱，不开发网页，不宣称岗位已核验、关系已确认或学生已达到职业能力。

课程依据唯一来自仓库实际PDF《{stats['课程PDF']}》，{stats['PDF物理页数']}个物理页。课程名称来自文件正文，非由文件名推测；自动遍历PDF并依据正文识别，没有假定文件名或位置。课程条目、引用、页码及坐标保留，未用教材常识补内容。条目ID及知识粒度由研究者结构化赋予，非标准自带编码。本文由AI辅助研究形成，全部匹配待人工审核。

## 二、输入与版本核对

基准main提交：`{stats['main提交']}`。PR #2合并状态：{stats['PR2合并状态']}，合并提交记录在phase3/input_snapshot.json。前置研究及PDF共{stats['原文件数']}个源文件在运行前后核对Git blob/SHA256，未修改。输入快照差异数{stats['输入快照差异数']}；若后续输入变化，应重新审阅规则与引文。

实际读取：岗位{stats['岗位数']}条、任务{stats['任务数']}条、原关系{stats['原关系数']}条；第二阶段三级能力{den}项、全层级能力{stats['优化全层级能力数']}项、新旧映射{stats['原新映射行数']}行、人工问题{stats['继承第二阶段问题数']}条、迁移判定{stats['第二阶段迁移判定数']}条。疑似原任务能力关系{stats['疑似原关系数']}条与人工问题清单互相核对一致。没有以26项第一阶段能力代替34项优化库；原能力仅作归档节点及映射端点。

实际文件、工作表、字段、记录数、各字段空值及版本指纹详见“课程标准结构化解析.xlsx”的输入盘点与缺失统计页。空值包含设计上留空的审核人、无正向匹配及无招聘来源等，不能一概判为数据错误。岗位原文摘录缺失{stats['原文摘录缺失岗位数']}条，原职位状态与初评等级完整保留，未升级为已核验。

## 三、课程解析与关键疑点

提取目标{stats['课程目标数']}条（知识、能力、素质分列）、教学项目{stats['教学项目数']}项、项目学生要求{stats['项目要求数']}条；知识条目共{stats['知识条目总数']}条，其中教学项目知识条目{stats['教学项目知识条目数']}条、目标层提及2条、思政层提及1条。知识项命名和细分是研究者对原文的结构化，不是新增教学内容，也不表示课标正式认定这些知识点编号。

实操条目{stats['实操条目总数']}条包括8条可见项目内实操要求和3条被遮挡文本层实训表候选，不能把11条当成独立实训项目数。目标和实操原文跨页时分别保留每个PDF物理页的摘录。共登记{stats['课程证据摘录数']}条证据，其中{stats['遮挡证据摘录数']}条被标记为遮挡候选。

### 学时口径

可见8项目学时依次为4、8、12、12、16、4、4、4，总计64学时，其中理论32、项目内实践32。实践已经在64总数中，不得另加。

第13页文本层还能提取实训一“数据定义”10学时、实训二“数据操作”12学时、实训三“数据查询”10学时及合计32；但是该区域位于后绘制白色矩形下，实际渲染不可见。它只是遮挡文本层候选内容，不能视为发布版可见正文已确认。原PDF未明确说明这32是否对应前述内含实践；既不能相加为96，也不能断言其完全等同于前述32。表内引文保留单元格坐标和遮挡状态，未用于确认34项能力的实训支撑。

可见性核对图见“课程PDF末页可见性核对.png”。脚本用PDF绘制顺序及后绘制不透明白矩形检测遮挡，并保留原PDF不做去遮挡处理。PDF文本可提取不等于发布页面可见。

### 原文问题与考核边界

被遮挡实训表中“INERT INTO”按原字保留并列问题，未静默改成INSERT INTO。第8页DELETE/TRUNCATE的技术概括及部分乱码标题由人工核对，不依据外部知识改写课程标准。

第13页考核章节标题也被遮挡；未提取到可确认的考核组成、比例与评分细则。第3页可见课堂即时反馈和操作检查，作为教学方式条目保存，不冒充正式考核方案。目标层提及DCL与权限，思政提及用户权限、加密与审计，但具体技术实训落点不清，列为人工核验项。

## 四、逐项能力匹配原则

以34项三级能力的可观察工作行为为候选库，分别比较动作、对象、工作产出与课程实践要求。匹配方案显式记录在matching_plan.json，课程条目引用在course_specs.json；运行只校验和重建这些研究判断，未用关键词算法自动定级。

- 直接支撑：可见课程内容及相应项目实操明确对应核心技术行为；仅限文本支撑，不等同培养效果。
- 部分支撑：课程明确涉及某个子环节、基础技术或目标，但缺少完整业务对象、过程或成果。
- 未见明确支撑：本PDF可见内容范围内未找到足够对应证据；不等于课程绝对不涉及，也不能凭缺失新增虚构知识点。
- 待人工审核：表述层次或实施落点不足以可靠判定，如思政提及审计不能直接证明技术检查训练。

同一能力可有多个知识条目，同一知识条目可对应多项能力。矩阵每行保留匹配ID、能力、项目、知识、实践或目标ID、引文、PDF页、理由与审核状态；目标层或未见支撑的项目/知识/实践ID可有意留空，不造占位节点。未见支撑不生成正向图谱支撑边。被遮挡的32学时表不用于直接/部分支撑。

招聘原文明示、岗位摘要支持、研究推断、无招聘证据的课程建议分开。课程PDF原文不是招聘原文；PDF能证实课程表述，不能补足岗位真实性。原109条疑似关系及全部未审关系继续待审，不能因为能力与课程相近而自动确认岗位需求。

## 五、匹配结果与两类证据缺口

矩阵{stats['匹配矩阵行数']}行，按唯一能力ID归并为34项判断：

| 候选状态 | 能力数 | 比例与分母 |
| --- | ---: | --- |
{table}

直接或部分文本支撑共{status.get('直接支撑',0)+status.get('部分支撑',0)}/{den}（{(status.get('直接支撑',0)+status.get('部分支撑',0))/den:.2%}），只能称候选文本关联比例，不是完整能力覆盖率或学生培养达成率。上述分母仅34项建议版职业能力，未混入3项课程建议，也未按可能重复的岗位频次加权。

{cases}

### 岗位需求证据不足

全部{stats['岗位数']}/{stats['岗位数']}岗位仍待核验；已核验招聘原文明示0/{stats['岗位数']}。原任务能力关系中{stats['疑似原关系数']}/{stats['原关系数']}为疑似待审，疑似不等于确认错误。不同量的分母分别是岗位数和原关系数，不能与课程匹配分母混用。摘要支持候选岗位频次仍包含未解决的重复风险，不能解释成全国独立真实需求分布。

### 课程标准支撑不足

{status.get('未见明确支撑',0)}项能力在本PDF范围内未见明确支撑、{status.get('部分支撑',0)}项仅有部分环节、{status.get('待人工审核',0)}项需进一步判读。{status.get('直接支撑',0)}项直接支撑也只说明标准中存在基本技术内容与实操要求。旅游业务情境、实施深度、学生表现及考核证据尚不充分。教学项目知识条目中有{stats['有34能力正向关联的项目知识条目数']}/{stats['教学项目知识条目数']}项建立了34能力库的直接或部分候选关联；未关联条目不意味着无用，它们可能服务基础教学或单列课程建议。

### 反向证据缺口

第二阶段单列的S111关系数据库结构、S112建表改表、S113明确SQL增删改在本PDF确有相应内容；但它们仍无具体招聘证据，不属于34项岗位能力。另表列示而不纳入岗位频次，不能从课程已教反推岗位需要。这类缺口与“岗位有摘要、课程未见明确支撑”方向相反，必须分开处理。

所有课程改进建议只是研究讨论项，不是原课标内容。建议先审定岗位证据和单门课程边界，再决定是否需要增加或调整案例、实操与评价；不因旅游背景需求而把全部经营能力强行纳入数据库课程。

## 六、图谱数据与人工审核

节点{stats['图谱节点数']}个、候选关系{stats['图谱候选关系数']}行、隔离关系{stats['图谱隔离关系数']}行。原能力归档节点与新能力节点独立，旧→新映射可追溯。岗位与任务ID保留原编码，课程条目采用P/K/EP/RP/G/CR/AS前缀，所有节点ID唯一、关系端点有效。合并同端点同类型证据边保留源记录ID并集，不能以边行数作为需求频次。

隔离关系包含第二阶段不建议迁移的匹配以及被遮挡课程条目的结构边，不进入默认候选关系集。候选边也不是已确认事实；任务到能力、能力到课程均为待审研究联系。所有边“已人工确认”为否、不能用于已确认招聘频次。课程匹配边的方向为能力→项目/知识/实操，知识→实操可接续查询；关系名称注明候选语义，不能把一次多跳路径当成已验证培养链。

人工清单继承第二阶段{stats['继承第二阶段问题数']}条未决问题，新增{stats['新增第三阶段问题数']}条课程与匹配问题，并单列{stats['图谱全部待审关系数']}条全部待审关系索引（其中{stats['图谱语义待审关系数']}条为新能力映射或课程匹配等语义联系）。三张清单有交叉，不能简单相加作为独立问题数量。审核人和结论均留空，没有模拟人工审定。

## 七、复现、测试与局限

从仓库根目录运行 `python phase3/build.py`，再运行 `python -m unittest discover -s phase3 -p 'test_*.py' -v`。依赖、路径发现、规则边界及输出说明见phase3/README.md，实际测试结果见测试结果.txt。程序只在独立第三阶段目录写文件；检查原文件指纹、引用原文与页码、编码映射、节点唯一和端点有效，重复运行保证单元格及统计一致，不要求Excel ZIP元数据相同。

本研究只有一个PDF版本；部分内容被覆盖、缺少正式考核细则。职业样本仍为候选，摘要与疑似关系限制了需求结论。知识粒度和34项匹配是AI辅助研究判断，未经教师、行业专家和学习评价验证。输出是可审阅的结构化研究数据，而非已核验行业能力标准或课程培养成效证据。后续先取得完整可见课程标准、补招聘原文、逐条审核109项疑问及34项匹配，再开展教学项目改进和实际评价验证。
'''


def run(input_root=ROOT,output_dir=None,pdf=None):
    root=Path(input_root).resolve();out=Path(output_dir or HERE/'output').resolve()
    protected=[root,root/'output',root/'phase2',root/'scripts',root/'tests']
    if out==root or any(out==p or p in out.parents for p in protected[1:]):
        raise ValueError('禁止覆盖前两阶段文件，请选独立第三阶段目录')
    snap=json.loads((HERE/'input_snapshot.json').read_text())
    before={e['path']:sha(root/e['path']) for e in snap['files']}
    drift=[]
    for e in snap['files']:
        b=(root/e['path']).read_bytes();actual=hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()
        if actual!=e['sha']:
            drift.append(e['path'])
    if drift:
        raise ValueError('源文件与锁定main快照不同，需重新审核：'+str(drift))
    data=load_inputs(root,snap)
    pdf_path,pdf_files=discover_pdf(root,pdf)
    corpus=Corpus(pdf_path)
    specs=json.loads((HERE/'course_specs.json').read_text());plan=json.loads((HERE/'matching_plan.json').read_text())
    catalog=parse_course(corpus,specs)
    matrix,summaries=match_course(data,plan,catalog,corpus)
    graph,suggestions=build_graph(data,catalog,matrix,corpus)
    validate(corpus,data,catalog,matrix,summaries,graph)
    issues=make_issues(data,catalog,matrix,summaries,corpus,graph)
    inv,missing=inventory(root,[x[0] for x in SHEETS],snap['main_commit'])
    coverage,recruit,known=coverage_tables(data,catalog,matrix,summaries,suggestions)
    semantic=[e for e in graph.edges+graph.quarantine if any(x in e['关系类型'] for x in ['候选','建议映射','无招聘证据'])]
    stats={'main提交':snap['main_commit'],'PR2合并状态':'已合并，实际第二阶段文件已核对','PR2合并提交':snap['pr2_merge_commit'],
           '原文件数':len(before),'输入SHA256':before,'输入快照差异数':len(drift),'课程PDF':str(pdf_path.relative_to(root)),
           '自动发现PDF文件':[str(p.relative_to(root)) for p in pdf_files],'课程PDF_SHA256':sha(pdf_path),'PDF物理页数':len(corpus.pages),
           '岗位数':len(data['jobs']),'任务数':len(data['tasks']),'原关系数':len(data['old_rel']),
           '三级候选能力数':len(data['leaf']),'优化全层级能力数':len(data['caps']),'原新映射行数':len(data['mapping']),
           '继承第二阶段问题数':len(data['inherited']),'第二阶段迁移判定数':len(data['proposals']),'疑似原关系数':len(data['suspect']),
           '原文摘录缺失岗位数':sum(not j['证据摘录'] for j in data['jobs']),
           '课程目标数':len(catalog[0]),'教学项目数':len(catalog[1]),'项目要求数':len(catalog[4]),'知识条目总数':len(catalog[2]),
           '教学项目知识条目数':sum(k['课程条目定位']=='教学项目知识条目' for k in catalog[2]),
           '实操条目总数':len(catalog[3]),'课程证据摘录数':len(corpus.evidence),'遮挡证据摘录数':sum(e['是否遮挡']=='是' for e in corpus.evidence),
           '匹配矩阵行数':len(matrix),'能力匹配状态':dict(Counter(s['最终候选匹配状态'] for s in summaries)),
           '图谱节点数':len(graph.nodes),'图谱节点类型':dict(Counter(n['节点类型'] for n in graph.nodes)),
           '图谱候选关系数':len(graph.edges),'图谱隔离关系数':len(graph.quarantine),'图谱语义待审关系数':len(semantic),'图谱全部待审关系数':len(graph.edges)+len(graph.quarantine),
           '新增第三阶段问题数':len(issues),
           '有34能力正向关联的项目知识条目数':sum(k['课程条目定位']=='教学项目知识条目' and k['关联能力去重数']>0 for k in known),
           '可见教学项目学时':catalog[-1][0], '遮挡实践表候选学时':catalog[-1][1],
           '已核验招聘原文明示数':0,'已人工批准图谱关系数':0,
           '方法文件SHA256':{str(p.relative_to(root)):sha(p) for p in [HERE/'build.py',HERE/'course_specs.json',HERE/'matching_plan.json']}}
    notes=[{'事项':'研究状态','说明':'岗位全部待核验，课程匹配及图谱语义联系全部待人工审核；文本支撑不等同实际培养'},
           {'事项':'证据区分','说明':'招聘原文明示本批0；岗位摘要支持、研究推断、无招聘证据课程建议分开。课程PDF原文不能替代招聘原文。'},
           {'事项':'课程范围','说明':'仅取实际PDF；知识点ID和粒度为研究结构化，无通用教材补写'},
           {'事项':'遮挡内容','说明':'第13页实践表及考核标题被白色矩形遮挡；可提取文本层仅作待核验候选，不用于确认实训支撑'},
           {'事项':'学时','说明':'64包含理论32与实践32；遮挡表32单列候选，不得相加为96或重复计数'},
           {'事项':'空值','说明':'未见明确支撑行没有正向引用或节点ID；目标层可无项目/实践ID；不是虚构占位节点。审核字段故意留空'},
           {'事项':'状态分母','说明':'能力状态统计按实际34项候选能力去重，不包含3项课程建议；部分支撑不代表完整覆盖'},
           {'事项':'重复风险','说明':'岗位与任务支撑数仍含未解决的同岗异源风险，不能作为全国独立招聘需求频次'},
           {'事项':'重新运行','说明':'覆盖第三阶段同名文件，人工已填审核版本须另存。前两阶段文件不修改'}]
    out.mkdir(parents=True,exist_ok=True)
    goals,projects,knowledge,practice,requirements,assessment,hours_info=catalog
    write_book(out/'课程标准结构化解析.xlsx',[('课程目标',goals),('教学项目',projects),('知识点',knowledge),('实训任务',practice),
         ('课程能力要求',requirements),('考核相关要求',assessment),('学时口径',hours_info),('PDF原文证据',corpus.evidence),
         ('输入文件工作表盘点',inv),('输入字段缺失统计',missing),('说明',notes)])
    write_book(out/'职业能力课程匹配矩阵.xlsx',[('34项能力汇总',summaries),('多对多匹配明细',matrix),('PDF原文证据',corpus.evidence),('说明',notes)])
    write_book(out/'课程能力覆盖与缺口分析.xlsx',[('能力状态统计',coverage),('逐项支撑与缺口',summaries),('岗位需求证据不足',recruit),
         ('知识条目关联情况',known),('课程有内容但招聘证据不足',suggestions),('说明',notes)])
    write_book(out/'知识图谱节点表.xlsx',[('全部节点',graph.nodes),('节点类型统计',[{'节点类型':k,'节点数':v} for k,v in stats['图谱节点类型'].items()]),('说明',notes)])
    write_book(out/'知识图谱关系表.xlsx',[('候选图谱关系',graph.edges),('隔离关系不作确认',graph.quarantine),('PDF原文证据',corpus.evidence),('说明',notes)])
    write_book(out/'人工审核清单_第三阶段.xlsx',[('新增第三阶段问题',issues),('继承第二阶段未决问题',data['inherited']),('图谱全部待审关系索引',graph.edges+graph.quarantine),('说明',notes)])
    (out/'运行统计.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2))
    (out/'第三阶段研究报告.md').write_text(render_report(stats,summaries,catalog,corpus))
    (out/'PDF逐页文本.json').write_text(json.dumps([{'PDF页码':p['page'],'文本层原始文本':p['raw'],'后绘制白色遮挡矩形':p['masks'],
        '说明':'文本层可能含页面不可见内容；以证据可见性状态区分'} for p in corpus.pages],ensure_ascii=False,indent=2))
    corpus.doc[-1].get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(out/'课程PDF末页可见性核对.png')
    if before!={p:sha(root/p) for p in before}:
        raise RuntimeError('原文件指纹发生变化')
    corpus.doc.close()
    return stats


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-root',type=Path,default=ROOT)
    parser.add_argument('--output',type=Path,default=HERE/'output')
    parser.add_argument('--pdf',type=Path,help='仅在自动发现不能唯一识别时人工指定，不默认文件名')
    args=parser.parse_args()
    print(json.dumps(run(args.input_root,args.output,args.pdf),ensure_ascii=False,indent=2))
