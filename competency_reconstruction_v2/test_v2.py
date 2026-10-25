import hashlib,json,re,tempfile,unittest
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
import build

class V2Tests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory();cls.out=Path(cls.tmp.name);cls.d=build.build(cls.out)
  cls.jobs={j['id']:j for j in cls.d['jobs']};cls.caps={c['id']:c for c in cls.d['capabilities']};cls.tasks={t['id']:t for t in cls.d['tasks']}
 @classmethod
 def tearDownClass(cls):cls.tmp.cleanup()
 def test_sources_required_and_distinct(self):
  self.assertEqual(len(self.jobs),len(self.d['jobs']))
  for j in self.jobs.values():
   for key in ['id','name','publisher','industry','url','collected_at','page_state','decision','reason','duties_summary','requirements_summary','locator']:
    self.assertTrue(j[key],(j['id'],key))
   self.assertTrue(j['url'].startswith('https://'))
  pairs=[(j['publisher'],j['name']) for j in self.jobs.values()];self.assertEqual(len(pairs),len(set(pairs)))
 def test_no_unchecked_sources_in_tasks(self):
  for t in self.tasks.values():
   self.assertTrue(t['evidence'])
   for e in t['evidence']:
    j=self.jobs[e['job_id']];self.assertIn(j['decision'],['纳入','条件纳入']);self.assertNotEqual(j['page_state'],'页面不可访问或正文缺失');self.assertTrue(e['location']);self.assertTrue(e['fragment']);self.assertTrue(j['excerpt'])
 def test_advanced_jobs_excluded_despite_title(self):
  for jid in ['VJ009','VJ010','VJ022']:self.assertEqual(self.jobs[jid]['decision'],'排除')
  linked={e['job_id'] for e in self.d['cross_links']};self.assertTrue(linked.isdisjoint({'VJ009','VJ010','VJ022'}))
 def test_unique_ids_and_acyclic_hierarchy(self):
  h={x['id']:x for x in self.d['hierarchy']};self.assertEqual(len(h),len(self.d['hierarchy']))
  for n in h.values():
   seen=set();cur=n
   while cur['parent_id']:
    self.assertNotIn(cur['id'],seen);seen.add(cur['id']);parent=h[cur['parent_id']];self.assertEqual(parent['level'],cur['level']-1);cur=parent
  for cid in self.caps:self.assertRegex(cid,r'^V2C\d{3}$')
 def test_relation_endpoints_and_separate_evidence_dimensions(self):
  ids=[]
  for e in self.d['cross_links']:
   ids.append(e['id']);self.assertIn(e['job_id'],self.jobs);self.assertIn(e['task_id'],self.tasks);self.assertIn(e['capability_id'],self.caps);self.assertIn(e['abc'],'ABC');self.assertNotIn('support_type',e);self.assertIn('推断',e['nature'])
  for e in self.d['course_links']:
   ids.append(e['id']);self.assertIn(e['capability_id'],self.caps);self.assertIn(e['support_type'],['K1','K2','K3','K4']);self.assertNotIn('abc',e);self.assertIn('推断',e['nature']);self.assertTrue(e['boundary'])
  self.assertEqual(len(ids),len(set(ids)))
 def test_no_duplicate_edges(self):
  for group,keys in [('cross_links',['job_id','task_id','capability_id']),('course_links',['ck_id','capability_id']),('tourism_links',['task_id','capability_id'])]:
   triples=[tuple(e[k] for k in keys) for e in self.d[group]];self.assertEqual(len(triples),len(set(triples)))
 def test_course_evidence_valid_and_visible(self):
  course=build.read_json(build.ROOT/build.COURSE);ce={e['证据ID']:e for e in course['evidence']};ks={k['知识点ID'] for k in course['knowledge']};pr={p['项目ID'] for p in course['projects']}
  for e in self.d['course_links']:
   self.assertIn(e['ck_id'],ks)
   for q in e['source_evidence']:
    self.assertIn(q['id'],ce);self.assertIn(q['project'],pr);self.assertTrue(1<=int(q['page'])<=13);self.assertNotEqual(int(q['page']),13);self.assertEqual(q['quote'],ce[q['id']]['原文短摘录'])
 def test_knowledge_occurrences_not_collapsed(self):
  for e in self.d['course_links']:
   if e['ck_id']=='CK027':self.assertEqual({x['project'] for x in e['source_evidence']},{'P04','P05'})
 def test_capabilities_have_jobs_or_course_basis(self):
  for c in self.caps.values():
   self.assertTrue(c['task_ids'] or c['course_ids']);self.assertTrue(c['definition']);self.assertTrue(c['behaviors']);self.assertTrue(c['expected_product']);self.assertTrue(c['evaluation_basis']);self.assertTrue(c['adaptation']);self.assertIn('候选',c['review_status'])
   if not c['task_ids']:self.assertIn('课程依据',c['origin']);self.assertEqual(c['job_ids'],[])
 def test_tourism_keeps_summary_status(self):
  tour={t['典型工作任务ID'] for t in build.rows(build.ROOT/build.TOURISM,'典型工作任务候选')}
  for e in self.d['tourism_links']:
   self.assertIn(e['task_id'],tour);self.assertIn(e['capability_id'],self.caps);self.assertIn('摘要',e['nature']);self.assertIn('未核验',e['review_status']);self.assertTrue(e['source_evidence'])
   for s in e['source_evidence']:self.assertTrue(s['岗位ID']);self.assertTrue(s['摘要证据文本']);self.assertTrue(s['源单元格'])
 def test_no_false_full_support_or_import_claim(self):
  for e in self.d['course_links']:
   if e['capability_id'] in ['V2C002','V2C006','V2C007'] and e['ck_id']!='CK018':self.assertEqual(e['support_type'],'K4')
  self.assertFalse(any('权限' in c['name'] for c in self.caps.values()))
 def test_summary_counts_computed(self):
  s=self.d['summary'];self.assertEqual(s['collected'],len(self.jobs));self.assertEqual(s['capabilities'],len(self.caps));self.assertEqual(s['cross_abc'],dict(Counter(e['abc'] for e in self.d['cross_links'])));self.assertEqual(sum(s['decisions'].values()),len(self.jobs));self.assertEqual(s['expert_validated'],0)
 def test_excel_json_consistency_and_openable(self):
  cases={'02_跨行业岗位原始资料索引.xlsx':('岗位来源索引','jobs'),'04_跨行业典型工作任务.xlsx':('归纳任务','tasks'),'06_V2三级候选能力清单.xlsx':('候选能力','capabilities'),'07_V2一级二级三级能力体系.xlsx':('自下而上层级','hierarchy'),'08_跨行业任务能力证据关系.xlsx':('岗位任务能力关系','cross_links'),'09_课程知识能力支撑候选关系.xlsx':('候选支撑','course_links'),'10_旅游任务能力应用情境关系.xlsx':('应用情境候选','tourism_links'),'13_待研究者审查事项.xlsx':('待审问题','issues')}
  self.assertEqual(len(list(self.out.glob('*.xlsx'))),13)
  for f in self.out.glob('*.xlsx'):
   w=load_workbook(f,read_only=True);self.assertTrue(w.sheetnames);w.close()
  for file,(sheet,key) in cases.items():
   w=load_workbook(self.out/file,read_only=True);rr=list(w[sheet].values);self.assertEqual(len(rr)-1,len(self.d[key]));self.assertEqual([r[0] for r in rr[1:]],[r['id'] for r in self.d[key]]);w.close()
 def test_old_comparison_correct(self):
  comp={r['旧能力ID']:r for r in self.d['comparison']};self.assertEqual(len(comp),17);self.assertEqual(comp['DB008']['V2概念对应'],['V2C010']);self.assertEqual(comp['DB012']['V2概念对应'],['V2C009']);self.assertEqual(comp['DB015']['V2概念对应'],[])
 def test_rebuild_deterministic_and_committed_json_matches(self):
  with tempfile.TemporaryDirectory() as d:
   other=build.build(Path(d));self.assertEqual(self.d,other)
   for f in self.out.glob('*.xlsx'):
    a=load_workbook(f,read_only=True);b=load_workbook(Path(d)/f.name,read_only=True)
    self.assertEqual(a.sheetnames,b.sheetnames)
    for s in a.sheetnames:self.assertEqual(list(a[s].values),list(b[s].values))
    a.close();b.close()
  self.assertEqual(self.d,build.read_json(build.BASE/'output/v2_research.json'))
 def test_original_inputs_and_app_unchanged(self):
  for path,digest in build.read_json(build.BASE/'protected_inputs.json').items():self.assertEqual(hashlib.sha256((build.ROOT/path).read_bytes()).hexdigest(),digest,path)
 def test_original_input_counts(self):
  s=self.d['summary'];self.assertEqual((s['old_jobs'],s['tourism_tasks'],s['course_knowledge']),(50,48,60));pdf=next(x for x in self.d['inventory'] if x['path'].endswith('.pdf'));self.assertEqual(pdf['pages'],13)

if __name__=='__main__':unittest.main()
