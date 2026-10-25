import unittest,tempfile,json,hashlib
from pathlib import Path
from openpyxl import load_workbook
import build_four as b
class FourTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.t=tempfile.TemporaryDirectory();cls.out=Path(cls.t.name);cls.d=b.run(cls.out)
 @classmethod
 def tearDownClass(cls):cls.t.cleanup()
 def test_classification(self):
  self.assertEqual(set(s['category'] for s in self.d['sources']),{'S1','S2','S3','S4'})
  self.assertFalse(any(s['category']=='S5' for s in self.d['sources']))
 def test_no_unverified_standard_clauses(self):
  for r in self.d['requirements']:self.assertNotIn(r['source_id'],['S3-001','S3-PUB','S4-LEAD'])
  self.assertFalse(any(e['category']=='S3' for e in self.d['relations']))
 def test_references_and_exclusions(self):
  sources={s['id']:s for s in self.d['sources']};raw={r['id']:r for r in self.d['requirements']}
  self.assertEqual(len(raw),len(self.d['requirements']))
  for r in raw.values():self.assertIn(r['source_id'],sources);self.assertTrue(r['locator']);self.assertTrue(r['adaptation_reason']);self.assertTrue(r['review'])
  for e in self.d['relations']:
   self.assertIn(e['requirement_id'],raw);self.assertNotEqual(raw[e['requirement_id']]['decision'],'排除');self.assertIn(e['capability_id'],{r['capability_id'] for r in self.d['matrix']})
 def test_s4_level_and_scope(self):
  for r in self.d['requirements']:
   if r['source_id']=='S4-001' and r['decision']!='排除':self.assertEqual(r['level'],'初级');self.assertIn('PDF实际',r['locator'])
  self.assertFalse(next(r for r in self.d['requirements'] if r['id']=='FR009')['candidate_capabilities'])
 def test_matrix_independence(self):
  self.assertEqual(len(self.d['matrix']),13)
  for r in self.d['matrix']:self.assertEqual(r['independent_employers'],len(set(r['employers'])));self.assertFalse(r['S3'])
 def test_no_duplicate_relations(self):
  es=[(e['requirement_id'],e['capability_id']) for e in self.d['relations']];self.assertEqual(len(es),len(set(es)))
 def test_excel_json(self):
  fs=list(self.out.glob('*.xlsx'));self.assertEqual(len(fs),7)
  for f in fs:
   w=load_workbook(f,read_only=True);self.assertTrue(w.sheetnames);w.close()
  for f,key,sheet in [('01_四源资料索引及核验状态.xlsx','sources','来源索引'),('03_四源原始能力需求条目池.xlsx','requirements','全部条目'),('06_四源能力证据关系.xlsx','relations','候选关系')]:
   w=load_workbook(self.out/f,read_only=True);rr=list(w[sheet].values);self.assertEqual(len(rr)-1,len(self.d[key]));self.assertEqual([r[0] for r in rr[1:]],[r['id'] for r in self.d[key]]);w.close()
 def test_reproducible(self):
  with tempfile.TemporaryDirectory() as t:self.assertEqual(self.d,b.run(t))
  self.assertEqual(self.d,b.load(b.BASE/'output/four_sources.json'))
 def test_protected_baseline(self):
  for f,h in b.load(b.BASE/'baseline_hashes.json').items():self.assertEqual(hashlib.sha256((b.BASE.parent/f).read_bytes()).hexdigest(),h)
if __name__=='__main__':unittest.main()
