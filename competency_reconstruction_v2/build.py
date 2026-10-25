"""离线、可复现的V2候选研究数据构建；不联网补写，不修改旧成果。"""
from pathlib import Path
import json, hashlib, re, argparse
from collections import Counter, defaultdict
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
BASE=Path(__file__).resolve().parent
ROOT=BASE.parent
COURSE='stage4_course_extraction/output/课程教学要素数据库.json'
TOURISM='research_revision/database_competencies_v1/output/02_典型工作任务修订与归纳表.xlsx'
OLD='research_revision/database_competencies_v1/stage3_final_review/output/01_第三阶段规范化三级候选能力体系.xlsx'
JOBS='旅游行业招聘岗位样本库_50条候选_待核验.xlsx'
OLD_LINKS='research_revision/database_competencies_v1/stage3_final_review/output/03_任务能力ABC证据关联复核.xlsx'
MAIN='e65a059ac696ca69fdfa4fb8aadb33f73fd74570'
LABELS={'source_platform':'来源平台域名','education_requirement':'学历要求（保留冲突）','experience_requirement':'经验要求（保留冲突）','id':'编号','name':'名称','publisher':'招聘单位或发布主体','industry':'发布主体行业（研究分类）','application_context':'业务应用领域（与雇主行业分开）','url':'来源URL','page_state':'页面核对状态','decision':'纳入决定','reason':'判断依据','duties_summary':'岗位职责摘要（研究转述）','requirements_summary':'学历经验与任职要求摘要（研究转述）','technology':'数据库或相关技术及出现位置','excerpt':'原文短摘录','locator':'页面章节或段落定位','published_at':'页面发布时间','collected_at':'采集日期','review_status':'研究者审查状态','limitations':'证据限制','object':'工作对象','goal':'工作目标','action':'关键行动（研究归纳）','result':'预期成果（研究分析）','quality':'可检查质量（研究拟定）','result_nature':'成果与质量证据性质','evidence':'岗位证据片段','job_id':'跨行业岗位编号','abc':'岗位证据类别ABC','fragment':'职责或要求片段（转述非原文）','location':'职责或任职要求定位','interpretation':'解释与边界','domain':'一级领域','category':'二级类别','definition':'能力定义','context':'适用工作情境','behaviors':'关键行为表现','knowledge_skills':'所需知识技能','expected_product':'预期工作成果','evaluation_basis':'可观察评价依据（研究拟定）','task_ids':'跨行业任务编号','origin':'能力来源类型','inference':'研究推断及适用边界','adaptation':'高职适配说明','capability_id':'V2能力编号','ck_id':'课程知识编号','support_type':'课程功能K1—K4','boundary':'适用边界','task_id':'典型任务编号','nature':'关系证据性质','parent_id':'父节点编号','level':'层级','source_kind':'来源类别','source_evidence':'来源证据','course_evidence_ids':'课程原文证据编号','page':'PDF实际页码','job_ids':'来源岗位编号','course_ids':'关联CK编号','domain_id':'一级编号','category_id':'二级编号'}
def read_json(path):return json.loads(Path(path).read_text(encoding='utf-8'))
def rows(path,sheet=None):
 w=load_workbook(path,read_only=True,data_only=True);s=w[sheet] if sheet else w.worksheets[0];it=iter(s.values);head=next(it);out=[dict(zip(head,row)) for row in it if any(v is not None for v in row)];w.close();return out

def blob(path):
 b=Path(path).read_bytes();return hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()
def cell(v):
 if v is None:return '暂无记录'
 if isinstance(v,(dict,list)):return json.dumps(v,ensure_ascii=False)
 return v

def xlsx(path,sheets):
 w=Workbook();w.remove(w.active);w.properties.created=w.properties.modified=datetime(2026,10,25)
 for title,data in sheets.items():
  s=w.create_sheet(title[:31]);keys=list(dict.fromkeys(k for row in data for k in row))
  if not keys:keys=['状态'];data=[{'状态':'本轮暂无记录；不是已完成或零需求的结论'}]
  s.append([LABELS.get(k,k) for k in keys])
  for row in data:
   s.append([cell(row.get(k)) for k in keys])
   for c in s[s.max_row]:
    if c.data_type=='f':c.data_type='s'
  s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
  for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='285B52')
  for col in s.columns:
   s.column_dimensions[col[0].column_letter].width=min(65,max(16,len(str(col[0].value))*1.8))
   for c in col:c.alignment=Alignment(vertical='top',wrap_text=True)
 w.save(path)

def build(out=BASE/'output'):
 out=Path(out);out.mkdir(parents=True,exist_ok=True)
 jobs=read_json(BASE/'sources.json');rules=read_json(BASE/'model_rules.json')
 course=read_json(ROOT/COURSE);ks={k['知识点ID']:k for k in course['knowledge']};ce={e['证据ID']:e for e in course['evidence']}
 tour=rows(ROOT/TOURISM,'典型工作任务候选');fragments=rows(ROOT/TOURISM,'任务片段提取');tour_index={t['典型工作任务ID']:t for t in tour}
 original_jobs=rows(ROOT/JOBS,'招聘岗位主库')
 tasks=rules['tasks'];caps=rules['capabilities'];ji={j['id']:j for j in jobs};ti={t['id']:t for t in tasks}
 # No old capability input above this line: hierarchy inferred only from V2 leaves.
 domains=list(dict.fromkeys(c['domain'] for c in caps));categories=list(dict.fromkeys((c['domain'],c['category']) for c in caps))
 ds={d:f'V2D{i:02}' for i,d in enumerate(domains,1)};bs={pair:f'V2B{i:02}' for i,pair in enumerate(categories,1)}
 hierarchy=[dict(id=ds[d],name=d,level=1,parent_id=None,review_status='自下而上归纳；待审') for d in domains]
 hierarchy += [dict(id=bs[(d,b)],name=b,level=2,parent_id=ds[d],review_status='自下而上归纳；待审') for d,b in categories]
 for c in caps:
  c.update(domain_id=ds[c['domain']],category_id=bs[(c['domain'],c['category'])],job_ids=sorted({e['job_id'] for tid in c['task_ids'] for e in ti[tid]['evidence']}),course_ids=[r['ck_id'] for r in rules['course_links'] if r['capability_id']==c['id']])
  hierarchy.append(dict(id=c['id'],name=c['name'],level=3,parent_id=c['category_id'],review_status=c['review_status']))
 cross=[]
 for c in caps:
  for tid in c['task_ids']:
   for e in ti[tid]['evidence']:
    j=ji[e['job_id']]
    cross.append(dict(id=f'VE{len(cross)+1:03}',source_kind='跨行业岗位',job_id=j['id'],task_id=tid,capability_id=c['id'],abc=e['abc'],excerpt=j['excerpt'],locator=e['location']+'；'+j['locator'],url=j['url'],page_state=j['page_state'],fragment=e['fragment'],reason=e['interpretation'],nature=('业务数据任务＋关系型实现研究推断' if e['abc']=='C' else '原文任务/要求经研究归纳与跨体系推断形成候选关系；并非原文明示能力编码'),boundary=c['inference'],review_status='待研究者审查'))
 course_links=[]
 for i,r in enumerate(rules['course_links'],1):
  k=ks[r['ck_id']]
  occ=[o for o in course['occurrences'] if o['知识点ID']==r['ck_id']]
  course_links.append(dict(id=f'VK{i:03}',source_kind='课程知识',**r,ck_name=k['规范化名称'],source_evidence=[dict(id=o['证据ID'],page=o['PDF实际页码'],project=o['项目ID'],quote=o['原文摘录'],state=o['来源状态']) for o in occ],course_evidence_ids=list(dict.fromkeys(o['证据ID'] for o in occ)),page=k['PDF实际页码'],nature='课程原文事实＋跨体系研究推断；非完整培养判定',review_status='待研究者审查'))
 tourism_links=[]
 for i,r in enumerate(rules['tourism_links'],1):
  t=tour_index[r['task_id']];f=[x for x in fragments if x['归纳典型任务ID']==r['task_id']];c=next(x for x in caps if x['id']==r['capability_id'])
  tourism_links.append(dict(id=f'VS{i:03}',source_kind='旅游应用情境',**r,task_name=t['典型工作任务'],job_ids=t['来源岗位ID'],business_context='沿用旅游样本中该任务的业务数据情境；未另造酒店或订单事实',possible_action=c['behaviors'],possible_product=c['expected_product'],source_evidence=[{k:x[k] for k in ['提取任务ID','岗位ID','摘要证据文本','原始招聘链接','原核验状态','源单元格']} for x in f],review_status='待研究者审查；候选岗位保持未核验'))
 # Compare old version AFTER construction; old IDs never seed V2 leaves.
 old=rows(ROOT/OLD)
 mapping={'DB001':[8],'DB002':[7],'DB003':[4],'DB004':[5],'DB005':[3],'DB006':[10],'DB007':[11,12],'DB008':[10],'DB009':[2],'DB010':[2],'DB011':[3],'DB012':[9],'DB013':[7],'DB014':[6],'DB015':[],'DB016':[12],'DB017':[8]}
 # Old labels are read, not assumed; separate conceptual comparison from evidence migration.
 compare=[]
 for o in old:
  oid=next((v for v in o.values() if isinstance(v,str) and re.fullmatch(r'DB\d{3}',v)),None)
  if not oid:continue
  ns=mapping.get(oid,[])
  compare.append({'旧能力ID':oid,'旧记录':o,'V2概念对应':[f'V2C{n:03}' for n in ns],'对照性质':'概念接近，非等价迁移；旧证据不复制','处理说明':'无充分新岗位任务或项目级课程操作依据，保持待补证' if not ns else ('区分模型设计与表结构实现' if oid=='DB007' else '按新任务目标重新归纳，范围需逐项人工确认'),'审查状态':'待审查'})
 issues=[]
 def issue(kind,ids,question):issues.append(dict(id=f'VR{len(issues)+1:03}',category=kind,objects=ids,question=question,review_status='待研究者审查'))
 for j in jobs:
  if j['decision']!='纳入':issue('岗位纳入与真实性',[j['id']],j['reason']+'；'+j['limitations'])
 for c in caps:issue('能力范围与粒度',[c['id']],c['inference'])
 issue('采样边界',[], '22条筛查线索仅8条可用于有限任务分析，未达30—50份参考规模；非概率样本，不能计算行业需求比例。运维行动高度集中单一雇主；需继续补充独立来源。')
 issue('来源分类',[], '雇主行业与业务应用领域分开；本轮医疗、物流岗位多来自软件服务企业，不能宣称各实体行业均有代表性。金融/零售直接雇主覆盖不足。')
 issue('课程边界',['CK057'],'备份与恢复共享知识点，分别限其技术子行为；不把同一个知识点重复计学时。')
 issue('权限边界',[], 'VJ011只在要求中出现基础授权；课程总体目标提到权限但60CK无专门操作项，暂不设正式V2权限能力，保留补证问题。')
 issue('技术适配',['V2C006'],'没有充分导入导出课程原文；INSERT仅提供记录写入环节，不等同批量导入或完整迁移。')
 issue('完整性边界',[], '课程知识未关联不表示冗余；候选能力无课程关系不表示课程缺陷；不进行学生达成率评价。')
 # Duplicate candidates are recorded, not silently deleted.
 duplicates=[]
 for field in ['url','publisher']:
  groups=defaultdict(list)
  for j in jobs:groups[j[field]].append(j['id'])
  duplicates += [dict(field=field,value=k,job_ids=v,nature='疑似重复需按职位及发布时间人工复核；未删除') for k,v in groups.items() if len(v)>1]
 keys=[(j['publisher'].strip().casefold(),j['name'].strip().casefold()) for j in jobs]
 for k,count in Counter(keys).items():
  if count>1:duplicates.append(dict(field='单位+岗位',value=str(k),count=count))
 inventory=[]
 for f in [JOBS,TOURISM,OLD,OLD_LINKS,COURSE,'stage4_competency_graph/output/graph_nodes.json','课标0913.pdf']:
  q=ROOT/f;rec=dict(path=f,git_blob_sha=blob(q),base_main=MAIN)
  if q.suffix=='.xlsx':
   w=load_workbook(q,read_only=True,data_only=True);rec['worksheets']=[dict(name=s.title,rows=len(list(s.values))-1,fields=list(next(s.values))) for s in w];w.close()
  elif q.suffix=='.json':
   d=read_json(q);rec['collections']={k:len(v) for k,v in d.items() if isinstance(v,list)} if isinstance(d,dict) else {'nodes':len(d)}
  else:
   import fitz
   with fitz.open(q) as pdf:rec['pages']=len(pdf)
  inventory.append(rec)
 summary=dict(base_main=MAIN,collected=len(jobs),page_checked=sum(j['page_state']!='页面不可访问或正文缺失' for j in jobs),original_page_checked=sum(j['page_state']=='原始发布页可见文本已核对' for j in jobs),secondary_page_checked=sum(j['page_state']=='二次发布页可见文本已核对' for j in jobs),decisions=dict(Counter(j['decision'] for j in jobs)),tasks=len(tasks),domains=len(domains),categories=len(categories),capabilities=len(caps),course_only=sum(not c['task_ids'] for c in caps),cross_links=len(cross),cross_abc=dict(Counter(e['abc'] for e in cross)),course_links=len(course_links),tourism_links=len(tourism_links),issues=len(issues),old_jobs=len(original_jobs),tourism_tasks=len(tour),course_knowledge=len(ks),expert_validated=0)
 data=dict(summary=summary,jobs=jobs,tasks=tasks,capabilities=caps,hierarchy=hierarchy,cross_links=cross,course_links=course_links,tourism_links=tourism_links,issues=issues,inventory=inventory,comparison=compare,duplicate_candidates=duplicates)
 (out/'v2_research.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 plan=[dict(行业=h,优先方向=d,采样方法='目的性采样，先开原页后判断；30—50份为参考非配额',检索表达=q,待补充缺口=gap) for h,d,q,gap in [('医疗信息化','实施助理','招聘 实施工程师 大专 SQL 数据库 医疗','独立SQL工作成果'),('物流/电商','WMS实施与数据助理','物流 SQL 实施 大专 招聘','初级真实岗位，不混入中高级'),('制造业','数据维护与ERP支持','招聘 SQL 数据统计 应届 制造','更多直接雇主专科岗位'),('教育信息化','初级数据支持','教育 SQL 实施 大专 招聘','技术深度适配与本科门槛'),('软件服务','基础DBA与数据质量','数据库 备份 恢复 应届 大专 招聘','多雇主备份恢复与质量任务'),('金融支持','数据核对助理','SQL 数据核对 招聘 大专','本轮未形成有效样本，继续补采'),('零售','基础数据处理','SQL 数据处理专员 经验不限','本轮缺少直接零售雇主样本')]]
 xlsx(out/'01_跨行业岗位采样方案.xlsx',{'采样方案':plan,'实际输入版本':inventory})
 xlsx(out/'02_跨行业岗位原始资料索引.xlsx',{'岗位来源索引':jobs,'疑似重复记录':duplicates,'缺失字段统计':[dict(field=k,missing=sum(j.get(k) in (None,'') for j in jobs),denominator=len(jobs)) for k in jobs[0]]})
 xlsx(out/'03_跨行业岗位纳入排除记录.xlsx',{'逐条决定':[{k:j[k] for k in ['id','name','industry','application_context','decision','reason','requirements_summary','page_state','limitations','review_status']} for j in jobs]})
 xlsx(out/'04_跨行业典型工作任务.xlsx',{'归纳任务':[{k:v for k,v in t.items() if k!='evidence'} for t in tasks],'来源片段':[dict(task_id=t['id'],**e,url=ji[e['job_id']]['url'],excerpt=ji[e['job_id']]['excerpt'],page_state=ji[e['job_id']]['page_state']) for t in tasks for e in t['evidence']]})
 xlsx(out/'05_职业行动与工作成果分析.xlsx',{'行动分析':[{k:v for k,v in t.items() if k!='evidence'} for t in tasks]})
 xlsx(out/'06_V2三级候选能力清单.xlsx',{'候选能力':caps})
 xlsx(out/'07_V2一级二级三级能力体系.xlsx',{'自下而上层级':hierarchy})
 xlsx(out/'08_跨行业任务能力证据关系.xlsx',{'岗位任务能力关系':cross})
 xlsx(out/'09_课程知识能力支撑候选关系.xlsx',{'候选支撑':course_links,'未关联CK':[k for k in ks.values() if k['知识点ID'] not in {e['ck_id'] for e in course_links}]})
 xlsx(out/'10_旅游任务能力应用情境关系.xlsx',{'应用情境候选':tourism_links,'全部48任务处置':[dict(**t,V2候选关联=[e['id'] for e in tourism_links if e['task_id']==t['典型工作任务ID']],处置='有候选关联' if any(e['task_id']==t['典型工作任务ID'] for e in tourism_links) else '暂无具体关联；未强制补线') for t in tour]})
 evidence_summary=[]
 for c in caps:
  es=[e for e in cross if e['capability_id']==c['id']]
  evidence_summary.append(dict(capability_id=c['id'],name=c['name'],origin=c['origin'],job_ids=c['job_ids'],cross_abc=dict(Counter(e['abc'] for e in es)),course_ids=c['course_ids'],旅游任务=[e['task_id'] for e in tourism_links if e['capability_id']==c['id']],计数说明='记录数量不是行业需求比例；ABC不等于核验状态，K不等于强度分值',review_status=c['review_status']))
 xlsx(out/'11_V2能力证据汇总.xlsx',{'三类来源分开':evidence_summary,'运行统计':[dict(指标=k,值=v) for k,v in summary.items()]})
 xlsx(out/'12_新旧能力体系对照.xlsx',{'旧17项概念对照':compare,'V2独立来源':caps})
 xlsx(out/'13_待研究者审查事项.xlsx',{'待审问题':issues})
 methods=f'''# V2专业能力体系建模说明

本轮为独立候选稿，未替换旧版平台。基线main `{MAIN}`。实际输入路径、工作表、字段、数量及Git blob校验值见01工作簿及v2_research.json.inventory。课程原件为根目录课标0913.pdf，共13页；使用既有60CK及其原文证据。本轮未重新解析或修改课程标准。页13遮挡内容不纳入正向证据。课程64学时=理论32+实践32；两处实践32同源，不累计。

## 方法与停止条件

先阅读新招聘资料，将职责与技能要求分开；按工作对象、目标、行动和成果归纳10项任务；再提炼{len(caps)}项候选能力，最后自下而上归纳{len(domains)}个领域、{len(categories)}个类别。旧17项仅在建模完成后用于概念对照，不迁移旧关系。model_rules.json为逐项人工可编辑的AI研究判断，不是自动关键词分类，也未经专家审定。

本轮22条筛查记录中，8条纳入或条件纳入资料能支持有限任务分析，因此继续建立小规模候选稿；不是以8条推断样本充分。只有两条无附加纳入条件，其余需审查。跨行业任务、课程知识和旅游场景分别建表，不相互替代。若后续审查否定这些有效证据，需重新评估建模依据，不能用课程填补伪装岗位需求。30—50有效样本参考目标尚未达到，不能称本轮采样充分或全国代表性。

岗位原页核对仅指读取公开可见职责/要求，不是对招聘单位真实性、招聘有效期或实际任职行为的独立认证。搜索摘要/打不开页面不进入任务关系；本科及经验要求原样保留。行业列表示雇主分类，应用领域另列，避免将软件供应商当医疗或物流实体企业。

## 能力粒度与高职适配

V2C008明细取数与V2C009指标汇总因目标和输出不同分开，均可源自VT008，不重复计岗位。连接、子查询、聚合是方法与行为，不因语法单独扩节点。V2C004备份与V2C005恢复目标及成果不同；同源职责拆解不是两个岗位。V2C002合并基础巡检与常见异常协助，不宣称独立复杂根因诊断。V2C011—013仅为课程依据支持：模型设计、表与约束实现、查询视图，各以可检查成果定义，不能统计为招聘需求。视图是否独立能力仍待审。

V2C006迁入迁出、V2C007质量处理的关系型实现是技术解释；V2C010只有录入任务证据，SQL更新删除和事务来自课程补充。成果、核查指标和单实例教学深度均为研究拟定，不是招聘方原文验收细则或学生评分表。本科岗位技术子活动可参考，不意味着高职学生符合该招聘学历门槛。VJ011的集群/高可用仅作为超范围说明，不纳入基础能力行为。

## 证据维度与课程系统性

A表示任务明确数据库技术活动；B表示数据库相关但操作/技术类型需解释；C表示业务数据任务采用关系数据库为研究推断。ABC与页面核对状态分开，来源A/B/C（三类资料角色）不能与岗位ABC证据等级混用。

K1基础知识、K2直接技术子操作、K3综合应用活动、K4条件性子支撑。本轮课程关系全部为跨体系研究推断；CK原文不等于完整能力明示或已经培养。知识实体与项目出现记录分开（CK027保留P04与P05语境）。WHERE不等于完整质量核查，INSERT不等于导入，创建索引不等于性能改善，外键不作为JOIN必要前提。CK057共享知识点分别引用备份和恢复范围，不重复学时。没有独立权限配置活动，不凭总体目标增补权限实践。

所有60CK均保留，未关联项目单独导出；没有岗位关联的基础课程内容不判为冗余。未建立关系也不判为课程缺陷。本轮没有能力得分、权重、行业需求覆盖率或学生达成率。

## 旅游特色与新旧对照

50岗位仍为候选摘要，48任务逐条保留处置结果。仅11条候选应用关联，摘要原句、岗位ID和来源单元格可追溯；数据库具体行动与成果属研究解释。没有虚构酒店订单等对象。权限等无充分依据的不强行补线。新旧对照是概念范围映射，非等价替换和证据迁移；旧数据一字未改。

## 重复检查、复现与局限

检查同URL、同发布主体、单位+岗位重复，疑似项只标记不删除。缺失日期保持空值，Excel显示暂无记录。采样是目的性、小规模、网络可达性偏倚明显；运维任务集中VJ011，技术范围和采样外推均须专家审查。缺少更多独立初级雇主、零售金融直接行业及权限等操作依据。生成器只读取冻结JSON和历史文件，不自动联网重写，事实摘要、研究判断和生成结果分开存放。补采应编辑sources.json并补齐原页核对记录，修改判断须编辑model_rules.json后运行脚本与测试。

V2是一套待研究者审查的候选专业能力体系，不代表已经完成行业专家验证，也不代表学生能力达成度评价。
'''
 (out/'V2专业能力体系建模说明.md').write_text(methods,encoding='utf-8')
 (out/'跨行业岗位采集与证据限制说明.md').write_text(f'''# 跨行业岗位采集与证据限制说明

采集日期2026-10-25。共{len(jobs)}条记录，{summary['original_page_checked']}条原始发布页可见文本核对，{summary['secondary_page_checked']}条高校等二次发布页核对，5条原页失败或正文缺失。页面已打开不等于原始企业直发；未联系发布方，未确认目前仍招。

决定分布：{summary['decisions']}。有效分析子集8条（2纳入、6条件纳入），其中原始发布页6条、二次发布页2条。未达到30—50有效样本参考数量。排除4条，其中3条因技术深度/学历经验定位不适配，1条缺数据库活动。其余5条背景保留不参加任务提取，5条待核验不参加任务提取。

检索覆盖医疗、物流、ERP、教育、制造、数据处理、初级运维方向；原始URL和岗位段落见02表，失败状态均保留。VJ008连续超时/503，VJ015两次超时，VJ016工具不可访问，VJ017仅返回输入框，VJ018为404。搜索片段未拼接为招聘原文。额外检索到的聚合列表、只读标题或无法确定主体的信息未生成正式样本。未登录或绕过访问控制。只保存必要短摘录和转述，不存整份招聘广告。

网页日期和搜索日期可能不同（VJ011）；VJ020学历标题与正文矛盾；VJ012单位简介及应届/经验条件矛盾。所有问题保留。原页14条统计包含被排除/背景岗位，不可称14份有效初级样本。校招本科的知识技术可参考，但不证明高职学历满足门槛。

后续优先补采多雇主的基础查询交付、数据维护、备份恢复、权限配置与结构设计职责；目前没有足够证据宣称能力覆盖完整。本次不是全国行业需求抽样估计，岗位数、关系数不得解释为需求占比。公开招聘文本真实性核对与候选能力专家验证是两件事，后者尚未进行。
''',encoding='utf-8')
 return data
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,default=BASE/'output');args=ap.parse_args();print(json.dumps(build(args.output)['summary'],ensure_ascii=False,indent=2))
