"""第四阶段第二轮：显式技术规则、继承关系及可审查多跳路径。"""
from pathlib import Path
from collections import Counter
import json,hashlib,importlib.util,tempfile,argparse
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent;ROOT=HERE.parent
spec=importlib.util.spec_from_file_location('course_helpers',ROOT/'stage4_course_extraction/build.py');helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
TYPES={'R01':('CONTAINS_COMPETENCY',None,None),'R02':('REQUIRES_COMPETENCY_CANDIDATE','WORK_TASK','COMPETENCY_L3'),'R03':('SUPPORTS_COMPETENCY_CANDIDATE','COURSE_KS','COMPETENCY_L3'),'R04':('APPEARS_IN_PROJECT','COURSE_KS','TEACHING_PROJECT'),'R05':('HAS_TEACHING_ACTIVITY','TEACHING_PROJECT','TEACHING_ACTIVITY'),'R06':('CONTENT_RELATED_TO_PRACTICE','TEACHING_PROJECT','PRACTICAL_TASK'),'R07':('MAY_PRODUCE_OUTPUT','TEACHING_ACTIVITY','LEARNING_OUTPUT'),'R08':('HAS_CHECK_EVIDENCE','TEACHING_ACTIVITY','ASSESSMENT_EVIDENCE'),'R09':('COMPETENCY_LOGIC_CANDIDATE','COMPETENCY_L3','COMPETENCY_L3')}
def read(p,s):
 w=load_workbook(p,read_only=True);it=iter(w[s].values);h=next(it);rows=[dict(zip(h,r)) for r in it if any(x is not None for x in r)];w.close();return rows

def run(out=None):
 snap=json.loads((HERE/'input_snapshot.json').read_text())
 for f in snap['files']:
  if helper.gitsha(ROOT/f['path'])!=f['sha']:raise ValueError('输入偏离main：'+f['path'])
 s3=next(ROOT.glob('research_revision/**/stage3_final_review/output'));cp=ROOT/'stage4_course_extraction/output';c=json.loads((cp/'课程教学要素数据库.json').read_text())
 caps=read(s3/'01_第三阶段规范化三级候选能力体系.xlsx','三级候选能力');hier=read(s3/'01_第三阶段规范化三级候选能力体系.xlsx','完整能力层级');jobs=read(s3/'03_任务能力ABC证据关联复核.xlsx','逐条关联复核');hs=read(s3/'04_三级能力层级包含关系.xlsx','层级包含关系');logic=read(s3/'05_三级能力非层级逻辑关系.xlsx','逐条关系判断')
 tasks=read(s3.parent.parent/'round2_abc/output/01_岗位任务ABC重分类与证据表.xlsx','典型任务类别集合')
 inventory=[]
 for folder in [s3,cp]:
  for p in sorted(folder.glob('*.xlsx')):
   w=load_workbook(p,read_only=True)
   for sheet in w:
    rows=list(sheet.values);inventory.append({'文件':str(p.relative_to(ROOT)),'工作表':sheet.title,'记录数':len(rows)-1,'字段':';'.join(str(x) for x in rows[0]),'来源main':snap['main_commit']})
   w.close()
 # 对照机器JSON与实际Excel的关键实体ID，差异停止，不默填。
 for table,file,sheet,key in [('knowledge','03_课程知识技能点规范化表.xlsx','规范化知识技能点','知识点ID'),('occurrences','03_课程知识技能点规范化表.xlsx','项目出现记录','出现记录ID'),('activities','04_教学活动与学习成果表.xlsx','教学活动','教学活动ID'),('outcomes','04_教学活动与学习成果表.xlsx','学习成果与潜在产物','学习成果ID'),('projects','02_八个教学项目结构化表.xlsx','教学项目','项目ID'),('practice','05_三项独立实训任务表.xlsx','独立实训任务','实训ID'),('checks','06_课程评价与检查证据表.xlsx','检查反馈与完成要求','评价检查ID')]:
  source=read(cp/file,sheet)
  if len(source)!=len(c[table]) or {x[key] for x in source}!={x[key] for x in c[table]}:raise ValueError('Excel/JSON实体不一致：'+table)
  keyed={x[key]:x for x in c[table]}
  for row in source:
   for field,value in row.items():
    if (value or '')!=(keyed[row[key]].get(field) or ''):raise ValueError('Excel/JSON字段不一致：'+table+':'+field)
 # 方法及审查原文只作为方法/限制输入，不从Markdown生成实体。
 for p in [s3/'06_第三阶段研究方法与划分规则.md',cp/'09_第四阶段第一轮研究方法说明.md',cp/'12_第四阶段第一轮人工审查修订记录.md']:p.read_text()
 old_questions=read(s3/'07_第三阶段研究者人工审查清单.xlsx','研究者审查问题')
 if len(caps)!=17 or len(c['knowledge'])!=60 or len(jobs)!=65:raise ValueError('关键输入数量不同，停止候选生成')
 nodes=[];edges=[];reviews=[];evidence=[];paths=[]
 def nid(t,id):return t+':'+id
 def node(t,id,name,desc,file,stage,state):
  nodes.append({'node_id':nid(t,id),'original_id':id,'node_type':t,'name':name,'description':desc,'source_stage':stage,'source_file':file,'evidence_state':state,'review_status':'待研究者审查'})
 s3file=str((s3/'01_第三阶段规范化三级候选能力体系.xlsx').relative_to(ROOT));coursefile=str((cp/'课程教学要素数据库.json').relative_to(ROOT))
 capby={r['能力编号']:r for r in caps};lev={'一级':'COMPETENCY_L1','二级':'COMPETENCY_L2','三级':'COMPETENCY_L3'}
 for h in hier:node(lev[h['能力层级']],h['节点ID'],h['名称'],capby[h['节点ID']]['能力定义（研究拟定）'] if h['节点ID'] in capby else '从三级候选归纳的上位能力领域',s3file,'第三阶段',h['研究状态'])
 for t in tasks:node('WORK_TASK',t['典型工作任务ID'],t['典型工作任务'],'候选岗位摘要归纳任务，原文未核验',str((s3.parent.parent/'round2_abc/output/01_岗位任务ABC重分类与证据表.xlsx').relative_to(ROOT)),'第二/三阶段','候选任务，未核验')
 specs=[('knowledge','COURSE_KS','知识点ID','规范化名称'),('projects','TEACHING_PROJECT','项目ID','项目原文名称'),('activities','TEACHING_ACTIVITY','教学活动ID','技术活动或教学方式名称'),('practice','PRACTICAL_TASK','实训ID','实训原文名称'),('outcomes','LEARNING_OUTPUT','学习成果ID','学习成果或产物'),('checks','ASSESSMENT_EVIDENCE','评价检查ID','类别')]
 for table,typ,key,name in specs:
  for r in c[table]:
   state=r.get('成果证据性质',r.get('内容来源状态',r.get('来源状态','候选课程要素')))
   node(typ,r[key],r[name],r.get('证据边界',r.get('正式提交物要求',r.get('内容性质',r[name]))),coursefile,'第四阶段第一轮',state)
 ce={r['证据ID']:r for r in c['evidence']}
 for id,r in ce.items():evidence.append({'evidence_id':'COURSE:'+id,'system':'课程PDF','source_record_id':id,'source_file':'课标0913.pdf','page':r['PDF实际页码'],'quote':r['原文短摘录'],'state':r['可信状态'],'source_attributes':r})
 for r in jobs:evidence.append({'evidence_id':'JOB:'+r['关联ID'],'system':'岗位摘要','source_record_id':r['关联ID'],'source_file':r['源文件'],'page':None,'quote':r['摘要证据文本'],'state':r['来源状态'],'source_attributes':r})
 def edge(id,rt,a,b,why,refs='',origin='SOURCE_DERIVED',sub='',uncertainty='',active=True,attrs=None):
  ids=['COURSE:'+x for x in refs.split(';') if x] if isinstance(refs,str) else refs
  pages=sorted({ce[x.split(':',1)[1]]['PDF实际页码'] for x in ids if x.startswith('COURSE:')})
  row={'edge_id':id,'source':a,'target':b,'relation_type':TYPES[rt][0],'relation_code':rt,'support_function_or_subtype':sub,'source_stage':'第四阶段第二轮' if rt=='R03' else '继承第三/四阶段第一轮','source_file':coursefile if rt in ['R03','R04','R05','R06','R07','R08'] else str((s3/({'R01':'04_三级能力层级包含关系.xlsx','R02':'03_任务能力ABC证据关联复核.xlsx','R09':'05_三级能力非层级逻辑关系.xlsx'}[rt])).relative_to(ROOT)),'evidence_ids':ids,'pdf_pages':pages,'quote_reference':';'.join(ids),'reason':why,'source_nature':origin,'research_inference':origin=='RESEARCH_INFERENCE','review_status':'候选、待人工审查','uncertainty':uncertainty,'default_visible':active,'source_attributes':attrs or {}}
  row['edge_status']='新增候选' if rt=='R03' else '继承候选关联'
  if rt=='R09':
   row['edge_status']=attrs['审查建议'];row['review_status']=attrs['人工审查状态']+'；处理：'+attrs['审查建议']
  edges.append(row);return row
 hid={n['original_id']:n['node_id'] for n in nodes if n['node_type'].startswith('COMPETENCY')}
 for r in hs:edge('R01:'+r['关系编号'],'R01',hid[r['源能力编号']],hid[r['目标能力编号']],r['判断依据'],attrs=r)
 for r in jobs:edge('R02:'+r['关联ID'],'R02',nid('WORK_TASK',r['典型工作任务ID']),hid[r['能力编号']],r['判断或推断理由'],['JOB:'+r['关联ID']],origin='SOURCE_DERIVED',sub=r['关联证据类别'],uncertainty=r['采用条件与替代解释'],active=r['关联证据类别']=='A',attrs=r)
 for r in c['occurrences']:edge('R04:'+r['出现记录ID'],'R04',nid('COURSE_KS',r['知识点ID']),nid('TEACHING_PROJECT',r['项目ID']),'复用项目出现记录与原文语境',r['证据ID'],attrs=r)
 for r in c['activities']:edge('R05:'+r['教学活动ID'],'R05',nid('TEACHING_PROJECT',r['项目ID']),nid('TEACHING_ACTIVITY',r['教学活动ID']),'原文项目内安排教学活动；非学生实际完成',r['证据ID'],origin='SOURCE_EXPLICIT',attrs=r)
 for i,r in enumerate(c['practice_project'],1):
  pt=next(x for x in c['practice'] if x['实训ID']==r['实训ID']);edge(f'R06:{i}','R06',nid('TEACHING_PROJECT',r['项目ID']),nid('PRACTICAL_TASK',r['实训ID']),r['关联理由'],pt['证据ID'],origin='RESEARCH_INFERENCE',uncertainty='内容相关，精确学时对应未确认；末页隐藏文本不得作确认正向证据',active=False,attrs=r)
 for key,rt,targettype,targetkey,lookup in [('activity_outcome','R07','LEARNING_OUTPUT','学习成果ID','outcomes'),('activity_check','R08','ASSESSMENT_EVIDENCE','评价检查ID','checks')]:
  for i,r in enumerate(c[key],1):
   item=next(x for x in c[lookup] if x[targetkey]==r[targetkey]);edge(f'{rt}:{i}',rt,nid('TEACHING_ACTIVITY',r['教学活动ID']),nid(targettype,r[targetkey]),r['关联性质'],item['证据ID'],origin='RESEARCH_INFERENCE',uncertainty=item.get('正式提交物要求',item.get('证据边界',''))+'；不是正式评分绑定',attrs=r)
 for r in logic:
  edge('R09:'+r['关系编号'],'R09',hid[r['源能力编号']],hid[r['目标能力编号']],r['判断依据'],origin='RESEARCH_INFERENCE',sub=r['关系类型'],uncertainty=r['适用条件'],active=r['审查建议']=='保留候选',attrs=r)
 ks={r['知识点ID']:r for r in c['knowledge']}
 rules=json.loads((HERE/'support_rules.json').read_text())
 for i,r in enumerate(rules,1):
  ck=ks[r['knowledge']];cap=capby[r['capability']];contexts=(["P04"] if r['capability']=='DB006' and r['knowledge']=='CK027' else ["P05"] if r['capability'] in ['DB001','DB013'] and r['knowledge']=='CK027' else ck['对应项目ID'].split(';'));refs=';'.join(x['证据ID'] for x in c['occurrences'] if x['知识点ID']==r['knowledge'] and x['项目ID'] in contexts);attrs={'knowledge_id':r['knowledge'],'competency_id':r['capability'],'knowledge_name':ck['规范化名称'],'competency_name':cap['三级能力名称'],'competency_behavior':cap['可观察行为（研究拟定）'],'interface_id':r.get('interface',''),'technical_behavior':r['behavior'],'course_context':';'.join(contexts)}
  if r.get('interface'):
   it=next(x for x in c['technical_behavior_interfaces'] if x['接口ID']==r['interface']);refs+=';'+it['证据ID'];attrs['interface_source']=it
  edge(f'R03:{i:03}','R03',nid('COURSE_KS',r['knowledge']),hid[r['capability']],r['reason'],refs,origin='RESEARCH_INFERENCE',sub=r['function'],uncertainty=r['limit'],attrs=attrs)
 support=[r for r in edges if r['relation_code']=='R03']
 # 每步保留边ID；路径末端包含实训、检查和产物，但不转为直接培养边。
 for s in support:
  for occurrence in [x for x in edges if x['relation_code']=='R04' and x['source']==s['source'] and x['target'].split(':',1)[1] in s['source_attributes']['course_context'].split(';')]:
   prefix=[s['edge_id'],occurrence['edge_id']];project=occurrence['target']
   def path(step,endpoint):
    chain=prefix+step;rs=[x for x in edges if x['edge_id'] in chain];paths.append({'path_id':f'PATH{len(paths)+1:05}','competency_id':s['target'],'knowledge_id':s['source'],'project_id':project,'endpoint':endpoint,'edge_path':chain,'evidence_ids':sorted({id for r in rs for id in r['evidence_ids']}),'nature':'基于知识点路径推导的候选关联，非直接原文能力培养','uncertainty':'项目内活动不保证针对本能力；实训隐藏文本、评分及提交限制沿用原记录','default_visible':all(r['default_visible'] for r in rs)})
   path([],project)
   for a in [x for x in edges if x['relation_code'] in ['R05','R06'] and x['source']==project]:
    path([a['edge_id']],a['target'])
    for z in [x for x in edges if x['relation_code'] in ['R07','R08'] and x['source']==a['target']]:path([a['edge_id'],z['edge_id']],z['target'])
 coverage=[]
 for id,k in ks.items():
  linked=[r for r in support if r['source']==nid('COURSE_KS',id)];coverage.append({'知识点ID':id,'名称':k['规范化名称'],'关联能力':[r['target'] for r in linked],'关系ID':[r['edge_id'] for r in linked],'状态':'有候选关联，待审' if linked else '当前未建立三级能力关联；不等于冗余或应删除'})
 overview=[]
 for id,cap in capby.items():
  ps=[p for p in paths if p['competency_id']==hid[id]];ss=[r for r in support if r['target']==hid[id]]
  overview.append({'能力ID':id,'名称':cap['三级能力名称'],'支撑候选关系ID':[r['edge_id'] for r in ss],'教学项目':sorted({p['project_id'] for p in ps}),'可追溯活动实训成果检查':sorted({p['endpoint'] for p in ps}),'路径ID':[p['path_id'] for p in ps],'状态':'候选技术关系，不是正式课程支持等级' if ss else '未建立候选支撑：原文技术对象/操作依据不足，非课程缺陷结论'})
 for r in support:reviews.append({'review_id':'Q:'+r['edge_id'],'object_id':r['edge_id'],'issue':'核对知识粒度与完整能力行为，审查支撑功能及条件','basis':r['reason'],'restriction':r['uncertainty'],'state':'待研究者审查'})
 for r in logic:
  if r['审查建议']!='保留候选':reviews.append({'review_id':'Q:'+r['关系编号'],'object_id':'R09:'+r['关系编号'],'issue':r['审查建议'],'basis':r['处理理由'],'restriction':'不进入默认有效候选边','state':r['人工审查状态']})
 for r in c['issues']:reviews.append({'review_id':'COURSE:'+r['问题ID'],'object_id':r['对象'],'issue':r['问题'],'basis':r['当前处理与所需审查'],'restriction':'继承课程第一轮限制，不自动关闭','state':r['审查状态']})
 for id in ['DB009','DB010','DB014','DB015']:reviews.append({'review_id':'NO:'+id,'object_id':id,'issue':'未建立课程知识支撑候选','basis':'一般环境部署不等于数据库运行检查/异常诊断；INSERT不等于数据导入；目标提及权限不等于具体配置技能','restriction':'可后续审查目标层证据，但不强制连线','state':'待研究者审查'})
 for r in old_questions:reviews.append({'review_id':'STAGE3:'+r['问题编号'],'object_id':r['涉及对象'],'issue':r['需研究者判断的问题'],'basis':r['可用依据'],'restriction':'继承第三阶段未决问题','state':r['审查状态']})
 bounds=json.loads((HERE/'boundary_cases.json').read_text())
 out=Path(out or HERE/'output').resolve()
 if out.is_relative_to(ROOT) and not out.is_relative_to(HERE):raise ValueError('禁止覆盖历史目录')
 out.mkdir(parents=True,exist_ok=True)
 tables={'01_能力图谱节点总表.xlsx':{'节点':nodes,'输入清单':inventory},'02_能力图谱关系总表.xlsx':{'全部关系记录':edges,'默认候选边':[e for e in edges if e['default_visible']],'默认隐藏':[e for e in edges if not e['default_visible']]},'03_三级能力与课程知识技能支撑关系.xlsx':{'候选支撑':support},'04_能力与教学项目活动关联路径.xlsx':{'多跳候选路径':paths},'05_课程知识技能点覆盖与未关联分析.xlsx':{'知识点关联概览':coverage},'06_三级能力课程支撑概览.xlsx':{'能力概览非等级评价':overview},'07_图谱关系证据索引.xlsx':{'证据':evidence,'关系证据引用':[{'关系ID':e['edge_id'],'证据ID':id} for e in edges for id in e['evidence_ids']]},'08_图谱关系人工审查清单.xlsx':{'人工审查':reviews},'12_能力图谱候选关系边界案例.xlsx':{'边界案例':bounds}}
 for name,sheets in tables.items():helper.write_book(out/name,sheets)
 summary={'main':snap['main_commit'],'nodes':len(nodes),'node_types':dict(Counter(n['node_type'] for n in nodes)),'edge_records':len(edges),'edge_types':dict(Counter(r['relation_code'] for r in edges)),'default_edges':sum(e['default_visible'] for e in edges),'support_edges':len(support),'linked_competencies':sum(bool(r['支撑候选关系ID']) for r in overview),'linked_knowledge':sum(bool(r['关系ID']) for r in coverage),'unlinked_knowledge':[r['知识点ID'] for r in coverage if not r['关系ID']],'evidence':len(evidence),'reviews':len(reviews),'paths':len(paths),'hours':c['summary']['学时']}
 for name,data in [('graph_nodes',nodes),('graph_edges',edges),('graph_evidence',evidence),('graph_review_items',reviews),('graph_paths',paths),('summary',summary)]:
  (out/(name+'.json')).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
 for src,dest in [('方法说明.md','09_第四阶段第二轮图谱构建方法说明.md'),('数据字典.md','11_图谱数据字典与关系模式.md')]: (out/dest).write_text((HERE/src).read_text())
 (out/'10_第四阶段第二轮执行总结.md').write_text('# 第四阶段第二轮执行总结\n\n实际main及各类数量：\n\n```json\n'+json.dumps(summary,ensure_ascii=False,indent=2)+'\n```\n\n第三阶段实际路径为research_revision/database_competencies_v1/stage3_final_review/output，并非仓库根目录。输入清单详见01。复用课程JSON与Excel，未从Markdown生成数据。所有R03为研究推断、待人工审查；R05记录原文活动，其他继承与路径按来源保留。R09共8判断，默认纳入5保留候选，1待复核及2排除隐藏。B/C岗位边与隐藏文本实训边默认隐藏。未关联知识不判冗余，未关联能力不判课程缺陷。测试实际结果见上级验证结果.txt。所有比例如后续计算，应分别以17能力或60知识点为分母；本轮不计算学生达成率。\n\n本次形成的是可追溯、待人工审查的专业能力图谱候选关系数据，不代表图谱已经完成专家验证，不代表课程教学效果或学生能力达成度已经得到证明。\n')
 return nodes,edges,evidence,reviews,summary,c,jobs,hs
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--output',type=Path);a=p.parse_args();print(json.dumps(run(a.output)[4],ensure_ascii=False,indent=2))
